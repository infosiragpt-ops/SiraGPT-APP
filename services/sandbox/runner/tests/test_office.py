"""Office deliverables: Word, Excel, PowerPoint, OpenDocument and their converters.

Every test builds a real file with one library and reopens it with an
independent reader (another library, LibreOffice or pandoc) and asserts on the
content. Spanish text with accents and ñ proves Unicode survives each
round-trip. Heavy imports live inside the tests so collection works anywhere.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

TITULO = "Informe anual de ventas — Región Andina"
PARRAFO = "La facturación creció un 12 % en Perú, Colombia y España durante el año."
FILAS = [("Producto", "Unidades", "Precio"), ("Café molido", 120, 18.5), ("Té de jazmín", 80, 9.75)]


def _write_docx(path: Path, titulo: str = TITULO, parrafo: str = PARRAFO) -> Path:
    """Helper shared by several tests: a .docx with heading, paragraph and table."""
    import docx

    document = docx.Document()
    document.add_heading(titulo, level=1)
    document.add_paragraph(parrafo)
    table = document.add_table(rows=0, cols=3)
    table.style = "Table Grid"
    for row in FILAS:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    document.save(str(path))
    return path


def _write_xlsx(path: Path) -> Path:
    """Helper: openpyxl workbook with a header, numbers and a SUM formula."""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Ventas"
    sheet.append(["Producto", "Unidades", "Precio"])
    sheet.append(["Café molido", 120, 18.5])
    sheet.append(["Té de jazmín", 80, 9.75])
    sheet["A4"] = "Total unidades"
    sheet["B4"] = "=SUM(B2:B3)"
    workbook.save(str(path))
    return path


# ----------------------------------------------------------------------------- Word


@pytest.mark.covers("docx", "docx2txt")
def test_docx_write_then_plain_text(out, check):
    """python-docx writes heading + paragraph + table; docx2txt reads them back as text."""
    import docx
    import docx2txt

    path = _write_docx(out / "informe.docx")
    check.nonempty(path, 5000)

    text = docx2txt.process(str(path))
    assert TITULO in text
    assert "Perú" in text and "Té de jazmín" in text

    again = docx.Document(str(path))
    assert again.paragraphs[0].text == TITULO
    assert len(again.tables) == 1
    assert again.tables[0].cell(2, 0).text == "Té de jazmín"
    assert again.tables[0].cell(1, 1).text == "120"


@pytest.mark.covers("docxtpl", "docx")
def test_docxtpl_fills_jinja_template(out):
    """docxtpl renders a Word template with {{ tags }} and a row loop; python-docx verifies the result."""
    import docx
    from docxtpl import DocxTemplate

    template = docx.Document()
    template.add_heading("Carta para {{ nombre }}", level=1)
    template.add_paragraph("Estimado/a {{ nombre }}, su saldo es {{ saldo }} €.")
    table = template.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Concepto"
    table.cell(0, 1).text = "Importe"
    table.cell(1, 0).text = "{%tr for item in items %}"  # a row-level tag replaces its whole row
    table.cell(2, 0).text = "{{ item.concepto }}"
    table.cell(2, 1).text = "{{ item.importe }}"
    table.cell(3, 0).text = "{%tr endfor %}"
    template_path = out / "plantilla.docx"
    template.save(str(template_path))

    rendered_path = out / "carta.docx"
    tpl = DocxTemplate(str(template_path))
    tpl.render({
        "nombre": "Señora Muñoz",
        "saldo": "1.250,50",
        "items": [{"concepto": "Matrícula", "importe": "900"}, {"concepto": "Libros", "importe": "350,50"}],
    })
    tpl.save(str(rendered_path))

    result = docx.Document(str(rendered_path))
    full_text = "\n".join(p.text for p in result.paragraphs)
    assert "Carta para Señora Muñoz" in full_text
    assert "su saldo es 1.250,50 €" in full_text
    assert "{{" not in full_text and "{%" not in full_text
    rows = result.tables[0].rows
    assert len(rows) == 3, "the {%tr for %} loop must expand into one row per item"
    assert rows[1].cells[0].text == "Matrícula"
    assert rows[2].cells[1].text == "350,50"


@pytest.mark.covers("docxcompose", "docx")
def test_docxcompose_merges_two_documents(out):
    """docxcompose appends one .docx to another; the merged file keeps both texts and tables."""
    import docx
    from docxcompose.composer import Composer

    first = _write_docx(out / "uno.docx", titulo="Primera parte", parrafo="Introducción con acentos: á é í ó ú ñ.")
    second = _write_docx(out / "dos.docx", titulo="Segunda parte", parrafo="Conclusión del informe.")

    master = docx.Document(str(first))
    composer = Composer(master)
    composer.append(docx.Document(str(second)))
    merged_path = out / "fusionado.docx"
    composer.save(str(merged_path))

    merged = docx.Document(str(merged_path))
    text = "\n".join(p.text for p in merged.paragraphs)
    assert "Primera parte" in text and "Segunda parte" in text
    assert text.index("Primera parte") < text.index("Segunda parte")
    assert "á é í ó ú ñ" in text
    assert len(merged.tables) == 2


@pytest.mark.covers("mammoth", "docx")
def test_mammoth_converts_docx_to_html(out):
    """mammoth turns a .docx into semantic HTML (h1, p, table) keeping the Spanish text."""
    import mammoth

    path = _write_docx(out / "para_html.docx")
    with path.open("rb") as handle:
        result = mammoth.convert_to_html(handle)
    html = result.value
    assert f"<h1>{TITULO}</h1>" in html
    assert "<table>" in html and "<td><p>Té de jazmín</p></td>" in html
    assert "Perú" in html

    with path.open("rb") as handle:
        raw = mammoth.extract_raw_text(handle).value
    assert PARRAFO in raw

    html_path = out / "informe.html"
    html_path.write_text(html, encoding="utf-8")
    assert html_path.stat().st_size > 200


@pytest.mark.covers("lxml")
def test_lxml_parses_docx_xml_and_validates_schema(out):
    """lxml reads the WordprocessingML inside a .docx with XPath and validates XML against an XSD."""
    from lxml import etree

    path = _write_docx(out / "xml.docx")
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    texts = root.xpath("//w:t/text()", namespaces=ns)
    assert TITULO in texts
    assert len(root.xpath("//w:tbl//w:tr", namespaces=ns)) == len(FILAS)

    # Build, serialise and validate an XML document of our own.
    schema_doc = etree.fromstring(
        b"""<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema">
          <xs:element name="pedido">
            <xs:complexType><xs:sequence>
              <xs:element name="cliente" type="xs:string"/>
              <xs:element name="total" type="xs:decimal"/>
            </xs:sequence></xs:complexType>
          </xs:element>
        </xs:schema>"""
    )
    schema = etree.XMLSchema(schema_doc)
    pedido = etree.Element("pedido")
    etree.SubElement(pedido, "cliente").text = "Señora Ibáñez"
    etree.SubElement(pedido, "total").text = "1250.50"
    xml_path = out / "pedido.xml"
    etree.ElementTree(pedido).write(str(xml_path), encoding="utf-8", xml_declaration=True, pretty_print=True)
    parsed = etree.parse(str(xml_path))
    schema.assertValid(parsed)
    assert parsed.findtext("cliente") == "Señora Ibáñez"

    bad = etree.fromstring(b"<pedido><cliente>x</cliente><total>no-es-numero</total></pedido>")
    assert not schema.validate(bad)


@pytest.mark.covers("soffice", "docx")
def test_soffice_converts_docx_to_pdf(out, check):
    """LibreOffice headless renders a .docx to PDF whose text layer keeps the accents."""
    path = _write_docx(out / "render.docx")
    check.soffice(["--convert-to", "pdf", "--outdir", str(out), str(path)], out)
    pdf = out / "render.pdf"
    assert pdf.exists(), "soffice produced no PDF"
    assert check.pdf_pages(pdf) == 1
    text = check.pdf_text(pdf)
    assert "Región Andina" in text
    assert "Té de jazmín" in text


# ---------------------------------------------------------------------------- Excel


@pytest.mark.covers("openpyxl", "xlsxwriter")
def test_xlsxwriter_writes_rich_workbook_openpyxl_reads_it(out, check):
    """xlsxwriter builds a styled sheet with formula, conditional format and chart; openpyxl reads it back."""
    import openpyxl
    import xlsxwriter

    path = out / "ventas.xlsx"
    workbook = xlsxwriter.Workbook(str(path))
    sheet = workbook.add_worksheet("Año 2026")
    bold = workbook.add_format({"bold": True, "bg_color": "#DDEBF7"})
    money = workbook.add_format({"num_format": "#,##0.00 €"})
    sheet.write_row(0, 0, ["Mes", "Ingresos", "Gastos", "Margen"], bold)
    datos = [("Enero", 1200.5, 800), ("Febrero", 1350, 910.25), ("Marzo", 990, 1020)]
    for index, (mes, ingresos, gastos) in enumerate(datos, start=1):
        sheet.write(index, 0, mes)
        sheet.write_number(index, 1, ingresos, money)
        sheet.write_number(index, 2, gastos, money)
        sheet.write_formula(index, 3, f"=B{index + 1}-C{index + 1}", money, ingresos - gastos)
    sheet.write(4, 0, "Total", bold)
    sheet.write_formula(4, 1, "=SUM(B2:B4)", money, 3540.5)
    sheet.conditional_format("D2:D4", {"type": "cell", "criteria": "<", "value": 0, "format": workbook.add_format({"font_color": "#9C0006"})})
    chart = workbook.add_chart({"type": "column"})
    chart.add_series({"name": "Ingresos", "categories": "='Año 2026'!$A$2:$A$4", "values": "='Año 2026'!$B$2:$B$4"})
    chart.set_title({"name": "Ingresos por mes"})
    sheet.insert_chart("F2", chart)
    sheet.set_column(0, 3, 14)
    workbook.close()
    check.nonempty(path, 5000)

    loaded = openpyxl.load_workbook(str(path))
    assert loaded.sheetnames == ["Año 2026"]
    ws = loaded["Año 2026"]
    assert ws["A1"].value == "Mes" and ws["A1"].font.bold
    assert ws["B2"].value == 1200.5
    assert ws["D2"].value == "=B2-C2"
    assert ws["B5"].value == "=SUM(B2:B4)"
    assert ws["B2"].number_format == "#,##0.00 €"
    assert len(ws.conditional_formatting) == 1
    assert len(ws._charts) == 1

    cached = openpyxl.load_workbook(str(path), data_only=True)["Año 2026"]
    assert cached["B5"].value == 3540.5
    assert cached["D4"].value == pytest.approx(-30)


@pytest.mark.covers("formulas", "openpyxl")
def test_formulas_evaluates_excel_formulas_without_excel(out):
    """formulas loads an openpyxl workbook and computes SUM/IF/AVERAGE results itself."""
    import formulas
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    sheet.append(["Región", "Ventas"])
    sheet.append(["Norte", 100])
    sheet.append(["Sur", 250])
    sheet.append(["Centro", 50])
    sheet["D1"] = "=SUM(B2:B4)"
    sheet["D2"] = "=AVERAGE(B2:B4)"
    sheet["D3"] = '=IF(D1>300,"objetivo cumplido","por debajo")'
    sheet["D4"] = "=MAX(B2:B4)-MIN(B2:B4)"
    path = out / "modelo.xlsx"
    workbook.save(str(path))

    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()

    def value(result, cell: str):
        key = f"'[{path.name}]DATOS'!{cell}".upper()
        matches = [k for k in result if k.upper() == key]
        assert matches, f"{key} missing from {sorted(result)[:8]}"
        return result[matches[0]].value[0, 0]

    assert value(solution, "D1") == 400
    assert value(solution, "D2") == pytest.approx(400 / 3)
    assert value(solution, "D3") == "objetivo cumplido"
    assert value(solution, "D4") == 200

    # Change an input and recompute: the engine, not cached values, produces the result.
    changed = model.calculate(inputs={f"'[{path.name}]DATOS'!B3": 10})
    assert value(changed, "D1") == 160
    assert value(changed, "D3") == "por debajo"


@pytest.mark.covers("xlrd", "soffice", "openpyxl")
def test_xlrd_reads_legacy_xls_made_by_libreoffice(out, check):
    """openpyxl writes .xlsx, LibreOffice converts it to legacy .xls, xlrd reads cells and formula results."""
    import xlrd

    xlsx = _write_xlsx(out / "legado.xlsx")
    check.soffice(["--convert-to", "xls", "--outdir", str(out), str(xlsx)], out)
    xls = out / "legado.xls"
    assert xls.exists(), "soffice did not write an .xls"
    check.nonempty(xls, 2000)

    book = xlrd.open_workbook(str(xls))
    assert book.sheet_names() == ["Ventas"]
    sheet = book.sheet_by_index(0)
    assert sheet.nrows == 4 and sheet.ncols == 3
    assert sheet.cell_value(0, 0) == "Producto"
    assert sheet.cell_value(2, 0) == "Té de jazmín"
    assert sheet.cell_value(1, 1) == 120
    assert sheet.cell_value(2, 2) == 9.75
    assert sheet.cell_value(3, 1) == 200, "SUM(B2:B3) must be recalculated by LibreOffice on export"


@pytest.mark.covers("pyxlsb", "soffice", "openpyxl")
def test_pyxlsb_reads_binary_xlsb_made_by_libreoffice(out, check):
    """openpyxl writes .xlsx, LibreOffice exports Excel 2007 Binary .xlsb, pyxlsb reads the rows."""
    import pyxlsb

    xlsx = _write_xlsx(out / "binario.xlsx")
    try:
        check.soffice(["--convert-to", "xlsb:Calc MS Excel 2007 Binary", "--outdir", str(out), str(xlsx)], out)
    except AssertionError as exc:
        pytest.skip(f"this LibreOffice cannot export the xlsb filter: {exc}")
    xlsb = out / "binario.xlsb"
    if not xlsb.exists():
        pytest.skip("LibreOffice accepted the command but wrote no .xlsb (export filter unavailable)")
    check.nonempty(xlsb, 1000)

    with pyxlsb.open_workbook(str(xlsb)) as book:
        assert "Ventas" in book.sheets
        with book.get_sheet("Ventas") as sheet:
            rows = [[cell.v for cell in row] for row in sheet.rows()]
    assert rows[0][:3] == ["Producto", "Unidades", "Precio"]
    assert rows[2][0] == "Té de jazmín"
    assert rows[1][1] == 120
    assert rows[2][2] == 9.75
    assert rows[3][1] == 200


# ----------------------------------------------------------------------- PowerPoint


@pytest.mark.covers("pptx", "soffice")
def test_pptx_deck_with_table_and_picture_renders_to_pdf(out, check):
    """python-pptx builds a 3-slide deck (title, bullets, table, picture); reopened with python-pptx and rendered by LibreOffice."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches, Pt

    picture = out / "logo.png"
    Image.new("RGB", (200, 120), (30, 90, 160)).save(picture)

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    title_slide = deck.slides.add_slide(deck.slide_layouts[0])
    title_slide.shapes.title.text = "Resultados del año"
    title_slide.placeholders[1].text = "Dirección de Operaciones · Perú"

    bullets = deck.slides.add_slide(deck.slide_layouts[1])
    bullets.shapes.title.text = "Puntos clave"
    frame = bullets.placeholders[1].text_frame
    frame.text = "Crecimiento del 12 %"
    for line in ("Expansión a España", "Nuevo almacén en Medellín"):
        paragraph = frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 1
        paragraph.font.size = Pt(20)

    data_slide = deck.slides.add_slide(deck.slide_layouts[5])
    data_slide.shapes.title.text = "Ventas por producto"
    table_shape = data_slide.shapes.add_table(len(FILAS), 3, Inches(0.7), Inches(1.6), Inches(7), Inches(1.8))
    for r, row in enumerate(FILAS):
        for c, value in enumerate(row):
            table_shape.table.cell(r, c).text = str(value)
    data_slide.shapes.add_picture(str(picture), Inches(9), Inches(2), width=Inches(3))

    path = out / "resultados.pptx"
    deck.save(str(path))
    check.nonempty(path, 20000)

    again = Presentation(str(path))
    assert len(again.slides) == 3
    assert again.slides[0].shapes.title.text == "Resultados del año"
    tables = [s for s in again.slides[2].shapes if s.has_table]
    pictures = [s for s in again.slides[2].shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
    assert len(tables) == 1 and len(pictures) == 1
    assert tables[0].table.cell(2, 0).text == "Té de jazmín"

    check.soffice(["--convert-to", "pdf", "--outdir", str(out), str(path)], out)
    pdf = out / "resultados.pdf"
    assert check.pdf_pages(pdf) == 3
    text = check.pdf_text(pdf)
    assert "Resultados del año" in text
    assert "Medellín" in text and "Té de jazmín" in text


# --------------------------------------------------------------------- OpenDocument


@pytest.mark.covers("odf", "soffice", "docx")
def test_odfpy_writes_odt_that_libreoffice_converts_to_docx(out, check):
    """odfpy writes an .odt with heading, paragraph and table; LibreOffice converts it to .docx read by python-docx."""
    import docx
    from odf import teletype
    from odf.opendocument import OpenDocumentText, load
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import H, P

    document = OpenDocumentText()
    document.text.addElement(H(outlinelevel=1, text="Acta de reunión"))
    document.text.addElement(P(text="Asistieron: Ángela, Iñaki y José María."))
    table = Table(name="Acuerdos")
    table.addElement(TableColumn(numbercolumnsrepeated=2))
    for fila in (("Tema", "Responsable"), ("Presupuesto", "Ángela"), ("Logística", "Iñaki")):
        row = TableRow()
        for value in fila:
            cell = TableCell()
            cell.addElement(P(text=value))
            row.addElement(cell)
        table.addElement(row)
    document.text.addElement(table)
    odt = out / "acta.odt"
    document.save(str(odt))
    check.nonempty(odt, 1000)

    reloaded = load(str(odt))
    paragraphs = [teletype.extractText(p) for p in reloaded.getElementsByType(P)]
    assert "Asistieron: Ángela, Iñaki y José María." in paragraphs
    assert len(reloaded.getElementsByType(TableRow)) == 3

    check.soffice(["--convert-to", "docx", "--outdir", str(out), str(odt)], out)
    converted = docx.Document(str(out / "acta.docx"))
    text = "\n".join(p.text for p in converted.paragraphs)
    assert "Acta de reunión" in text and "José María" in text
    assert len(converted.tables) == 1
    assert converted.tables[0].cell(2, 1).text == "Iñaki"


# ------------------------------------------------------------------- Converters


@pytest.mark.covers("pdf2docx", "docx")
def test_pdf2docx_converts_pdf_text_and_table_to_word(out):
    """pdf2docx turns a one-page PDF (paragraph + ruled table) into an editable .docx with the same text."""
    import docx
    from pdf2docx import Converter
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    pdf = out / "origen.pdf"
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Informe de inventario", styles["Heading1"]),
        Paragraph("Existencias al cierre del trimestre en el almacén de Málaga.", styles["Normal"]),
        Spacer(1, 12),
    ]
    table = Table([list(map(str, row)) for row in FILAS])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.75, colors.black)]))
    story.append(table)
    SimpleDocTemplate(str(pdf), pagesize=A4).build(story)

    target = out / "convertido.docx"
    converter = Converter(str(pdf))
    try:
        converter.convert(str(target), start=0, end=1)
    finally:
        converter.close()

    result = docx.Document(str(target))
    paragraphs = "\n".join(p.text for p in result.paragraphs)
    cells = " ".join(cell.text for table in result.tables for row in table.rows for cell in row.cells)
    everything = paragraphs + "\n" + cells
    assert "Informe de inventario" in everything
    assert "Málaga" in everything
    assert "Unidades" in everything and "120" in everything
    assert result.tables, "the ruled table should become a Word table"


@pytest.mark.covers("pypandoc", "pandoc", "docx")
def test_pypandoc_markdown_to_docx_and_pandoc_cli_back_to_markdown(out, check):
    """pypandoc converts Markdown (heading, list, table) to .docx; the pandoc CLI converts it back to Markdown."""
    import docx
    import pypandoc

    markdown = (
        "# Plan de capacitación\n\n"
        "Objetivos para el **próximo semestre**:\n\n"
        "- Formación en seguridad\n- Taller de redacción\n\n"
        "| Módulo | Horas |\n|---|---|\n| Introducción | 4 |\n| Práctica guiada | 12 |\n"
    )
    target = out / "plan.docx"
    pypandoc.convert_text(markdown, "docx", format="md", outputfile=str(target))
    check.nonempty(target, 5000)

    document = docx.Document(str(target))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Plan de capacitación" in text
    assert "próximo semestre" in text
    assert "Taller de redacción" in text
    assert len(document.tables) == 1
    assert document.tables[0].cell(2, 0).text == "Práctica guiada"

    done = check.run(["pandoc", str(target), "-t", "gfm", "--wrap=none"], timeout=60)
    assert "# Plan de capacitación" in done.stdout
    assert "**próximo semestre**" in done.stdout
    assert "| Práctica guiada | 12 |" in done.stdout.replace("  ", " ") or "Práctica guiada" in done.stdout
    assert pypandoc.get_pandoc_version().split(".")[0].isdigit()
