"""Areas "dev" and "files", plus the Node.js toolchain.

Developer tooling (pytest, notebooks, CLIs, HTTP frameworks on loopback only),
archive and type-detection libraries, the zip/7z/ruff/jupyter commands, and the
global Node modules: each JS script writes an Office/PDF/zip file that a Python
library then reads back, so both sides of the bridge are proven.
"""
import http.server
import io
import json
import os
import socket
import subprocess
import sys
import textwrap
import threading
import zipfile
from pathlib import Path

import pytest

PNG_1X1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478"
    "9c6360f8cfc00000020001e221bc330000000049454e44ae426082")


# --------------------------------------------------------------------------- helpers
def free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


@pytest.fixture
def no_proxy(monkeypatch):
    """Loopback requests must never be sent through an outbound proxy."""
    for var in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "http_proxy", "https_proxy", "all_proxy"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("NO_PROXY", "127.0.0.1,localhost")
    monkeypatch.setenv("no_proxy", "127.0.0.1,localhost")


class _Handler(http.server.BaseHTTPRequestHandler):
    def do_GET(self):  # noqa: N802 - http.server API
        body = json.dumps({"path": self.path, "saludo": "hola"}).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):  # silence
        pass


@pytest.fixture
def loopback_server():
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_address[1]}"
    finally:
        server.shutdown()
        server.server_close()


def make_pdf(path):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as handle:
        writer.write(handle)
    return path


def node_env():
    """NODE_PATH as the image sets it; on a workstation also the repo's node_modules."""
    env = dict(os.environ)
    paths = [env.get("NODE_PATH") or "/usr/local/lib/node_modules", os.path.join(os.getcwd(), "node_modules")]
    env["NODE_PATH"] = os.pathsep.join(dict.fromkeys(p for p in paths if p))
    return env


def run_node(check, script_path, *args, timeout=120):
    return check.run(["node", str(script_path), *map(str, args)], timeout=timeout, env=node_env())


def build_notebook():
    import nbformat
    nb = nbformat.v4.new_notebook()
    nb.cells = [
        nbformat.v4.new_markdown_cell("# Informe de prueba\n\nSuma de los primeros números."),
        nbformat.v4.new_code_cell("total = sum(range(10))\nprint('total =', total)\ntotal"),
    ]
    nb.metadata["kernelspec"] = {"name": "python3", "display_name": "Python 3", "language": "python"}
    return nb


# =========================================================================== dev
@pytest.mark.covers("pytest")
def test_pytest_runs_generated_suites_in_a_subprocess(out):
    passing = out / "test_ok.py"
    passing.write_text(textwrap.dedent("""
        import pytest

        @pytest.mark.parametrize("a,b", [(1, 2), (2, 3)])
        def test_sum(a, b):
            assert a + 1 == b
    """), encoding="utf-8")
    failing = out / "test_bad.py"
    failing.write_text("def test_wrong():\n    assert 1 + 1 == 3\n", encoding="utf-8")
    base = [sys.executable, "-m", "pytest", "-q", "-p", "no:cacheprovider", "--rootdir", str(out), "-c", os.devnull]
    good = subprocess.run([*base, str(passing)], capture_output=True, text=True, timeout=60, check=False)
    assert good.returncode == 0, good.stdout[-800:] + good.stderr[-400:]
    assert "2 passed" in good.stdout
    bad = subprocess.run([*base, str(failing)], capture_output=True, text=True, timeout=60, check=False)
    assert bad.returncode == 1, bad.stdout[-800:]
    assert "1 failed" in bad.stdout and "assert (1 + 1) == 3" in bad.stdout


@pytest.mark.covers("nbformat", "nbclient", "ipykernel", "nbconvert", "jupytext")
def test_notebook_is_executed_offline_and_converted(out, check):
    import nbformat
    from nbclient import NotebookClient
    from nbconvert import HTMLExporter, MarkdownExporter
    import jupytext
    nb = build_notebook()
    path = out / "informe.ipynb"
    nbformat.write(nb, str(path))
    check.nonempty(path, 200)
    loaded = nbformat.read(str(path), as_version=4)
    nbformat.validate(loaded)
    assert len(loaded.cells) == 2
    client = NotebookClient(loaded, kernel_name="python3", timeout=60, allow_errors=False,
                            resources={"metadata": {"path": str(out)}})
    executed = client.execute()
    outputs = executed.cells[1].outputs
    streams = "".join(o.get("text", "") for o in outputs if o.get("output_type") == "stream")
    assert "total = 45" in streams
    results = [o for o in outputs if o.get("output_type") == "execute_result"]
    assert results and results[0]["data"]["text/plain"] == "45"
    assert executed.cells[1].execution_count == 1
    nbformat.write(executed, str(out / "informe.ejecutado.ipynb"))
    html, _ = HTMLExporter().from_notebook_node(executed)
    assert "<html" in html.lower() and "total = 45" in html and "Informe de prueba" in html
    (out / "informe.html").write_text(html, encoding="utf-8")
    markdown, _ = MarkdownExporter().from_notebook_node(executed)
    assert "# Informe de prueba" in markdown and "total = sum(range(10))" in markdown and "45" in markdown
    # jupytext round trips: .py (percent format) and .md.
    py_path = out / "informe.py"
    jupytext.write(executed, str(py_path), fmt="py:percent")
    py_text = py_path.read_text(encoding="utf-8")
    assert "# %% [markdown]" in py_text and "total = sum(range(10))" in py_text
    from_py = jupytext.read(str(py_path))
    assert [c.cell_type for c in from_py.cells] == ["markdown", "code"]
    assert from_py.cells[1].source.strip() == loaded.cells[1].source.strip()
    md_path = out / "informe.md"
    jupytext.write(executed, str(md_path), fmt="md")
    from_md = jupytext.read(str(md_path))
    assert from_md.cells[1].source.strip().startswith("total = sum(range(10))")


@pytest.mark.covers("jupyter")
def test_jupyter_cli_executes_and_converts_a_notebook(out, check):
    import nbformat
    path = out / "cli.ipynb"
    nbformat.write(build_notebook(), str(path))
    check.run(["jupyter", "nbconvert", "--to", "html", "--execute", "--ExecutePreprocessor.timeout=60",
               "--output", "cli_ejecutado", str(path)], timeout=120, cwd=str(out))
    html = (out / "cli_ejecutado.html").read_text(encoding="utf-8")
    assert "total = 45" in html and "Informe de prueba" in html
    check.run(["jupyter", "nbconvert", "--to", "markdown", "--output", "cli", str(path)], timeout=90, cwd=str(out))
    md = (out / "cli.md").read_text(encoding="utf-8")
    assert "# Informe de prueba" in md and "sum(range(10))" in md
    listing = check.run(["jupyter", "kernelspec", "list", "--json"], timeout=60).stdout
    assert "python3" in json.loads(listing)["kernelspecs"]


@pytest.mark.covers("rich")
def test_rich_renders_table_and_markdown_to_text():
    from rich.console import Console
    from rich.markdown import Markdown
    from rich.table import Table
    console = Console(record=True, width=80, force_terminal=False, color_system=None)
    table = Table(title="Ventas")
    table.add_column("Producto")
    table.add_column("Total", justify="right")
    table.add_row("Laptop", "1.200")
    table.add_row("Mouse", "25")
    console.print(table)
    console.print(Markdown("# Resumen\n\n- uno\n- dos"))
    text = console.export_text()
    assert "Ventas" in text and "Laptop" in text and "1.200" in text and "Mouse" in text
    assert "Resumen" in text and "uno" in text
    html = Console(record=True, width=60).export_html() or ""
    assert isinstance(html, str)
    svg_console = Console(record=True, width=40)
    svg_console.print("[bold]hola[/bold]")
    assert svg_console.export_svg().startswith("<svg")


@pytest.mark.covers("tqdm")
def test_tqdm_reports_progress_to_a_stream():
    from tqdm import tqdm
    sink = io.StringIO()
    total = 0
    for n in tqdm(range(50), file=sink, ascii=True, mininterval=0, desc="procesando"):
        total += n
    assert total == 1225
    report = sink.getvalue()
    assert "procesando" in report and "100%" in report and "50/50" in report
    from tqdm.contrib import tenumerate
    assert [i for i, _ in tenumerate(["a", "b"], file=io.StringIO())] == [0, 1]


@pytest.mark.covers("click")
def test_click_command_with_options_and_arguments():
    import click
    from click.testing import CliRunner

    @click.command()
    @click.argument("nombre")
    @click.option("--veces", type=int, default=1, show_default=True)
    @click.option("--grito/--no-grito", default=False)
    def saludar(nombre, veces, grito):
        texto = f"Hola {nombre}" * veces
        click.echo(texto.upper() if grito else texto)

    runner = CliRunner()
    result = runner.invoke(saludar, ["Ana", "--veces", "2", "--grito"])
    assert result.exit_code == 0 and result.output == "HOLA ANAHOLA ANA\n"
    assert runner.invoke(saludar, ["--help"]).output.count("--veces") == 1
    bad = runner.invoke(saludar, ["Ana", "--veces", "dos"])
    assert bad.exit_code == 2 and "Invalid value" in bad.output


@pytest.mark.covers("typer")
def test_typer_app_with_subcommands():
    import typer
    from typer.testing import CliRunner
    app = typer.Typer()

    @app.command()
    def sumar(a: int, b: int = typer.Option(0, help="segundo sumando")):
        typer.echo(a + b)

    @app.command()
    def saludar(nombre: str, formal: bool = False):
        typer.echo(("Buenos días, " if formal else "Hola ") + nombre)

    runner = CliRunner()
    assert runner.invoke(app, ["sumar", "2", "--b", "40"]).output.strip() == "42"
    assert runner.invoke(app, ["saludar", "Ana", "--formal"]).output.strip() == "Buenos días, Ana"
    bad = runner.invoke(app, ["sumar", "x"])
    assert bad.exit_code != 0


@pytest.mark.covers("attr")
def test_attrs_classes_validate_and_serialise():
    import attr
    import attrs

    @attr.s(auto_attribs=True, frozen=True)
    class Factura:
        numero: str
        total: float = attr.ib(validator=attr.validators.instance_of(float))
        items: list = attr.Factory(list)

    f = Factura("F-001", 120.5, ["a", "b"])
    assert attr.asdict(f) == {"numero": "F-001", "total": 120.5, "items": ["a", "b"]}
    g = attr.evolve(f, total=99.0)
    assert g.total == 99.0 and f.total == 120.5
    with pytest.raises(attr.exceptions.FrozenInstanceError):
        f.total = 1.0  # type: ignore[misc]
    with pytest.raises(TypeError):
        Factura("F-002", "no es float")

    @attrs.define(order=True)
    class Punto:
        x: int
        y: int = 0

    assert sorted([Punto(2, 1), Punto(1, 5)])[0] == Punto(1, 5)
    assert attrs.fields(Punto).y.default == 0


@pytest.mark.covers("more_itertools", "toolz")
def test_more_itertools_and_toolz_functional_helpers():
    import more_itertools as mit
    import toolz
    from toolz import curried
    assert list(mit.chunked(range(7), 3)) == [[0, 1, 2], [3, 4, 5], [6]]
    assert list(mit.windowed([1, 2, 3, 4], 2)) == [(1, 2), (2, 3), (3, 4)]
    assert mit.first([], default="vacío") == "vacío"
    assert list(mit.unique_everseen("abracadabra")) == ["a", "b", "r", "c", "d"]
    evens, odds = mit.partition(lambda n: n % 2, range(6))
    assert (list(evens), list(odds)) == ([0, 2, 4], [1, 3, 5])
    assert list(mit.flatten([[1], [2, 3]])) == [1, 2, 3]
    grouped = toolz.groupby(len, ["a", "bb", "cc", "d"])
    assert grouped == {1: ["a", "d"], 2: ["bb", "cc"]}
    pipeline = toolz.compose(sum, curried.map(lambda n: n * n), curried.filter(lambda n: n % 2))
    assert pipeline(range(6)) == 1 + 9 + 25
    assert toolz.pipe({"a": 1, "b": 2}, curried.valmap(lambda v: v * 10), curried.itemmap(lambda kv: (kv[0].upper(), kv[1]))) == {"A": 10, "B": 20}
    assert toolz.frequencies("hola") == {"h": 1, "o": 1, "l": 1, "a": 1}


@pytest.mark.covers("cloudpickle", "dill")
def test_cloudpickle_and_dill_serialise_lambdas_and_closures(out):
    import cloudpickle
    import dill
    factor = 3
    scale = lambda x: x * factor  # noqa: E731 - a lambda is the point
    blob = cloudpickle.dumps(scale)
    assert cloudpickle.loads(blob)(5) == 15
    path = out / "fn.cpkl"
    path.write_bytes(blob)
    assert cloudpickle.loads(path.read_bytes())(2) == 6

    def make_counter():
        count = {"n": 0}

        def step():
            count["n"] += 1
            return count["n"]
        return step

    counter = make_counter()
    counter()
    revived = dill.loads(dill.dumps(counter))
    assert revived() == 2  # closure state survived
    dpath = out / "session.dill"
    with dpath.open("wb") as handle:
        dill.dump({"fn": lambda s: s.upper(), "data": [1, 2]}, handle)
    with dpath.open("rb") as handle:
        state = dill.load(handle)
    assert state["fn"]("hola") == "HOLA" and state["data"] == [1, 2]


@pytest.mark.covers("psutil")
def test_psutil_reports_process_and_system_metrics(out):
    import psutil
    me = psutil.Process()
    assert me.pid == os.getpid()
    assert me.memory_info().rss > 1_000_000
    assert me.num_threads() >= 1
    assert psutil.cpu_count() >= 1
    vm = psutil.virtual_memory()
    assert vm.total > vm.available > 0
    usage = psutil.disk_usage(str(out))
    assert usage.total >= usage.used
    assert me.cpu_times().user >= 0
    assert "python" in me.name().lower() or me.exe()
    assert isinstance(psutil.pids(), list) and me.pid in psutil.pids()


@pytest.mark.covers("requests", "httpx")
def test_requests_and_httpx_fetch_from_loopback_server(no_proxy, loopback_server):
    import httpx
    import requests
    response = requests.get(f"{loopback_server}/datos?x=1", timeout=10)
    assert response.status_code == 200
    assert response.headers["Content-Type"] == "application/json"
    assert response.json() == {"path": "/datos?x=1", "saludo": "hola"}
    with requests.Session() as session:
        assert session.get(f"{loopback_server}/otra", timeout=10).json()["path"] == "/otra"
    with httpx.Client(trust_env=False, timeout=10) as client:
        res = client.get(f"{loopback_server}/httpx")
        assert res.status_code == 200 and res.json()["path"] == "/httpx"
        assert res.http_version == "HTTP/1.0" or res.http_version.startswith("HTTP/1")
    import asyncio

    async def fetch():
        async with httpx.AsyncClient(trust_env=False, timeout=10) as client:
            return (await client.get(f"{loopback_server}/async")).json()

    assert asyncio.run(fetch())["path"] == "/async"


@pytest.mark.covers("flask")
def test_flask_app_with_test_client():
    from flask import Flask, jsonify, request
    app = Flask("prueba")

    @app.get("/saludo/<nombre>")
    def saludo(nombre):
        return jsonify(mensaje=f"Hola {nombre}", veces=request.args.get("veces", 1, type=int))

    @app.post("/suma")
    def suma():
        data = request.get_json()
        return {"total": data["a"] + data["b"]}

    client = app.test_client()
    res = client.get("/saludo/Ana?veces=3")
    assert res.status_code == 200 and res.get_json() == {"mensaje": "Hola Ana", "veces": 3}
    assert client.post("/suma", json={"a": 2, "b": 5}).get_json() == {"total": 7}
    assert client.get("/no-existe").status_code == 404


def _fastapi_app():
    from fastapi import FastAPI, HTTPException
    from pydantic import BaseModel

    class Item(BaseModel):
        nombre: str
        precio: float

    app = FastAPI(title="Inventario")
    items: dict[int, Item] = {}

    @app.post("/items/{item_id}", status_code=201)
    def crear(item_id: int, item: Item):
        items[item_id] = item
        return {"id": item_id, **item.model_dump()}

    @app.get("/items/{item_id}")
    def leer(item_id: int):
        if item_id not in items:
            raise HTTPException(404, "no existe")
        return items[item_id]

    @app.get("/salud")
    def salud():
        return {"ok": True}

    return app


@pytest.mark.covers("fastapi")
def test_fastapi_api_with_test_client():
    from fastapi.testclient import TestClient
    client = TestClient(_fastapi_app())
    created = client.post("/items/1", json={"nombre": "Laptop", "precio": 1200.5})
    assert created.status_code == 201 and created.json() == {"id": 1, "nombre": "Laptop", "precio": 1200.5}
    assert client.get("/items/1").json()["nombre"] == "Laptop"
    assert client.get("/items/9").status_code == 404
    invalid = client.post("/items/2", json={"nombre": "x"})
    assert invalid.status_code == 422 and invalid.json()["detail"][0]["loc"] == ["body", "precio"]
    schema = client.get("/openapi.json").json()
    assert schema["info"]["title"] == "Inventario" and "/items/{item_id}" in schema["paths"]


@pytest.mark.covers("uvicorn")
def test_uvicorn_serves_fastapi_on_loopback(no_proxy):
    import time

    import httpx
    import uvicorn
    port = free_port()
    config = uvicorn.Config(_fastapi_app(), host="127.0.0.1", port=port, log_level="warning", lifespan="off")
    assert config.load() is None and config.loaded
    server = uvicorn.Server(config)
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    deadline = time.time() + 20
    while not server.started and time.time() < deadline:
        time.sleep(0.05)
    assert server.started, "uvicorn did not start within 20 s"
    try:
        with httpx.Client(trust_env=False, timeout=10) as client:
            res = client.get(f"http://127.0.0.1:{port}/salud")
            assert res.status_code == 200 and res.json() == {"ok": True}
            assert "uvicorn" in res.headers.get("server", "")
            assert client.get(f"http://127.0.0.1:{port}/items/5").status_code == 404
    finally:
        server.should_exit = True
        thread.join(timeout=15)
    assert not thread.is_alive()


@pytest.mark.covers("colorama")
def test_colorama_codes_and_stripping():
    import colorama
    from colorama import Back, Fore, Style
    assert Fore.RED == "\x1b[31m" and Back.GREEN == "\x1b[42m" and Style.RESET_ALL == "\x1b[0m"
    coloured = Fore.RED + "error" + Style.RESET_ALL + " listo"
    sink = io.StringIO()
    colorama.AnsiToWin32(sink, strip=True, convert=False).write(coloured)
    assert sink.getvalue() == "error listo"
    colorama.init(strip=True, autoreset=True)
    try:
        assert "error" in coloured
    finally:
        colorama.deinit()
    assert colorama.ansi.clear_line() == "\x1b[2K"


# =========================================================================== files
@pytest.mark.covers("magic")
def test_python_magic_detects_mime_from_buffer_and_file(out):
    import magic
    pdf = make_pdf(out / "hoja.pdf")
    assert magic.from_buffer(PNG_1X1, mime=True) == "image/png"
    assert magic.from_file(str(pdf), mime=True) == "application/pdf"
    assert "PDF document" in magic.from_file(str(pdf))
    text = out / "datos.csv"
    text.write_text("a,b\n1,2\n", encoding="utf-8")
    assert magic.from_file(str(text), mime=True) in ("text/csv", "text/plain")
    assert magic.Magic(mime_encoding=True).from_buffer("hola".encode("utf-8")) in ("us-ascii", "utf-8")


@pytest.mark.covers("filetype")
def test_filetype_guesses_from_bytes_and_paths(out):
    import filetype
    kind = filetype.guess(PNG_1X1)
    assert kind.mime == "image/png" and kind.extension == "png"
    pdf = make_pdf(out / "hoja.pdf")
    assert filetype.guess(str(pdf)).mime == "application/pdf"
    assert filetype.is_image(PNG_1X1) and not filetype.is_image(str(pdf))
    assert filetype.guess(b"solo texto plano") is None
    with zipfile.ZipFile(out / "a.zip", "w") as zf:
        zf.writestr("x.txt", "x")
    assert filetype.guess(str(out / "a.zip")).mime == "application/zip"


@pytest.mark.covers("py7zr")
def test_py7zr_writes_and_reads_7z(out, check):
    import py7zr
    src = out / "src"
    src.mkdir()
    (src / "informe.txt").write_text("informe anual\n" * 100, encoding="utf-8")
    (src / "datos.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    archive = out / "paquete.7z"
    with py7zr.SevenZipFile(str(archive), "w") as zf:
        zf.writeall(str(src), "paquete")
    check.nonempty(archive, 100)
    assert archive.stat().st_size < 1300  # the repeated text compressed
    assert py7zr.is_7zfile(str(archive))
    with py7zr.SevenZipFile(str(archive), "r") as zf:
        names = set(zf.getnames())
        assert {"paquete/informe.txt", "paquete/datos.csv"} <= names
        dest = out / "dest"
        zf.extractall(str(dest))
    assert (dest / "paquete" / "informe.txt").read_text(encoding="utf-8").count("informe anual") == 100
    encrypted = out / "secreto.7z"
    with py7zr.SevenZipFile(str(encrypted), "w", password="clave") as zf:
        zf.write(str(src / "datos.csv"), "datos.csv")
    with py7zr.SevenZipFile(str(encrypted), "r", password="clave") as zf:
        zf.extract(path=str(out / "dec"), targets=["datos.csv"])
    assert (out / "dec" / "datos.csv").read_bytes() == b"a,b\n1,2\n"


@pytest.mark.covers("zstandard")
def test_zstandard_compresses_one_shot_and_streaming(out):
    import zstandard as zstd
    payload = ("línea de registro repetida %d\n" % 7).encode("utf-8") * 2000
    compressed = zstd.ZstdCompressor(level=10).compress(payload)
    assert len(compressed) < len(payload) // 20
    assert zstd.ZstdDecompressor().decompress(compressed) == payload
    path = out / "logs.zst"
    with path.open("wb") as handle, zstd.ZstdCompressor().stream_writer(handle) as writer:
        for _ in range(10):
            writer.write(payload[:5000])
    with path.open("rb") as handle, zstd.ZstdDecompressor().stream_reader(handle) as reader:
        back = reader.read()
    assert back == payload[:5000] * 10
    assert zstd.get_frame_parameters(compressed).content_size == len(payload)


@pytest.mark.covers("rarfile")
def test_rarfile_recognises_non_rar_and_can_use_7z_backend(out):
    import shutil

    import rarfile
    with zipfile.ZipFile(out / "no.zip", "w") as zf:
        zf.writestr("x.txt", "x")
    assert not rarfile.is_rarfile(str(out / "no.zip"))
    (out / "texto.rar").write_bytes(b"esto no es un rar")
    assert not rarfile.is_rarfile(str(out / "texto.rar"))
    with pytest.raises(rarfile.Error):
        rarfile.RarFile(str(out / "texto.rar"))
    # A genuine RAR5 signature is recognised even without an extractor present.
    (out / "firma.rar").write_bytes(b"Rar!\x1a\x07\x01\x00" + b"\x00" * 64)
    assert rarfile.is_rarfile(str(out / "firma.rar"))
    assert rarfile.RAR_ID == b"Rar!\x1a\x07\x00" and rarfile.RAR5_ID == b"Rar!\x1a\x07\x01\x00"
    if shutil.which("7z"):
        tool = rarfile.tool_setup(unrar=False, unar=False, bsdtar=False, sevenzip=True, force=True)
        assert tool is not None
        cmdline = tool.get_cmdline("open_cmd", None)
        assert cmdline[0].endswith("7z") and "e" in cmdline, cmdline
    else:
        pytest.skip("7z not on PATH: rarfile's extraction backend cannot be checked here")


@pytest.mark.covers("zip")
def test_zip_cli_archives_files_that_python_reads(out, check):
    (out / "uno.txt").write_text("uno\n" * 50, encoding="utf-8")
    (out / "dos.txt").write_text("dos\n", encoding="utf-8")
    archive = out / "archivo.zip"
    done = check.run(["zip", "-j", "-q", str(archive), str(out / "uno.txt"), str(out / "dos.txt")])
    assert done.returncode == 0
    check.nonempty(archive, 100)
    with zipfile.ZipFile(archive) as zf:
        assert sorted(zf.namelist()) == ["dos.txt", "uno.txt"]
        assert zf.read("dos.txt") == b"dos\n"
        assert zf.getinfo("uno.txt").compress_size < zf.getinfo("uno.txt").file_size
        assert zf.testzip() is None
    listing = check.run(["unzip", "-l", str(archive)]).stdout if shutil_which("unzip") else ""
    assert "uno.txt" in listing or listing == ""


def shutil_which(name):
    import shutil
    return shutil.which(name)


@pytest.mark.covers("7z")
def test_7z_cli_archives_and_lists_what_py7zr_reads(out, check):
    import py7zr
    src = out / "docs"
    src.mkdir()
    (src / "a.txt").write_text("alfa\n" * 40, encoding="utf-8")
    (src / "b.txt").write_text("beta\n", encoding="utf-8")
    archive = out / "docs.7z"
    check.run(["7z", "a", "-bd", "-y", str(archive), str(src)], timeout=60)
    check.nonempty(archive, 100)
    listing = check.run(["7z", "l", "-ba", str(archive)]).stdout
    assert "a.txt" in listing and "b.txt" in listing
    with py7zr.SevenZipFile(str(archive), "r") as zf:
        names = set(zf.getnames())
        assert {"docs/a.txt", "docs/b.txt"} <= names
        zf.extract(path=str(out / "dec"), targets=["docs/b.txt"])
    assert (out / "dec" / "docs" / "b.txt").read_bytes() == b"beta\n"
    # 7z also handles zip, the other way round.
    zipped = out / "docs.zip"
    check.run(["7z", "a", "-bd", "-y", "-tzip", str(zipped), str(src / "b.txt")], timeout=60)
    with zipfile.ZipFile(zipped) as zf:
        assert zf.read("b.txt") == b"beta\n"


@pytest.mark.covers("ruff")
def test_ruff_flags_unused_import_and_checks_formatting(out, check):
    dirty = out / "sucio.py"
    dirty.write_text("import os\nimport sys\n\n\ndef f():\n    return sys.argv\n", encoding="utf-8")
    result = subprocess.run(["ruff", "check", "--no-cache", "--output-format", "concise", "--isolated", str(dirty)],
                            capture_output=True, text=True, timeout=60, check=False)
    assert result.returncode == 1, result.stdout + result.stderr
    assert "F401" in result.stdout and "`os` imported but unused" in result.stdout
    fixed = subprocess.run(["ruff", "check", "--no-cache", "--isolated", "--fix", str(dirty)],
                           capture_output=True, text=True, timeout=60, check=False)
    assert fixed.returncode == 0 and "import os" not in dirty.read_text(encoding="utf-8")
    pretty = out / "limpio.py"
    pretty.write_text('def saludar(nombre: str) -> str:\n    return f"Hola {nombre}"\n', encoding="utf-8")
    check.run(["ruff", "format", "--no-cache", "--isolated", "--check", str(pretty)], timeout=60)
    ugly = out / "feo.py"
    ugly.write_text("x = { 'a':1 ,'b':2 }\n", encoding="utf-8")
    bad = subprocess.run(["ruff", "format", "--no-cache", "--isolated", "--check", str(ugly)],
                         capture_output=True, text=True, timeout=60, check=False)
    assert bad.returncode == 1
    check.run(["ruff", "format", "--no-cache", "--isolated", str(ugly)], timeout=60)
    assert ugly.read_text(encoding="utf-8") == 'x = {"a": 1, "b": 2}\n'


# =========================================================================== node
@pytest.mark.covers("node")
def test_node_runs_a_script_with_builtin_modules(out, check):
    script = out / "hola.js"
    script.write_text(textwrap.dedent("""
        const fs = require("fs");
        const path = require("path");
        const crypto = require("crypto");
        const target = path.join(process.argv[2], "salida.json");
        const hash = crypto.createHash("sha256").update("hola").digest("hex");
        fs.writeFileSync(target, JSON.stringify({ hash, suma: [1, 2, 3].reduce((a, b) => a + b, 0), node: process.version }));
        console.log("listo");
    """), encoding="utf-8")
    done = run_node(check, script, out)
    assert done.stdout.strip() == "listo"
    data = json.loads((out / "salida.json").read_text(encoding="utf-8"))
    assert data["suma"] == 6
    assert data["hash"] == "b221d9dbb083a7f33428d7c2a3c3198ae925614d70210e28716ccaa7cd4ddb79"
    assert data["node"].startswith("v") and int(data["node"][1:].split(".")[0]) >= 18


@pytest.mark.covers("docx", "node")
def test_node_docx_writes_document_that_python_docx_reads(out, check):
    script = out / "docx.js"
    script.write_text(textwrap.dedent("""
        const fs = require("fs");
        const { Document, Packer, Paragraph, TextRun, HeadingLevel, Table, TableRow, TableCell } = require("docx");
        const doc = new Document({
          sections: [{
            children: [
              new Paragraph({ text: "Informe Sira", heading: HeadingLevel.HEADING_1 }),
              new Paragraph({ children: [new TextRun("Hola desde Node"), new TextRun({ text: " en negrita", bold: true })] }),
              new Table({ rows: [
                new TableRow({ children: [new TableCell({ children: [new Paragraph("Producto")] }), new TableCell({ children: [new Paragraph("Total")] })] }),
                new TableRow({ children: [new TableCell({ children: [new Paragraph("Laptop")] }), new TableCell({ children: [new Paragraph("1200")] })] }),
              ] }),
            ],
          }],
        });
        Packer.toBuffer(doc).then((buffer) => fs.writeFileSync(process.argv[2], buffer));
    """), encoding="utf-8")
    target = out / "informe.docx"
    run_node(check, script, target)
    check.nonempty(target, 2000)
    import docx
    document = docx.Document(str(target))
    paragraphs = [p.text for p in document.paragraphs]
    assert "Informe Sira" in paragraphs
    assert "Hola desde Node en negrita" in paragraphs
    assert document.paragraphs[0].style.name.lower().startswith("heading")
    assert any(run.bold for p in document.paragraphs for run in p.runs)
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "Laptop" and document.tables[0].cell(1, 1).text == "1200"


@pytest.mark.covers("exceljs", "node")
def test_node_exceljs_writes_workbook_that_openpyxl_reads(out, check):
    script = out / "xlsx.js"
    script.write_text(textwrap.dedent("""
        const ExcelJS = require("exceljs");
        const wb = new ExcelJS.Workbook();
        const ws = wb.addWorksheet("Ventas");
        ws.columns = [{ header: "Producto", key: "p", width: 20 }, { header: "Total", key: "t", width: 12 }];
        ws.addRow({ p: "Laptop", t: 1200 });
        ws.addRow({ p: "Mouse", t: 25 });
        ws.getCell("B4").value = { formula: "SUM(B2:B3)", result: 1225 };
        ws.getCell("B4").numFmt = "#,##0.00";
        ws.getRow(1).font = { bold: true };
        wb.xlsx.writeFile(process.argv[2]).then(() => console.log("xlsx-ok"));
    """), encoding="utf-8")
    target = out / "ventas.xlsx"
    assert "xlsx-ok" in run_node(check, script, target).stdout
    check.nonempty(target, 2000)
    import openpyxl
    wb = openpyxl.load_workbook(str(target))
    ws = wb["Ventas"]
    assert [c.value for c in ws[1]] == ["Producto", "Total"]
    assert ws["A2"].value == "Laptop" and ws["B2"].value == 1200 and ws["B3"].value == 25
    assert ws["B4"].value == "=SUM(B2:B3)"
    assert ws["B4"].number_format == "#,##0.00"
    assert ws["A1"].font.bold is True
    cached = openpyxl.load_workbook(str(target), data_only=True)["Ventas"]
    assert cached["B4"].value == 1225


@pytest.mark.covers("pptxgenjs", "node")
def test_node_pptxgenjs_writes_deck_that_python_pptx_reads(out, check):
    script = out / "pptx.js"
    script.write_text(textwrap.dedent("""
        const pptxgen = require("pptxgenjs");
        const pptx = new pptxgen();
        pptx.layout = "LAYOUT_16x9";
        const slide = pptx.addSlide();
        slide.addText("Hola Sira", { x: 1, y: 1, w: 8, h: 1, fontSize: 36, bold: true });
        slide.addText("Segunda caja", { x: 1, y: 2.5, w: 8, h: 1, fontSize: 18 });
        const second = pptx.addSlide();
        second.addText("Cierre", { x: 1, y: 1, w: 8, h: 1 });
        pptx.writeFile({ fileName: process.argv[2] }).then((name) => console.log("pptx-ok " + name));
    """), encoding="utf-8")
    target = out / "deck.pptx"
    assert "pptx-ok" in run_node(check, script, target).stdout
    check.nonempty(target, 5000)
    from pptx import Presentation
    deck = Presentation(str(target))
    assert len(deck.slides) == 2
    texts = [shape.text_frame.text for shape in deck.slides[0].shapes if shape.has_text_frame]
    assert "Hola Sira" in texts and "Segunda caja" in texts
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9, abs=0.01)
    assert any(run.font.bold for shape in deck.slides[0].shapes if shape.has_text_frame
               for p in shape.text_frame.paragraphs for run in p.runs)


@pytest.mark.covers("pdf-lib", "node")
def test_node_pdf_lib_creates_pdf_that_pypdf_reads(out, check):
    script = out / "pdf.js"
    script.write_text(textwrap.dedent("""
        const fs = require("fs");
        const { PDFDocument, StandardFonts, rgb } = require("pdf-lib");
        (async () => {
          const doc = await PDFDocument.create();
          doc.setTitle("Factura de prueba");
          const font = await doc.embedFont(StandardFonts.Helvetica);
          const page = doc.addPage([400, 300]);
          page.drawText("Factura 2026", { x: 50, y: 220, size: 24, font, color: rgb(0, 0, 0) });
          page.drawText("Total: 1200", { x: 50, y: 180, size: 14, font });
          doc.addPage([400, 300]).drawText("Anexo", { x: 50, y: 150, size: 14, font });
          fs.writeFileSync(process.argv[2], await doc.save());
          console.log("pdf-ok");
        })();
    """), encoding="utf-8")
    target = out / "factura.pdf"
    assert "pdf-ok" in run_node(check, script, target).stdout
    check.nonempty(target, 500)
    assert check.pdf_pages(target) == 2
    text = check.pdf_text(target)
    assert "Factura 2026" in text and "Total: 1200" in text and "Anexo" in text
    from pypdf import PdfReader
    reader = PdfReader(str(target))
    assert reader.metadata.title == "Factura de prueba"
    assert float(reader.pages[0].mediabox.width) == 400


@pytest.mark.covers("jszip", "node")
def test_node_jszip_creates_archive_that_zipfile_reads(out, check):
    script = out / "zip.js"
    script.write_text(textwrap.dedent("""
        const fs = require("fs");
        const JSZip = require("jszip");
        const zip = new JSZip();
        zip.file("hola.txt", "contenido jszip\\n");
        zip.folder("datos").file("tabla.csv", "a,b\\n1,2\\n");
        zip.generateAsync({ type: "nodebuffer", compression: "DEFLATE" }).then((buffer) => {
          fs.writeFileSync(process.argv[2], buffer);
          return JSZip.loadAsync(buffer);
        }).then((back) => back.file("datos/tabla.csv").async("string"))
          .then((csv) => console.log("jszip-ok " + csv.split("\\n")[0]));
    """), encoding="utf-8")
    target = out / "paquete.zip"
    assert "jszip-ok a,b" in run_node(check, script, target).stdout
    check.nonempty(target, 100)
    with zipfile.ZipFile(target) as zf:
        assert set(zf.namelist()) >= {"hola.txt", "datos/tabla.csv"}
        assert zf.read("hola.txt") == b"contenido jszip\n"
        assert zf.read("datos/tabla.csv") == b"a,b\n1,2\n"
        assert zf.testzip() is None


@pytest.mark.covers("markdown-it", "node")
def test_node_markdown_it_renders_markdown_to_html(out, check):
    """markdown-it (CommonJS, so require() works from the global modules) renders GFM tables, lists and code."""
    script = out / "md.js"
    script.write_text(textwrap.dedent("""
        const fs = require("fs");
        const MarkdownIt = require("markdown-it");
        const md = new MarkdownIt({ html: false, linkify: true });
        const src = fs.readFileSync(process.argv[2], "utf8");
        fs.writeFileSync(process.argv[3], md.render(src));
        console.log("markdown-it-ok");
    """), encoding="utf-8")
    source = out / "doc.md"
    source.write_text("# Título principal\n\nUn párrafo con **negrita** y [enlace](https://example.com).\n\n"
                      "- uno\n- dos\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n```js\nconsole.log(1)\n```\n",
                      encoding="utf-8")
    target = out / "doc.html"
    assert "markdown-it-ok" in run_node(check, script, source, target).stdout
    html = target.read_text(encoding="utf-8")
    assert "<h1>Título principal</h1>" in html
    assert "<strong>negrita</strong>" in html and 'href="https://example.com"' in html
    assert html.count("<li>") == 2
    assert "<table>" in html and "<td>2</td>" in html
    assert "<pre><code" in html and "console.log(1)" in html
