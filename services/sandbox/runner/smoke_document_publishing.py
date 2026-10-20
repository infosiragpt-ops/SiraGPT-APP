"""File-backed publishing proofs. Importing the registry needs only stdlib."""
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree


def register_proofs(proof):
    @proof("word-mailmerge")
    def word_mailmerge(root):
        from docx import Document
        from lxml import etree
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm
        from mailmerge import MailMerge, MailMergeOptions, OptionKeepFields
        from smoke_document_toolkit import _Checks

        checks = _Checks()
        source = root / "template.docx"
        target = root / "letters.docx"
        doc = Document()
        doc.sections[0].left_margin = Mm(23)
        doc.sections[0].header.paragraphs[0].text = "Encabezado que debe conservarse"
        doc.sections[0].footer.paragraphs[0].text = "Pie fijo"
        title = doc.add_paragraph("Carta personalizada")
        title.runs[0].bold = True

        def field(paragraph, name):
            simple = OxmlElement("w:fldSimple")
            simple.set(qn("w:instr"), " MERGEFIELD " + name + " ")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "«" + name + "»"
            run.append(text)
            simple.append(run)
            paragraph._p.append(simple)

        field(doc.add_paragraph(), "Recipient")
        field(doc.add_paragraph(), "Reference")
        field(doc.add_paragraph(), "Pending")
        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "Detalle"
        table.cell(0, 1).text = "Importe"
        field(table.cell(1, 0).paragraphs[0], "Item")
        field(table.cell(1, 1).paragraphs[0], "Amount")
        doc.save(source)
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        rows = [{"Item": "Edición técnica", "Amount": "12,375"},
                {"Item": "Revisión Ω", "Amount": "0,000"}]
        options = MailMergeOptions(keep_fields=OptionKeepFields.SOME)
        with MailMerge(source, options=options) as merged:
            checks.require(merged.get_merge_fields() == {
                "Recipient", "Reference", "Pending", "Item", "Amount"
            }, "template fields differ from the requested edit")
            merged.merge(Recipient="Ana Muñoz", Reference="000017")
            merged.merge_rows("Item", rows)
            merged.write(target)
        reopened = Document(target)
        paragraphs = [p.text for p in reopened.paragraphs]
        checks.require("Ana Muñoz" in paragraphs, "recipient was not merged")
        checks.require("000017" in paragraphs, "identifier lost leading zeroes")
        checks.require(reopened.paragraphs[0].runs[0].bold, "title formatting changed")
        checks.require(reopened.sections[0].left_margin == Document(source).sections[0].left_margin, "page margin changed")
        checks.require(reopened.tables[0].style.name == "Table Grid", "table style changed")
        actual_rows = [[c.text for c in r.cells] for r in reopened.tables[0].rows]
        checks.require(actual_rows == [["Detalle", "Importe"]] +
                       [[r["Item"], r["Amount"]] for r in rows], "merged rows changed")
        with MailMerge(target) as remaining:
            checks.require(remaining.get_merge_fields() == {"Pending"},
                           "an unrelated unresolved field was removed or a filled field remained")
        with ZipFile(source) as before, ZipFile(target) as after:
            for name in ("word/styles.xml", "word/header1.xml", "word/footer1.xml"):
                a, b = before.read(name), after.read(name)
                if name != "word/styles.xml":
                    # Mailmerge reserializes these parts; compare canonical XML,
                    # retaining all element attributes, text and child order.
                    a = etree.tostring(etree.fromstring(a), method="c14n")
                    b = etree.tostring(etree.fromstring(b), method="c14n")
                checks.require(a == b, name + " changed")
            xml = ElementTree.fromstring(after.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            instructions = [e.get("{%s}instr" % ns["w"], "") for e in xml.findall(".//w:fldSimple", ns)]
            checks.require(len(instructions) == 1 and "Pending" in instructions[0],
                           "unfilled merge field was not preserved in OOXML")
        checks.require(hashlib.sha256(source.read_bytes()).hexdigest() == original_hash,
                       "the source template was overwritten")
        return checks.result(source, target)

    @proof("bibliography-roundtrip")
    def bibliography_roundtrip(root):
        import bibtexparser
        import rispy
        from smoke_document_toolkit import _Checks

        checks = _Checks()
        source = root / "references.bib"
        target = root / "references-edited.bib"
        source.write_text(
            "@article{garcia2024,\n"
            " author={García, Ana and Muñoz, José},\n"
            " title={Datos y {Sira} en educación},\n"
            " year={2024}, journal={Revista de prueba},\n"
            " doi={10.0000/synthetic-a}, pages={001--017}}\n"
            "@book{lee2023,\n author={Lee, Mei},\n"
            " title={Métodos Ω}, year={2023}, publisher={Editorial sintética}}\n",
            encoding="utf8")
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

        def normalized(library):
            return {e.key: {"type": e.entry_type,
                            "fields": {f.key: f.value for f in e.fields}}
                    for e in library.entries}

        library = bibtexparser.parse_file(str(source))
        checks.require(not library.failed_blocks, "valid BibTeX was rejected")
        before = normalized(library)
        checks.require(set(before) == {"garcia2024", "lee2023"}, "reference keys were lost")
        checks.require(before["garcia2024"]["fields"]["pages"] == "001--017",
                       "page number formatting changed")
        entry = next(e for e in library.entries if e.key == "garcia2024")
        entry.fields_dict["title"].value = "Datos revisados y {Sira} en educación"
        bibtexparser.write_file(str(target), library)
        result = bibtexparser.parse_file(str(target))
        checks.require(not result.failed_blocks, "edited BibTeX cannot be read")
        expected = json.loads(json.dumps(before))
        expected["garcia2024"]["fields"]["title"] = "Datos revisados y {Sira} en educación"
        checks.require(normalized(result) == expected, "a BibTeX field outside the title changed")
        duplicate = bibtexparser.parse_string("@book{same,title={A}}\n@book{same,title={B}}")
        checks.require(bool(duplicate.failed_blocks), "duplicate reference key was silently accepted")
        malformed = bibtexparser.parse_string("@article{broken, title={unclosed")
        checks.require(bool(malformed.failed_blocks), "malformed entry was silently accepted")
        checks.require(hashlib.sha256(source.read_bytes()).hexdigest() == original_hash,
                       "the original bibliography was overwritten")

        ris_source, ris_target = root / "references.ris", root / "references-edited.ris"
        # This is a deliberately explicit field mapping, not a universal
        # BibTeX-to-RIS converter. Complex names and unknown fields need review.
        records = [
            {"type_of_reference": "JOUR", "id": "garcia2024",
             "title": expected["garcia2024"]["fields"]["title"],
             "authors": ["García, Ana", "Muñoz, José"], "year": "2024",
             "doi": "10.0000/synthetic-a", "start_page": "001", "end_page": "017"},
            {"type_of_reference": "BOOK", "id": "lee2023", "title": "Métodos Ω",
             "authors": ["Lee, Mei"], "year": "2023", "publisher": "Editorial sintética"}
        ]
        rispy.dump(records, ris_source, encoding="utf8")
        ris_hash = hashlib.sha256(ris_source.read_bytes()).hexdigest()
        reread = rispy.load(ris_source, encoding="utf8")
        checks.require(reread == records, "RIS metadata changed during export and reread")
        reread[0]["title"] = "Segunda revisión de datos"
        rispy.dump(reread, ris_target, encoding="utf8")
        checks.require(rispy.load(ris_target, encoding="utf8") == reread,
                       "edited RIS metadata cannot be recovered")
        checks.require(reread[1] == records[1], "an unrelated RIS reference changed")
        checks.require({k: v for k, v in reread[0].items() if k != "title"} ==
                       {k: v for k, v in records[0].items() if k != "title"},
                       "RIS authors, DOI, pages or identifier changed")
        checks.require(hashlib.sha256(ris_source.read_bytes()).hexdigest() == ris_hash,
                       "the original RIS was overwritten")
        checks.require(reread[0]["doi"] == expected["garcia2024"]["fields"]["doi"]
                       and reread[0]["id"] == "garcia2024", "interchange identity changed")
        return checks.result(source, target, ris_source, ris_target)
