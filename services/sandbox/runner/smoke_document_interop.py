"""Small, offline interoperability proofs with synthetic documents.

Discovery imports only the standard library. These fixtures exercise native
edits and independent readers; they do not claim arbitrary-document fidelity.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import zipfile
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree


class _Checks:
    def __init__(self):
        self.count = 0

    def require(self, condition, message):
        if not condition:
            raise AssertionError(message)
        self.count += 1

    def rejects(self, exception, function, *args, **kwargs):
        try:
            function(*args, **kwargs)
        except exception:
            self.count += 1
            return
        raise AssertionError("invalid input was silently accepted")

    def result(self, *paths):
        return {"checks": self.count, "artifacts": [Path(path).name for path in paths]}


def _sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _json_file(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, allow_nan=False, indent=2), encoding="utf8")
    return json.loads(path.read_text(encoding="utf8"))


def _zip_parts(path):
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def register_proofs(proof):
    @proof("odf-native-edits")
    def odf_native_edits(root):
        from odfdo import Document, Paragraph, Span, Style, Table
        from odf.opendocument import load
        from odf import teletype
        from odf.text import P, Span as OdfSpan
        from odf.table import TableCell
        from python_calamine import CalamineWorkbook

        checks = _Checks()
        original_odt, edited_odt = root / "original.odt", root / "edited.odt"
        document = Document("text")
        paragraph_style = Style("paragraph", name="PreciseParagraph")
        paragraph_style.set_properties({"fo:margin-left": "12.375mm", "fo:margin-right": "8.5mm"}, area="paragraph")
        document.insert_style(paragraph_style)
        document.insert_style(Style("text", name="ProtectedBold", bold=True))
        document.body.append(Paragraph("Campo: pendiente", style="PreciseParagraph"))
        protected = Paragraph("Protegido: ", style="PreciseParagraph")
        protected.append(Span("niñez Ω", style="ProtectedBold"))
        document.body.append(protected)
        table = Table("Identificadores", width=2, height=2)
        table.set_row_values(0, ["ID", "Nota"])
        table.set_row_values(1, ["00017", "No modificar"])
        document.body.append(table)
        document.save(original_odt)
        digest_odt = _sha(original_odt)
        before = load(str(original_odt))
        before_paragraphs = [teletype.extractText(item) for item in before.getElementsByType(P)]
        reopened = Document(original_odt)
        targets = reopened.body.get_paragraphs(content=r"^Campo: pendiente$")
        checks.require(len(targets) == 1, "the edit must select exactly one target paragraph")
        targets[0].text = "Campo: revisión verificada"
        checks.require(reopened.body.get_paragraphs(content=r"^Campo inexistente$") == [],
                       "a missing target must not select another paragraph")
        reopened.save(edited_odt)
        after = load(str(edited_odt))
        after_paragraphs = [teletype.extractText(item) for item in after.getElementsByType(P)]
        expected_paragraphs = ["Campo: revisión verificada" if value == "Campo: pendiente" else value
                               for value in before_paragraphs]
        checks.require(after_paragraphs == expected_paragraphs, "ODT edit changed text outside its target")
        styled = [item for item in after.getElementsByType(P)
                  if teletype.extractText(item) == "Campo: revisión verificada"]
        checks.require(len(styled) == 1 and styled[0].getAttribute("stylename") == "PreciseParagraph",
                       "ODT edit changed the paragraph style reference")
        spans = after.getElementsByType(OdfSpan)
        checks.require(any(teletype.extractText(item) == "niñez Ω" and
                           item.getAttribute("stylename") == "ProtectedBold" for item in spans),
                       "ODT edit changed the protected styled span")
        parts_before, parts_after = _zip_parts(original_odt), _zip_parts(edited_odt)
        checks.require(parts_before.keys() == parts_after.keys(), "ODT edit added or removed package parts")
        checks.require(parts_before["styles.xml"] == parts_after["styles.xml"], "ODT style definitions or margins changed")
        checks.require([teletype.extractText(item) for item in after.getElementsByType(TableCell)] ==
                       [teletype.extractText(item) for item in before.getElementsByType(TableCell)],
                       "ODT table identifiers changed")
        checks.require(_sha(original_odt) == digest_odt, "ODT source was overwritten")

        original_ods, edited_ods = root / "original.ods", root / "edited.ods"
        spreadsheet = Document("spreadsheet")
        cell_style = Style("table-cell", name="MeasurementCell")
        cell_style.set_properties({"fo:background-color": "#eef2ff", "fo:padding": "1.25mm"}, area="table-cell")
        spreadsheet.insert_style(cell_style)
        sheet = Table("Medidas", width=3, height=3)
        sheet.set_row_values(0, ["ID", "Medida", "Nota"])
        sheet.set_row_values(1, ["00017", Decimal("12.375"), "niñez Ω"])
        sheet.set_row_values(2, ["00018", Decimal("8.500"), "Sin cambiar"])
        sheet.set_value("B2", Decimal("12.375"), style="MeasurementCell")
        spreadsheet.body.append(sheet)
        spreadsheet.save(original_ods)
        digest_ods = _sha(original_ods)
        with CalamineWorkbook.from_path(str(original_ods)) as independent:
            old_values = independent.get_sheet_by_name("Medidas").to_python()
        revised = Document(original_ods)
        target_sheet = revised.body.get_table_by_name("Medidas")
        checks.require(target_sheet.get_cell("B2").value == Decimal("12.375"),
                       "ODS precondition did not match the original value")
        target_sheet.set_value("B2", Decimal("12.500"))
        revised.save(edited_ods)
        with CalamineWorkbook.from_path(str(edited_ods)) as independent:
            new_values = independent.get_sheet_by_name("Medidas").to_python()
        expected_values = copy.deepcopy(old_values)
        expected_values[1][1] = 12.5
        checks.require(new_values == expected_values, "ODS edit changed cells outside B2")
        checks.require(new_values[1][0] == "00017" and new_values[2][0] == "00018",
                       "ODS edit coerced zero-padded identifiers")
        checks.require(new_values[1][2] == "niñez Ω", "ODS edit lost Unicode text")
        native = Document(edited_ods).body.get_table_by_name("Medidas").get_cell("B2")
        checks.require(native.get_attribute("table:style-name") == "MeasurementCell", "ODS edit lost the target cell style")
        checks.require(Decimal(native.get_attribute("office:value")) == Decimal("12.500"), "ODS numeric XML value is wrong")
        checks.require(_zip_parts(original_ods)["styles.xml"] == _zip_parts(edited_ods)["styles.xml"],
                       "ODS edit changed style definitions")
        checks.require(_sha(original_ods) == digest_ods, "ODS source was overwritten")
        summary = root / "native-edits.json"
        report = _json_file(summary, {"odt_changed_paragraphs": 1, "ods_changed_cells": ["B2"],
                                     "originals_unchanged": True, "independent_readers": ["odfpy", "calamine"]})
        checks.require(report["ods_changed_cells"] == ["B2"], "edit evidence did not survive JSON reopening")
        return checks.result(original_odt, edited_odt, original_ods, edited_ods, summary)

    @proof("qr-roundtrip")
    def qr_roundtrip(root):
        import segno
        import zxingcpp
        import vl_convert
        import pypdfium2 as pdfium
        from PIL import Image
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import mm
        from pptx import Presentation
        from pptx.util import Mm

        checks = _Checks()
        payload = "SiraGPT|ID=00017|Nombre=Niñez Ω 😀|Medida=12.375 mm"
        other_payload = payload.replace("00017", "00018")
        png, svg, second_png = root / "qr.png", root / "qr.svg", root / "different-qr.png"
        code = segno.make_qr(payload, error="h", encoding="utf-8", eci=True)
        code.save(png, scale=5, border=4, dark="black", light="white")
        code.save(svg, scale=5, border=4, dark="black", light="white")
        segno.make_qr(other_payload, error="h", encoding="utf-8", eci=True).save(second_png, scale=5, border=4)

        def decode(image):
            return zxingcpp.read_barcodes(image.convert("RGB"), formats=zxingcpp.BarcodeFormat.QRCode)

        with Image.open(png) as image:
            found = decode(image)
            checks.require(len(found) == 1 and found[0].text == payload, "PNG QR did not preserve its Unicode payload")
            checks.require(image.width == image.height and image.width < 1000, "QR fixture exceeds its bounded canvas")
        tree = ElementTree.parse(svg)
        checks.require(tree.getroot().tag == "{http://www.w3.org/2000/svg}svg", "QR SVG is not valid vector content")
        svg_png = root / "qr-from-svg.png"
        svg_png.write_bytes(vl_convert.svg_to_png(svg.read_text(encoding="utf8")))
        with Image.open(svg_png) as image:
            found = decode(image)
            checks.require(len(found) == 1 and found[0].text == payload, "rendered SVG QR is not decodable")
        with Image.open(second_png) as image:
            found = decode(image)
            checks.require(len(found) == 1 and found[0].text == other_payload and found[0].text != payload,
                           "decoder returned the first payload for a different QR")
        blank = root / "blank.png"
        Image.new("RGB", (300, 300), "white").save(blank)
        with Image.open(blank) as image:
            checks.require(decode(image) == [], "blank image produced a false QR detection")
        checks.rejects(segno.DataOverflowError, segno.make_qr, "X" * 100, version=1, error="h")

        pdf = root / "qr-document.pdf"
        page = canvas.Canvas(str(pdf), pagesize=(120 * mm, 120 * mm))
        page.drawString(15 * mm, 105 * mm, "QR verification")
        page.drawImage(str(png), 25 * mm, 25 * mm, width=70 * mm, height=70 * mm)
        page.save()
        pdf_png = root / "qr-from-pdf.png"
        with pdfium.PdfDocument(pdf) as document:
            checks.require(len(document) == 1, "QR PDF has an unexpected page count")
            pdf_page = document[0]
            bitmap = pdf_page.render(scale=2)
            bitmap.to_pil().save(pdf_png)
            bitmap.close()
            pdf_page.close()
        with Image.open(pdf_png) as image:
            found = decode(image)
            checks.require(len(found) == 1 and found[0].text == payload, "QR did not survive PDF embedding and rendering")
        pptx = root / "qr-slide.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(png), Mm(10), Mm(15), width=Mm(50), height=Mm(50))
        deck.save(pptx)
        restored = Presentation(pptx).slides[0].shapes[0]
        checks.require((restored.left, restored.top, restored.width, restored.height) ==
                       (Mm(10), Mm(15), Mm(50), Mm(50)), "QR position or size changed in PPTX")
        checks.require(restored.image.blob == png.read_bytes(), "PPTX altered the embedded QR bytes")
        with Image.open(io.BytesIO(restored.image.blob)) as image:
            found = decode(image)
            checks.require(len(found) == 1 and found[0].text == payload, "PPTX QR image payload changed")
        report_path = root / "qr-evidence.json"
        report = _json_file(report_path, {"payload": payload, "decoded_from": ["png", "svg", "pdf", "pptx-image"],
                                         "blank_detections": 0, "pptx_rasterized": False})
        checks.require(report["payload"] == payload, "QR evidence lost Unicode text")
        return checks.result(png, svg, second_png, svg_png, blank, pdf, pdf_png, pptx, report_path)

    @proof("structured-diff")
    def structured_diff(root):
        from deepdiff import DeepDiff

        checks = _Checks()
        first_amount = Decimal("12.3750000000000000000000001")
        next_amount = Decimal("12.3750000000000000000000002")
        original = {"records": [{"id": "00017", "amount": first_amount, "label": "niñez Ω"},
                                {"id": "00018", "amount": Decimal("8.5"), "label": "Sin cambiar"}],
                    "order": ["first", "second"], "count": 20, "enabled": True}
        edited = copy.deepcopy(original)
        edited["records"][0]["amount"] = next_amount

        def compare(left, right, **options):
            return DeepDiff(left, right, ignore_order=False, ignore_numeric_type_changes=False,
                            significant_digits=None, verbose_level=2, max_diffs=options.get("max_diffs", 1000))

        difference = compare(original, edited)
        allowed_path = "root['records'][0]['amount']"
        checks.require(set(difference) == {"values_changed"} and set(difference["values_changed"]) == {allowed_path},
                       "the decimal edit must produce exactly one changed path")
        checks.require(difference["values_changed"][allowed_path]["old_value"] == first_amount and
                       difference["values_changed"][allowed_path]["new_value"] == next_amount,
                       "fine decimal differences were rounded away")
        checks.require(not difference.get_stats()["MAX DIFF LIMIT REACHED"], "a partial diff was treated as complete")
        checks.require(not compare(original, copy.deepcopy(original)), "identical structures produced a change")
        collateral = copy.deepcopy(edited)
        collateral["records"][1]["label"] = "Cambio no pedido"
        collateral_difference = compare(original, collateral)
        checks.require(set(collateral_difference["values_changed"]) - {allowed_path} ==
                       {"root['records'][1]['label']"}, "an unrequested change escaped detection")
        reordered = copy.deepcopy(original)
        reordered["order"].reverse()
        order_difference = compare(original, reordered)
        checks.require(set(order_difference["values_changed"]) == {"root['order'][0]", "root['order'][1]"},
                       "list ordering was ignored")
        type_changed = copy.deepcopy(original)
        type_changed["count"] = "20"
        type_changed["enabled"] = 1
        type_difference = compare(original, type_changed)
        checks.require(set(type_difference["type_changes"]) == {"root['count']", "root['enabled']"},
                       "string/integer or boolean/integer changes were coerced away")
        source_path, edited_path = root / "source.json", root / "edited.json"

        def persist_data(path, value):
            serialized = copy.deepcopy(value)
            for record in serialized["records"]:
                record["amount"] = {"decimal": str(record["amount"])}
            loaded = _json_file(path, serialized)
            for record in loaded["records"]:
                record["amount"] = Decimal(record["amount"]["decimal"])
            return loaded

        loaded_source = persist_data(source_path, original)
        digest = _sha(source_path)
        loaded_edit = persist_data(edited_path, edited)
        checks.require(not compare(loaded_source, original), "source serialization lost types or decimal precision")
        checks.require(compare(loaded_source, loaded_edit) == difference, "reopened files produced a different diff")
        evidence = root / "exact-diff.json"
        evidence.write_text(difference.to_json(default_mapping={Decimal: lambda value: {"decimal": str(value)}}),
                            encoding="utf8")
        restored = json.loads(evidence.read_text(encoding="utf8"))
        values = restored["values_changed"][allowed_path]
        checks.require(Decimal(values["old_value"]["decimal"]) == first_amount and
                       Decimal(values["new_value"]["decimal"]) == next_amount,
                       "diff JSON lost exact decimal values")
        # Demonstrate the explicit limit signal; never present a capped diff
        # as exhaustive. No Delta objects, pickle or executable deserialization.
        capped = compare(list(range(30)), list(range(100, 130)), max_diffs=1)
        checks.require(capped.get_stats()["MAX DIFF LIMIT REACHED"], "comparison budget exhaustion was not reported")
        limit_path = root / "limited-diff.json"
        limit_report = _json_file(limit_path, {"complete": False, "reason": "max_diffs",
                                               "limit_reached": capped.get_stats()["MAX DIFF LIMIT REACHED"]})
        checks.require(limit_report["complete"] is False and limit_report["limit_reached"] is True,
                       "limited evidence was misrepresented as complete")
        checks.require(_sha(source_path) == digest and original["records"][0]["amount"] == first_amount,
                       "read-only comparison changed source data")
        return checks.result(source_path, edited_path, evidence, limit_path)
