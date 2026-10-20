"""Offline, file-backed proofs for the document toolkit.

Registration needs only the Python standard library. Optional libraries are
imported inside each proof, so discovery works before installing dependencies.
All inputs are synthetic; all outputs stay inside the supplied proof directory.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree


class _Checks:
    """Count checks actually executed, including expected rejection cases."""

    def __init__(self):
        self.count = 0

    def require(self, condition, message):
        if not condition:
            raise AssertionError(message)
        self.count += 1

    def rejects(self, error_type, function, *args, **kwargs):
        try:
            function(*args, **kwargs)
        except error_type:
            self.count += 1
            return
        raise AssertionError("invalid input was silently accepted")

    def result(self, *artifacts):
        return {"checks": self.count, "artifacts": [Path(p).name for p in artifacts]}


def _json_file(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, allow_nan=False, indent=2), encoding="utf8")
    return json.loads(path.read_text(encoding="utf8"))


def _sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def register_proofs(proof):
    @proof("document-composition")
    def document_composition(root):
        from docx import Document
        from docx.shared import Mm
        from docxcompose.composer import Composer
        from docx2python import docx2python
        from striprtf.striprtf import rtf_to_text
        from PIL import Image, ImageDraw

        checks = _Checks()
        picture = root / "measured-image.png"
        image = Image.new("RGB", (120, 60), "white")
        ImageDraw.Draw(image).rectangle((10, 10, 100, 40), fill="#225588")
        image.save(picture)
        first, second, combined = (root / name for name in ("first.docx", "second.docx", "combined.docx"))
        master = Document()
        master.sections[0].header.paragraphs[0].text = "Cabecera principal"
        master.add_paragraph("Primer documento — precisión").runs[0].bold = True
        master.add_table(rows=1, cols=2).rows[0].cells[0].text = "00017"
        master.tables[0].cell(0, 1).text = "12.375"
        master.save(first)
        appendix = Document()
        appendix.sections[0].header.paragraphs[0].text = "Cabecera secundaria"
        appendix.add_paragraph("Segundo documento — niñez Ω").runs[0].italic = True
        appendix.add_table(rows=1, cols=2).rows[0].cells[0].text = "00018"
        appendix.tables[0].cell(0, 1).text = "Información protegida"
        appendix.add_picture(str(picture), width=Mm(37.5), height=Mm(18.75))
        appendix.save(second)
        originals = (_sha(first), _sha(second))
        composer = Composer(Document(first))
        composer.append(Document(second))
        composer.save(combined)
        reopened = Document(combined)
        texts = [paragraph.text for paragraph in reopened.paragraphs]
        checks.require("Primer documento — precisión" in texts, "composition lost the first document")
        checks.require("Segundo documento — niñez Ω" in texts, "composition lost Unicode text")
        checks.require(len(reopened.tables) == 2, "composition lost a table")
        checks.require([[cell.text for cell in table.rows[0].cells] for table in reopened.tables] ==
                       [["00017", "12.375"], ["00018", "Información protegida"]], "table values changed")
        checks.require(len(reopened.inline_shapes) == 1, "the appended image was lost")
        shape = reopened.inline_shapes[0]
        checks.require((shape.width, shape.height) == (Mm(37.5), Mm(18.75)), "image EMU dimensions changed")
        checks.require(reopened.sections[0].header.paragraphs[0].text == "Cabecera principal",
                       "the documented master-header behavior changed")
        checks.require(next(p for p in reopened.paragraphs if p.text == texts[0]).runs[0].bold,
                       "direct bold formatting changed")
        checks.require(next(p for p in reopened.paragraphs if p.text == "Segundo documento — niñez Ω").runs[0].italic,
                       "direct italic formatting changed")
        with docx2python(combined) as extracted:
            checks.require(all(value in extracted.text for value in
                               ("00017", "00018", "Información protegida", "niñez Ω")), "independent text extraction lost data")
            checks.require(any(hashlib.sha256(data).hexdigest() == _sha(picture)
                               for data in extracted.images.values()), "image bytes changed during composition")
            with Image.open(io.BytesIO(next(iter(extracted.images.values())))) as restored:
                checks.require(restored.size == (120, 60), "extracted image dimensions changed")
        checks.require((_sha(first), _sha(second)) == originals, "composition modified source files")
        rich_text = root / "unicode.rtf"
        rich_text.write_text(r"{\rtf1\ansi\ansicpg1252\uc1 Espa\u241?a \u945? caf\'e9 \{ok\}\par Fin}", encoding="ascii")
        plain = root / "unicode.txt"
        plain.write_text(rtf_to_text(rich_text.read_text(encoding="ascii")), encoding="utf8")
        recovered = plain.read_text(encoding="utf8")
        checks.require(all(value in recovered for value in ("España", "α", "café", "{ok}", "\nFin")),
                       "RTF Unicode, codepage, escaping or paragraph extraction failed")
        return checks.result(first, second, combined, picture, rich_text, plain)

    @proof("structured-json")
    def structured_json(root):
        import ijson
        import jmespath
        import jsonpatch
        from jsonschema import Draft202012Validator, ValidationError
        from packaging.requirements import Requirement, InvalidRequirement
        from packaging.version import Version

        checks = _Checks()
        schema = {
            "$schema": "https://json-schema.org/draft/2020-12/schema",
            "type": "object", "additionalProperties": False,
            "required": ["records", "requirements"],
            "properties": {
                "records": {"type": "array", "items": {
                    "type": "object", "additionalProperties": False,
                    "required": ["id", "amount", "sequence"],
                    "properties": {"id": {"type": "string", "pattern": "^[0-9]{5}$"},
                                   "amount": {"type": ["number", "null"]},
                                   "sequence": {"type": "integer"}}}},
                "requirements": {"type": "array", "items": {"type": "string"}},
            },
        }
        source = root / "source.json"
        original = _json_file(source, {"records": [{"id": "00017", "amount": 12.375, "sequence": 9007199254740993},
                                                    {"id": "00018", "amount": None, "sequence": 2}],
                                      "requirements": ["packaging==26.3", "jsonschema>=4,<5"]})
        digest = _sha(source)
        schema_path = root / "schema.json"
        validator = Draft202012Validator(_json_file(schema_path, schema))
        validator.validate(original)
        checks.require(validator.is_valid(original), "the processing manifest did not validate")
        invalid = json.loads(json.dumps(original))
        invalid["records"][0]["id"] = 17
        checks.rejects(ValidationError, validator.validate, invalid)
        patch_path = root / "edit.jsonpatch"
        operations = _json_file(patch_path, [{"op": "test", "path": "/records/0/id", "value": "00017"},
                                            {"op": "test", "path": "/records/0/amount", "value": 12.375},
                                            {"op": "replace", "path": "/records/0/amount", "value": 12.5}])
        patch = jsonpatch.JsonPatch(operations)
        changed = patch.apply(original, in_place=False)
        checks.require(original["records"][0]["amount"] == 12.375, "patch modified its input")
        checks.require(changed["records"][0]["amount"] == 12.5, "the requested field was not changed")
        checks.require(changed["records"][1] == original["records"][1], "an unrelated record changed")
        stale = json.loads(json.dumps(original))
        stale["records"][0]["amount"] = 99
        checks.rejects(jsonpatch.JsonPatchTestFailed, patch.apply, stale, in_place=False)
        checks.require(stale["records"][0]["amount"] == 99, "failed precondition changed input data")
        checks.require(jmespath.search("records[?id == '00017'].id", changed) == ["00017"],
                       "query coerced the identifier or selected the wrong record")
        edited = root / "edited.json"
        restored = _json_file(edited, changed)
        checks.require(validator.is_valid(restored), "the edited manifest violates its schema")
        with source.open("rb") as stream:
            records = list(ijson.items(stream, "records.item", use_float=False))
        checks.require(records[0]["id"] == "00017" and isinstance(records[0]["id"], str),
                       "streaming extraction coerced a zero-padded identifier")
        checks.require(records[0]["amount"] == Decimal("12.375") and isinstance(records[0]["amount"], Decimal),
                       "streaming extraction lost decimal precision")
        checks.require(records[0]["sequence"] == 9007199254740993 and records[1]["amount"] is None,
                       "streaming extraction lost a large integer or null")
        checks.rejects(ijson.JSONError, lambda: list(ijson.items(io.BytesIO(b'{"records":['), "records.item")))
        requirement = Requirement(restored["requirements"][0])
        checks.require(Version("26.3") in requirement.specifier and Version("26.4") not in requirement.specifier,
                       "the package pin does not enforce its version")
        ranged = Requirement(restored["requirements"][1])
        checks.require(Version("4.25.1") in ranged.specifier and Version("5.0") not in ranged.specifier,
                       "PEP 440 range validation changed")
        checks.rejects(InvalidRequirement, Requirement, "not a valid requirement !!!")
        checks.require(_sha(source) == digest, "source JSON was modified")
        return checks.result(source, schema_path, patch_path, edited)

    @proof("fuzzy-candidates")
    def fuzzy_candidates(root):
        from rapidfuzz import fuzz, process

        checks = _Checks()
        source = root / "names.json"
        names = _json_file(source, {"record-a": "Jorge Carrera", "record-b": "Jorge Carera",
                                   "record-c": "María Fernández"})
        digest = _sha(source)
        candidates = process.extract("Jorge Carrera", names, scorer=fuzz.WRatio, score_cutoff=80, limit=2)
        checks.require(len(candidates) == 2, "near duplicates should remain distinct candidates")
        checks.require(candidates[0][2] == "record-a" and candidates[0][1] == 100, "exact candidate was not ranked first")
        checks.require(candidates[1][2] == "record-b" and 80 <= candidates[1][1] < 100, "near-match score is not distinguishable")
        checks.require(process.extract("zzzzzz", names, scorer=fuzz.WRatio, score_cutoff=95) == [],
                       "unrelated text was suggested above the cutoff")
        output = root / "candidates.json"
        report = _json_file(output, {"usage": "candidate_suggestions_only", "requires_review": True,
                                    "selected_id": None, "candidates": [
                                        {"id": key, "name": name, "score": score} for name, score, key in candidates]})
        checks.require(report["selected_id"] is None and len(report["candidates"]) == 2,
                       "the candidate workflow must not make an automatic correction")
        checks.require(_sha(source) == digest and json.loads(source.read_text()) == names, "matching changed source records")
        return checks.result(source, output)

    @proof("physical-units")
    def physical_units(root):
        import pint
        from uncertainties import ufloat

        checks = _Checks()
        units = pint.UnitRegistry()
        length = ufloat(120.0, 0.5) * units.mm
        width = ufloat(80.0, 0.25) * units.mm
        area = (length * width).to("cm**2")
        checks.require(math.isclose(area.magnitude.nominal_value, 96), "area units were not converted")
        checks.require(math.isclose(area.magnitude.std_dev, 0.5), "independent measurement uncertainty was not propagated")
        back = length.to("inch").to("mm").magnitude
        checks.require(math.isclose(back.nominal_value, 120) and math.isclose(back.std_dev, 0.5),
                       "unit roundtrip lost value or uncertainty")
        checks.require((length - length).magnitude.std_dev == 0, "correlated measurements lost their shared uncertainty")
        checks.rejects(pint.DimensionalityError, length.to, "second")
        checks.require(math.isclose(units.Quantity(20, units.degC).to("kelvin").magnitude, 293.15),
                       "offset temperature conversion failed")
        output = root / "measurement.json"
        saved = _json_file(output, {"nominal": area.magnitude.nominal_value, "std_dev": area.magnitude.std_dev,
                                   "unit": str(area.units), "assumption": "independent length and width"})
        restored = units.Quantity(ufloat(saved["nominal"], saved["std_dev"]), saved["unit"]).to("mm**2")
        checks.require(math.isclose(restored.magnitude.nominal_value, 9600) and
                       math.isclose(restored.magnitude.std_dev, 50), "persisted uncertainty cannot be reconstructed")
        return checks.result(output)

    @proof("svg-geometry")
    def svg_geometry(root):
        from svgpathtools import parse_path, svg2paths2
        import vl_convert
        from PIL import Image

        checks = _Checks()
        curve = parse_path("M 10 50 C 10 10 90 10 90 50 L 90 70 L 10 70 Z")
        checks.require(all(math.isclose(a, b) for a, b in zip(curve.bbox(), (10, 90, 20, 70))),
                       "curve bounds must include extrema, not just control points")
        checks.require(curve.isclosed() and curve.length() > 200, "the closed SVG path was not parsed")
        document = ElementTree.Element("svg", xmlns="http://www.w3.org/2000/svg",
                                       width="120", height="100", viewBox="0 0 120 100")
        ElementTree.SubElement(document, "path", d=curve.d(), fill="#225588")
        vector = root / "geometry.svg"
        ElementTree.ElementTree(document).write(vector, encoding="utf8", xml_declaration=True)
        paths, attributes, svg_attributes = svg2paths2(str(vector))
        checks.require(len(paths) == 1 and paths[0].bbox() == curve.bbox(), "written SVG changed geometric bounds")
        checks.require(attributes[0]["fill"] == "#225588" and svg_attributes["viewBox"] == "0 0 120 100",
                       "written SVG lost its fill or coordinate system")
        raster = root / "geometry.png"
        raster.write_bytes(vl_convert.svg_to_png(vector.read_text(encoding="utf8")))
        with Image.open(raster) as image:
            rgba = image.convert("RGBA")
            checks.require(rgba.size == (120, 100), "rendered SVG has the wrong canvas")
            checks.require(rgba.getpixel((50, 50)) == (34, 85, 136, 255), "interior shape was not rendered")
            checks.require(rgba.getpixel((5, 5))[3] == 0, "renderer painted outside the expected shape")
        return checks.result(vector, raster)

    @proof("declarative-charts")
    def declarative_charts(root):
        import altair as alt
        import vl_convert
        from PIL import Image
        from pypdf import PdfReader
        from pptx import Presentation
        from pptx.util import Mm
        import matplotlib

        checks = _Checks()
        font_dir = Path("/usr/share/fonts/truetype/dejavu")
        if not font_dir.is_dir():
            font_dir = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        vl_convert.register_font_directory(str(font_dir))
        records = [{"Grupo": "Niñez", "Medida": 12.375}, {"Grupo": "Adultos", "Medida": 18.5}]
        chart = (alt.Chart(alt.InlineData(values=records)).mark_bar(color="#225588")
                 .encode(x=alt.X("Grupo:N", sort=None), y=alt.Y("Medida:Q"))
                 .properties(width=320, height=180, title="Medición verificada")
                 .configure(font="DejaVu Sans"))
        specification = root / "chart.vl.json"
        spec = _json_file(specification, chart.to_dict(validate=True))

        def no_remote_data(value):
            if isinstance(value, dict):
                return "url" not in value and all(no_remote_data(item) for item in value.values())
            if isinstance(value, list):
                return all(no_remote_data(item) for item in value)
            return True

        checks.require(no_remote_data(spec), "the chart must embed all data and assets")
        embedded = spec["data"].get("values", spec.get("datasets", {}).get(spec["data"].get("name")))
        checks.require(embedded == records, "the persisted chart lost its source values")
        vector, raster, pdf = (root / name for name in ("chart.svg", "chart.png", "chart.pdf"))
        vector.write_text(vl_convert.vegalite_to_svg(spec, allowed_base_urls=[]), encoding="utf8")
        raster.write_bytes(vl_convert.vegalite_to_png(spec, allowed_base_urls=[]))
        pdf.write_bytes(vl_convert.vegalite_to_pdf(spec, allowed_base_urls=[]))
        tree = ElementTree.parse(vector)
        labels = "".join(tree.getroot().itertext())
        checks.require("Medición verificada" in labels and "Niñez" in labels, "SVG text labels were lost")
        with Image.open(raster) as image:
            checks.require(image.width >= 320 and image.height >= 180, "PNG chart is too small")
            aspect_ratio = image.height / image.width
            extrema = image.convert("RGB").getextrema()
            checks.require(any(low < high for low, high in extrema), "PNG chart is blank")
        rendered_pdf = PdfReader(pdf)
        checks.require(len(rendered_pdf.pages) == 1, "PDF export has the wrong page count")
        checks.require("Medición verificada" in rendered_pdf.pages[0].extract_text(), "PDF text is not readable")
        checks.require(float(rendered_pdf.pages[0].mediabox.width) > 200, "PDF page geometry is invalid")
        # The plotted graphic is an image; the separate Vega-Lite spec stays
        # editable. The slide title is a native, independently editable textbox.
        deck = Presentation()
        deck.slide_width, deck.slide_height = Mm(254), Mm(190.5)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Mm(10), Mm(10), Mm(230), Mm(15))
        title.name, title.text = "editable-title", "Medición verificada"
        image_height = round(Mm(160) * aspect_ratio)
        graphic = slide.shapes.add_picture(str(raster), Mm(10), Mm(35), width=Mm(160), height=image_height)
        graphic.name = "chart-image"
        presentation = root / "chart.pptx"
        deck.save(presentation)
        reopened = Presentation(presentation)
        checks.require(len(reopened.slides) == 1, "the chart slide was lost")
        shapes = {shape.name: shape for shape in reopened.slides[0].shapes}
        checks.require((shapes["chart-image"].left, shapes["chart-image"].top,
                        shapes["chart-image"].width, shapes["chart-image"].height) ==
                       (Mm(10), Mm(35), Mm(160), image_height), "PPTX chart geometry changed")
        checks.require(shapes["chart-image"].image.blob == raster.read_bytes(), "PPTX embedding changed the chart image")
        checks.require(shapes["editable-title"].has_text_frame and
                       shapes["editable-title"].text == "Medición verificada", "the slide title is not native text")
        shapes["editable-title"].text = "Título editado"
        revised_presentation = root / "chart-title-edited.pptx"
        reopened.save(revised_presentation)
        revised_title = next(shape for shape in Presentation(revised_presentation).slides[0].shapes
                             if shape.name == "editable-title")
        checks.require(revised_title.text == "Título editado", "the slide title edit did not survive reopening")
        return checks.result(specification, vector, raster, pdf, presentation, revised_presentation)

    @proof("columnar-data")
    def columnar_data(root):
        import duckdb
        import pyarrow as pa
        import pyarrow.parquet as pq

        checks = _Checks()
        root = root.resolve()
        schema = pa.schema([("id", pa.string()), ("amount", pa.decimal128(16, 3)), ("sequence", pa.int64()),
                            ("label", pa.string()), ("enabled", pa.bool_())])
        rows = [{"id": "00017", "amount": Decimal("12.375"), "sequence": 9007199254740993, "label": "Niñez — Ω", "enabled": True},
                {"id": "00018", "amount": None, "sequence": -1, "label": None, "enabled": False},
                {"id": "00019", "amount": Decimal("-0.125"), "sequence": 0, "label": "", "enabled": None}]
        table = pa.Table.from_pylist(rows, schema=schema)
        source, edited = root / "source.parquet", root / "edited.parquet"
        pq.write_table(table, source, compression="zstd", row_group_size=2)
        digest = _sha(source)
        checks.require(pq.read_table(source, use_threads=False).equals(table), "Arrow Parquet roundtrip changed types or values")
        configuration = {"threads": 1, "memory_limit": "128MB", "autoinstall_known_extensions": False,
                         "autoload_known_extensions": False}
        with duckdb.connect(":memory:", config=configuration) as database:
            # DuckDB requires defining the local allowlist before disabling
            # external access. No file query runs until both gates are set.
            database.execute("SET allowed_directories = ?", [[str(root)]])
            database.execute("SET enable_external_access = false")
            database.execute("SET lock_configuration = true")
            checks.require(database.execute("SELECT current_setting('threads')").fetchone()[0] == 1, "DuckDB is not thread bounded")
            checks.require(database.execute("SELECT current_setting('enable_external_access')").fetchone()[0] is False,
                           "DuckDB external access is enabled")
            database.execute("CREATE TABLE measurements AS SELECT * FROM read_parquet(?)", [str(source)])
            decoded = database.execute("SELECT id, amount, sequence, label, enabled FROM measurements ORDER BY id").fetchall()
            checks.require(decoded == [tuple(row[field] for field in schema.names) for row in rows],
                           "DuckDB coerced a decimal, null, identifier or Unicode value")
            database.execute("UPDATE measurements SET amount = ? WHERE id = ?", [Decimal("12.500"), "00017"])
            database.table("measurements").order("id").write_parquet(str(edited), compression="zstd")
        expected = [dict(row) for row in rows]
        expected[0]["amount"] = Decimal("12.500")
        restored = pq.read_table(edited, use_threads=False)
        checks.require(restored.schema.equals(schema, check_metadata=False), "DuckDB changed the Arrow field types")
        checks.require(restored.to_pylist() == expected, "Parquet edit changed fields beyond the requested value")
        checks.require(_sha(source) == digest, "columnar editing overwrote its source")
        checks.rejects(pa.ArrowInvalid, pa.array, [Decimal("100000000000000")], type=pa.decimal128(16, 3))
        return checks.result(source, edited)

    @proof("visual-fingerprints")
    def visual_fingerprints(root):
        import imagehash
        from PIL import Image, ImageDraw, ImageOps

        checks = _Checks()
        source, shifted = root / "original.png", root / "mirrored.png"
        image = Image.new("RGB", (128, 128), "white")
        ImageDraw.Draw(image).rectangle((8, 12, 52, 90), fill="black")
        image.save(source)
        ImageOps.mirror(image).save(shifted)
        with Image.open(source) as restored, Image.open(shifted) as different:
            original_hash = imagehash.phash(image)
            restored_hash = imagehash.phash(restored)
            changed_hash = imagehash.phash(different)
            checks.require(original_hash - restored_hash == 0, "lossless reopening changed the perceptual fingerprint")
            checks.require(original_hash - changed_hash > 0, "the mirrored asymmetric shape was not distinguished")
            checks.require(restored.convert("RGB").tobytes() == image.tobytes(), "lossless reopening changed pixels")
        red, blue = root / "red.png", root / "blue.png"
        Image.new("RGB", (64, 64), "red").save(red)
        Image.new("RGB", (64, 64), "blue").save(blue)
        with Image.open(red) as first, Image.open(blue) as second:
            collision = bool(imagehash.average_hash(first) == imagehash.average_hash(second))
            checks.require(collision, "uniform colours should demonstrate a perceptual-hash collision")
            checks.require(first.tobytes() != second.tobytes(), "the collision fixture must contain different pixels")
        report = root / "fingerprints.json"
        saved = _json_file(report, {"original_phash": str(original_hash), "mirrored_phash": str(changed_hash),
                                   "distance": int(original_hash - changed_hash), "usage": "visual_similarity_signal_only",
                                   "collision_proven": collision, "pixel_equality_verified_separately": True})
        checks.require(imagehash.hex_to_hash(saved["original_phash"]) == original_hash, "persisted fingerprint changed")
        checks.require(saved["distance"] > 0 and _sha(red) != _sha(blue), "a similarity signal was treated as file equality")
        return checks.result(source, shifted, red, blue, report)
