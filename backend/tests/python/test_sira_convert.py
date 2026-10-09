"""Real saved-file acceptance: no provider mocks, renamed files, or skipped tools."""
import hashlib
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.shared import Inches
from PIL import Image
from pypdf import PdfReader, PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src/services/agent-runner"))
import sira_convert as converter


class RealConversions(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # These are mandatory in this acceptance suite and installed on CI's
        # validator shard. Absence must fail visibly, not become a green skip.
        for executable in ("soffice", "ffmpeg"):
            if not shutil.which(executable):
                raise AssertionError(f"Required conversion runtime missing: {executable}")

    def setUp(self):
        self.previous = Path.cwd()
        self.temp = tempfile.TemporaryDirectory(prefix="sira-conversion-test-")
        self.root = Path(self.temp.name)
        os.chdir(self.root)
        Path("uploads").mkdir()
        Path("outputs").mkdir()

    def tearDown(self):
        os.chdir(self.previous)
        self.temp.cleanup()

    def make_docx(self):
        path = Path("uploads/Informe con espacios.docx")
        doc = Document()
        doc.add_heading("Informe de prueba", 0)
        doc.add_paragraph("Primer texto: información, acción y revisión profesional.")
        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        for cell, value in zip((cell for row in table.rows for cell in row.cells), ("Producto", "Cantidad", "Bicicletas", "27")):
            cell.text = value
        Image.new("RGB", (160, 80), (22, 100, 120)).save("uploads/reference.png")
        doc.add_picture("uploads/reference.png", width=Inches(1.6))
        doc.add_page_break()
        doc.add_heading("Segunda página", 1)
        doc.add_paragraph("Texto final que debe sobrevivir a la conversión.")
        doc.save(path)
        return path

    def make_mp3(self):
        path = Path("uploads/audio con espacios.mp3")
        subprocess.run([shutil.which("ffmpeg"), "-hide_banner", "-loglevel", "error", "-nostdin",
                        "-f", "lavfi", "-i", "sine=frequency=440:duration=1.2", "-c:a", "libmp3lame", str(path)], check=True)
        return path

    def assert_report(self, result, expected):
        self.assertTrue(result["ok"], result)
        path = Path(result["output"])
        self.assertEqual(path, Path(expected))
        self.assertGreater(result["bytes"], 0)
        self.assertEqual(result["bytes"], path.stat().st_size)
        self.assertEqual(result["sha256"], hashlib.sha256(path.read_bytes()).hexdigest())
        self.assertTrue(result["validated"])

    def test_word_pdf_word_preserves_real_text_tables_and_sources(self):
        source = self.make_docx()
        original = source.read_bytes()
        pdf = converter.convert(str(source), "outputs/informe.pdf")
        self.assert_report(pdf, "outputs/informe.pdf")
        reader = PdfReader("outputs/informe.pdf")
        self.assertEqual(len(reader.pages), 2)
        self.assertTrue(Path("outputs/informe.pdf").read_bytes().startswith(b"%PDF-"))
        self.assertIn("Texto final", reader.pages[1].extract_text())
        word = converter.convert("outputs/informe.pdf", "outputs/reconstruido.docx")
        self.assert_report(word, "outputs/reconstruido.docx")
        self.assertEqual(word["fidelity"], "editable_reconstruction")
        self.assertTrue(word["warnings"])
        self.assertTrue(word["text_preserved"])
        reconstructed = Document("outputs/reconstruido.docx")
        self.assertIn("Texto final", "\n".join(p.text for p in reconstructed.paragraphs))
        self.assertEqual(len(reconstructed.tables), 1)
        self.assertEqual(reconstructed.tables[0].cell(1, 1).text, "27")
        self.assertEqual(len(reconstructed.inline_shapes), 1)
        self.assertEqual(word["images"], 1)
        self.assertEqual(source.read_bytes(), original)

    def test_mp3_mp4_mp3_contains_actual_video_and_audio_and_decodes(self):
        source = self.make_mp3()
        original = source.read_bytes()
        video = converter.convert(str(source), "outputs/video.mp4")
        self.assert_report(video, "outputs/video.mp4")
        self.assertTrue(video["decoded"])
        self.assertEqual({(s["type"], s["codec"]) for s in video["streams"]}, {("video", "h264"), ("audio", "aac")})
        self.assertIn(b"ftyp", Path("outputs/video.mp4").read_bytes()[:40])
        audio = converter.convert("outputs/video.mp4", "outputs/audio.mp3")
        self.assert_report(audio, "outputs/audio.mp3")
        self.assertEqual(audio["streams"], [{"type": "audio", "codec": "mp3"}])
        self.assertAlmostEqual(video["duration_seconds"], audio["duration_seconds"], delta=0.2)
        self.assertTrue(audio["decoded"])
        self.assertEqual(source.read_bytes(), original)
        # Independent decode proves a playable signal is present, not just a
        # matching extension, container header or report.
        pcm = subprocess.run([shutil.which("ffmpeg"), "-v", "error", "-i", "outputs/audio.mp3",
                              "-f", "s16le", "-acodec", "pcm_s16le", "-"], check=True, capture_output=True).stdout
        self.assertGreater(len(pcm), 50_000)
        self.assertGreater(len(set(pcm)), 100)

    def test_silent_video_does_not_fabricate_audio(self):
        subprocess.run([shutil.which("ffmpeg"), "-hide_banner", "-loglevel", "error", "-f", "lavfi",
                        "-i", "color=c=black:s=64x64:r=25", "-t", "0.5", "-c:v", "libx264", "uploads/silent.mp4"], check=True)
        result = converter.convert("uploads/silent.mp4", "outputs/audio.mp3")
        self.assertFalse(result["ok"])
        self.assertEqual(result["code"], "E_CONTENT")
        self.assertIn("no contiene una pista de audio", result["error"])
        self.assertFalse(Path("outputs/audio.mp3").exists())

    def test_scan_only_pdf_requests_ocr_without_publishing_empty_word(self):
        writer = PdfWriter()
        writer.add_blank_page(width=600, height=800)
        writer.write("uploads/scan.pdf")
        result = converter.convert("uploads/scan.pdf", "outputs/scan.docx")
        self.assertFalse(result["ok"])
        self.assertIn("OCR", result["error"])
        self.assertFalse(Path("outputs/scan.docx").exists())

    def test_encrypted_pdf_is_actionable(self):
        writer = PdfWriter()
        writer.add_blank_page(width=600, height=800)
        writer.encrypt("test-only-password")
        writer.write("uploads/private.pdf")
        result = converter.convert("uploads/private.pdf", "outputs/result.docx")
        self.assertFalse(result["ok"])
        self.assertIn("protegido", result["error"])

    def test_pdf_with_one_scan_page_reports_ocr_gap(self):
        source = self.make_docx()
        self.assertTrue(converter.convert(str(source), "outputs/base.pdf")["ok"])
        writer = PdfWriter()
        writer.append("outputs/base.pdf")
        writer.add_blank_page(width=600, height=800)
        writer.write("uploads/partial.pdf")
        result = converter.convert("uploads/partial.pdf", "outputs/partial.docx")
        self.assert_report(result, "outputs/partial.docx")
        self.assertTrue(any("requieren OCR" in warning and "3" in warning for warning in result["warnings"]))

    def test_existing_output_cannot_be_replaced(self):
        source = self.make_docx()
        Path("outputs/existing.pdf").write_bytes(b"KEEP EXISTING")
        result = converter.convert(str(source), "outputs/existing.pdf")
        self.assertEqual(result["code"], "E_PARAMS")
        self.assertEqual(Path("outputs/existing.pdf").read_bytes(), b"KEEP EXISTING")

    def test_path_escape_and_write_to_uploads_are_rejected(self):
        source = self.make_docx()
        for destination in ("../escape.pdf", "uploads/overwrite.pdf", "/tmp/escape.pdf"):
            with self.subTest(destination=destination):
                result = converter.convert(str(source), destination)
                self.assertEqual(result["code"], "E_PARAMS")

    def test_symlink_escape_is_rejected(self):
        source = self.make_docx()
        Path("outputs/link").symlink_to(self.root.parent, target_is_directory=True)
        result = converter.convert(str(source), "outputs/link/escape.pdf")
        self.assertEqual(result["code"], "E_PARAMS")

    def test_wrong_extension_and_malicious_playlist_are_not_output(self):
        for name, content in (("fake.pdf", b"not a pdf"),
                              ("fake.mp3", b"#EXTM3U\nhttp://127.0.0.1:1/private\n")):
            Path("uploads", name).write_bytes(content)
            destination = "outputs/fake.docx" if name.endswith("pdf") else "outputs/fake.mp4"
            result = converter.convert("uploads/" + name, destination)
            self.assertFalse(result["ok"], result)
            self.assertFalse(Path(destination).exists())

    def test_missing_dependency_returns_recovery_not_success(self):
        source = self.make_mp3()
        with patch.object(converter.shutil, "which", return_value=None):
            result = converter.convert(str(source), "outputs/video.mp4")
        self.assertEqual(result["code"], "E_PROVIDER")
        self.assertIn("imagen oficial", result["next_action"])
        self.assertFalse(Path("outputs/video.mp4").exists())

    def test_timeout_is_bounded_and_does_not_publish(self):
        source = self.make_docx()
        result = converter.convert(str(source), "outputs/result.pdf", timeout=0.000001)
        self.assertEqual(result["code"], "E_TIMEOUT")
        self.assertFalse(Path("outputs/result.pdf").exists())

    def test_unsupported_pair_is_not_renamed(self):
        source = self.make_docx()
        result = converter.convert(str(source), "outputs/result.mp4")
        self.assertEqual(result["code"], "E_PARAMS")
        self.assertFalse(Path("outputs/result.mp4").exists())


if __name__ == "__main__":
    unittest.main()
