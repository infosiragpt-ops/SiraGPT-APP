"""Native chart acceptance: workbook bytes reopened by openpyxl, no model or DB."""

import copy
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'src/services/agent-runner'))

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, PieChart, Reference
from sira_charts import add_xlsx_chart, apply_xlsx_chart_colors
import sira_design as design

NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
COLORS = ['1122CC', 'EE8800', '229944']


def fixture():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resultados reales'
    for row in [('Mes', 'Ingresos', 'Costes', 'Medición X'),
                ('Enero', 120, 80, 1), ('Febrero', 150, 95, 2), ('Marzo', 180, 110, 3)]:
        ws.append(row)
    return wb, ws


def data_signature(ws):
    return [(cell.coordinate, cell.value) for row in ws for cell in row if cell.value is not None]


def reference_signature(chart):
    refs = []
    for plot in chart._charts:
        for series in plot.series:
            for prop in ('tx', 'cat', 'val', 'xVal', 'yVal'):
                node = getattr(series, prop, None)
                if node is not None:
                    refs.append((prop, ET.tostring(node.to_tree(tagname=prop))))
    return refs


class NativeXlsxCharts(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='sira-native-charts-')
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_all_seven_types_are_native_editable_charts_with_real_references_and_palette(self):
        wb, ws = fixture()
        before = data_signature(ws)
        kinds = ['column', 'bar', 'line', 'area', 'pie', 'doughnut', 'scatter']
        for index, kind in enumerate(kinds):
            kwargs = {'x_range': 'D2:D4'} if kind == 'scatter' else {'category_range': 'A2:A4'}
            add_xlsx_chart(ws, chart_type=kind,
                           data_range='B1:B4' if kind in ('pie', 'doughnut') else 'B1:C4',
                           colors=COLORS, title='Datos originales - ' + kind, legend='b',
                           data_labels={'showPercent': True} if kind in ('pie', 'doughnut') else True,
                           anchor='F%d' % (index * 18 + 2), width=17, height=9, **kwargs)
        output = self.root / 'seven-types.xlsx'
        wb.save(output)
        actual = load_workbook(output).active
        self.assertEqual(data_signature(actual), before)
        self.assertEqual(len(actual._charts), 7)
        self.assertEqual([type(c).__name__ for c in actual._charts],
                         ['BarChart', 'BarChart', 'LineChart', 'AreaChart', 'PieChart', 'DoughnutChart', 'ScatterChart'])
        self.assertEqual(actual._charts[0].type, 'col')
        self.assertEqual(actual._charts[1].type, 'bar')
        self.assertIs(actual._charts[2].smooth, False)
        self.assertTrue(all(s.smooth is False for s in actual._charts[2].series), 'line defaults must not invent smoothed interpolation')
        for index, chart in enumerate(actual._charts):
            self.assertEqual(chart.legend.position, 'b')
            self.assertEqual(chart.anchor._from.col, 5)
            self.assertEqual(chart.anchor._from.row, index * 18 + 1)
            self.assertEqual(chart.anchor.ext.cx, 17 * 360000)
            self.assertEqual(chart.anchor.ext.cy, 9 * 360000)
            self.assertEqual(chart.series[0].graphicalProperties.solidFill.srgbClr, COLORS[0])
            if kinds[index] in ('pie', 'doughnut'):
                self.assertEqual([p.graphicalProperties.solidFill.srgbClr for p in chart.series[0].data_points], COLORS)
                self.assertTrue(chart.dataLabels.showPercent)
            else:
                self.assertEqual(len(chart.series), 2)
                self.assertEqual(chart.series[1].graphicalProperties.line.solidFill.srgbClr, COLORS[1])
                self.assertTrue(chart.dataLabels.showVal)
            value = chart.series[0].yVal if kinds[index] == 'scatter' else chart.series[0].val
            self.assertEqual(value.numRef.f, "'Resultados reales'!$B$2:$B$4")
            self.assertEqual(chart.series[0].tx.strRef.f, "'Resultados reales'!B1")
        self.assertEqual(actual._charts[-1].series[0].xVal.numRef.f, "'Resultados reales'!$D$2:$D$4")
        with zipfile.ZipFile(output) as z:
            self.assertIsNone(z.testzip())
            chart_parts = [n for n in z.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')]
            self.assertEqual(len(chart_parts), 7)
            self.assertFalse(any(n.startswith('xl/media/') for n in z.namelist()), 'editable charts must not be PNG stand-ins')
            for part in chart_parts:
                xml = ET.fromstring(z.read(part))
                self.assertTrue(xml.findall('.//c:f', NS), 'native series must reference the source worksheet')

    def test_formula_and_null_cells_remain_formulas_and_gaps_without_data_coercion(self):
        wb, ws = fixture()
        ws['B3'] = None
        ws['C3'] = '=B2*0.5'
        before = data_signature(ws)
        add_xlsx_chart(ws, chart_type='line', data_range=['B1:B4', 'C1:C4'], category_range='A2:A4', colors=COLORS)
        out = self.root / 'gaps.xlsx'
        wb.save(out)
        result = load_workbook(out).active
        self.assertEqual(data_signature(result), before)
        self.assertIsNone(result['B3'].value)
        self.assertEqual(result['C3'].value, '=B2*0.5')
        self.assertEqual(result._charts[0].display_blanks, 'gap')

    def test_recolor_existing_chart_preserves_type_data_refs_geometry_and_values(self):
        wb, ws = fixture()
        chart = LineChart()
        chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
        ws.add_chart(chart, 'G3')
        src = self.root / 'before.xlsx'
        wb.save(src)
        reopened = load_workbook(src)
        chart = reopened.active._charts[0]
        refs = reference_signature(chart)
        anchor = ET.tostring(chart.anchor.to_tree())
        before = data_signature(reopened.active)
        apply_xlsx_chart_colors(chart, COLORS)
        dst = self.root / 'after.xlsx'
        reopened.save(dst)
        edited = load_workbook(dst).active
        self.assertEqual(data_signature(edited), before)
        self.assertIsInstance(edited._charts[0], LineChart)
        self.assertEqual(reference_signature(edited._charts[0]), refs)
        self.assertEqual(ET.tostring(edited._charts[0].anchor.to_tree()), anchor)
        self.assertEqual([s.graphicalProperties.line.solidFill.srgbClr for s in edited._charts[0].series], COLORS[:2])

    def test_point_palette_is_explicitly_supported_for_a_single_bar_series(self):
        wb, ws = fixture()
        chart = add_xlsx_chart(ws, chart_type='bar', data_range='B1:B4', category_range='A2:A4', colors=COLORS,
                               color_by='point', legend=False, style=13)
        self.assertIsNone(chart.legend)
        self.assertEqual(chart.style, 13)
        self.assertEqual([p.graphicalProperties.solidFill.srgbClr for p in chart.series[0].data_points], COLORS)

    def test_recolor_replaces_existing_fill_choices_instead_of_serializing_contradictory_colors(self):
        from openpyxl.chart.marker import DataPoint
        from openpyxl.drawing.fill import GradientFillProperties, PatternFillProperties
        wb, ws = fixture()
        chart = LineChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
        series = chart.series[0]
        series.graphicalProperties.gradFill = GradientFillProperties()
        series.graphicalProperties.line.noFill = True
        series.marker.symbol = 'circle'
        series.marker.graphicalProperties.pattFill = PatternFillProperties(prst='pct50')
        series.marker.graphicalProperties.line.gradFill = GradientFillProperties()
        point = DataPoint(idx=0)
        point.graphicalProperties.noFill = True
        point.graphicalProperties.line.pattFill = PatternFillProperties(prst='pct50')
        series.data_points = [point]
        ws.add_chart(chart, 'G3')
        src, dst = self.root / 'old-fills.xlsx', self.root / 'solid-colors.xlsx'
        wb.save(src)
        edited = load_workbook(src)
        before = reference_signature(edited.active._charts[0])
        apply_xlsx_chart_colors(edited.active._charts[0], [COLORS[0]])
        edited.save(dst)
        self.assertEqual(reference_signature(load_workbook(dst).active._charts[0]), before)
        with zipfile.ZipFile(dst) as z:
            series = ET.fromstring(z.read('xl/charts/chart1.xml')).find('.//c:ser', NS)
            styles = series.findall('.//c:spPr', NS)
            self.assertEqual(len(styles), 3, 'series, marker and data-point styles remain present')
            for style in styles:
                for props in [style, style.find('a:ln', NS)]:
                    fills = [el for el in props if el.tag.rsplit('}', 1)[-1] in ('noFill', 'solidFill', 'gradFill', 'pattFill')]
                    self.assertEqual([el.tag.rsplit('}', 1)[-1] for el in fills], ['solidFill'])
                    self.assertEqual(fills[0].find('a:srgbClr', NS).get('val'), COLORS[0])

    def test_invalid_ranges_types_colors_and_options_fail_without_attaching_or_changing_data(self):
        cases = [
            {'chart_type': 'waterfall'}, {'data_range': 'B1:C99'}, {'data_range': 'B4:B1'},
            {'data_range': ['B1:B4', 'C1:C3']}, {'data_range': ['B1:B4', 'B1:B4']},
            {'category_range': 'A1:A4'}, {'category_range': None}, {'category_range': 'Other!A2:A4'},
            {'chart_type': 'pie'}, {'chart_type': 'scatter'}, {'colors': ['not-a-color']}, {'colors': []},
            {'color_by': 'random'}, {'width': -1}, {'height': float('nan')}, {'legend': 'middle'},
            {'data_labels': {'unsupported': True}}, {'data_labels': {'showPercent': True}},
            {'title': 42}, {'anchor': 'bad'}, {'anchor': 'ZZZ2'}, {'anchor': 'F1048577'}, {'style': 49},
        ]
        for options in cases:
            with self.subTest(options=options):
                _, ws = fixture()
                before = data_signature(ws)
                args = {'chart_type': 'line', 'data_range': 'B1:C4', 'category_range': 'A2:A4'}
                args.update(options)
                with self.assertRaises(ValueError):
                    add_xlsx_chart(ws, **args)
                self.assertEqual(ws._charts, [])
                self.assertEqual(data_signature(ws), before)
        for invalid in ['sin dato', True, float('inf')]:
            _, ws = fixture()
            ws['B2'] = invalid
            with self.assertRaises(ValueError):
                add_xlsx_chart(ws, chart_type='column', data_range='B1:B4', category_range='A2:A4')
            self.assertEqual(ws._charts, [])
        for kind in ('pie', 'doughnut'):
            for data in [(0, 0, 0), (10, -1, 20), (10, None, 20)]:
                _, ws = fixture()
                for row, value in enumerate(data, start=2):
                    ws.cell(row, 2).value = value
                with self.assertRaises(ValueError):
                    add_xlsx_chart(ws, chart_type=kind, data_range='B1:B4', category_range='A2:A4')
                self.assertEqual(ws._charts, [])

    def test_restyle_updates_existing_chart_palette_without_replacing_chart_or_data(self):
        wb, ws = fixture()
        ws['C3'] = '=B3*0.5'
        native = LineChart()
        native.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
        native.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
        for series in native.series:
            series.graphicalProperties.line.solidFill = 'FF0000'
        ws.add_chart(native, 'G3')
        src, dst = self.root / 'styled-source.xlsx', self.root / 'styled-output.xlsx'
        wb.save(src)
        old = load_workbook(src).active
        refs, values = reference_signature(old._charts[0]), data_signature(old)
        theme = copy.deepcopy(design.DEFAULT_THEME)
        theme['chartColors'] = COLORS
        report = design.restyle_xlsx(str(src), str(dst), theme)
        self.assertTrue(report['ok'])
        self.assertEqual(report['warnings'], [])
        result = load_workbook(dst).active
        self.assertEqual(len(result._charts), 1)
        self.assertIsInstance(result._charts[0], LineChart)
        self.assertEqual(reference_signature(result._charts[0]), refs)
        self.assertEqual(data_signature(result), values)
        self.assertEqual([s.graphicalProperties.line.solidFill.srgbClr for s in result._charts[0].series], COLORS[:2])
        self.assertEqual(report['sheets'][0]['charts_recolored'], 1)

    def test_restyle_adds_temporal_multiseries_chart_with_requested_palette(self):
        wb, ws = fixture()
        # An unrelated numeric ID must not become a measured series.
        ws['D1'] = 'ID'
        src, dst = self.root / 'table.xlsx', self.root / 'redesigned.xlsx'
        before = data_signature(ws)
        wb.save(src)
        theme = copy.deepcopy(design.DEFAULT_THEME)
        theme['chartColors'] = COLORS
        report = design.restyle_xlsx(str(src), str(dst), theme)
        self.assertTrue(report['ok'])
        self.assertEqual(report['warnings'], [])
        result = load_workbook(dst).active
        self.assertEqual(data_signature(result), before)
        self.assertEqual(len(result._charts), 1)
        chart = result._charts[0]
        self.assertIsInstance(chart, LineChart)
        self.assertEqual(len(chart.series), 2)
        self.assertEqual(chart.legend.position, 'b')
        self.assertEqual([s.graphicalProperties.line.solidFill.srgbClr for s in chart.series], COLORS[:2])
        self.assertEqual([s.val.numRef.f for s in chart.series], ["'Resultados reales'!$B$2:$B$4", "'Resultados reales'!$C$2:$C$4"])


if __name__ == '__main__':
    unittest.main()
