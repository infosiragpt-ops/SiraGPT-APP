"""Gráficas nativas: artefactos OOXML reales, fuentes seguras y verificación exacta."""
import copy
import io
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import zipfile
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Mm

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'src/services/agent-runner'))
import sira_office as so

COLORS = ['1F4E78', 'ED7D31', '70AD47']
CATEGORIES = ['T1', 'T2', 'T3', 'T4']
NAMES = ['Norte', 'Centro', 'Sur']
VALUES = [[120, 135, 128, 150], [90, 95, 105, 120], [60, 72, 75, 90]]
HAS_RENDER = bool(shutil.which('soffice') and shutil.which('pdftoppm'))


def fixtures(folder):
    xlsx, pptx = str(Path(folder) / 'charts.xlsx'), str(Path(folder) / 'charts.pptx')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Datos'
    ws.append(['Periodo', *NAMES])
    for i, label in enumerate(CATEGORIES):
        ws.append([label, *(values[i] for values in VALUES)])
    chart = BarChart()
    chart.type, chart.grouping = 'col', 'clustered'
    chart.title = 'Ventas trimestrales'
    chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    for series, color in zip(chart.series, COLORS):
        series.graphicalProperties.solidFill = color
    ws.add_chart(chart, 'F2')
    wb.save(xlsx)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = CATEGORIES
    for name, values in zip(NAMES, VALUES):
        data.add_series(name, values)
    graphic = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Mm(20), Mm(25), Mm(180), Mm(110), data)
    chart = graphic.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = 'Ventas trimestrales'
    chart.has_legend = True
    for series, color in zip(chart.series, COLORS):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
    prs.save(pptx)
    return xlsx, pptx


def rewrite(src, dst, change):
    with zipfile.ZipFile(src) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    change(parts)
    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return dst


def xml_change(parts, part, fn):
    root = etree.fromstring(parts[part])
    fn(root)
    parts[part] = etree.tostring(root)


class TestNativeCharts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.mkdtemp(prefix='sira-charts-')
        cls.xlsx, cls.pptx = fixtures(cls.folder)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.folder)

    def path(self, name):
        return str(Path(self.folder) / (self._testMethodName + name))

    def expected(self, fmt):
        return {**({'sheet': 'Datos'} if fmt == 'xlsx' else {'slide': 1}), 'chart': 1,
                'type': 'column', 'grouping': 'clustered', 'editable': True,
                'title': 'Ventas trimestrales', 'legend': True, 'categories': CATEGORIES,
                'series': [{'name': name, 'values': values, 'color': color} for name, values, color in zip(NAMES, VALUES, COLORS)]}

    def inspect_chart(self, path, fmt):
        out = so.inspect(path)
        return out['sheets' if fmt == 'xlsx' else 'slides'][0]['charts'][0]

    def checks(self, path, fmt, expected):
        return so.chart_checks(so.OfficePackage(path), fmt, [expected])

    def test_real_xlsx_native_data_color_anchor_and_stable_selector(self):
        chart = self.inspect_chart(self.xlsx, 'xlsx')
        self.assertTrue(chart['complete'])
        self.assertTrue(chart['editable'])
        self.assertEqual(chart['type'], 'column')
        self.assertEqual(chart['categories'], CATEGORIES)
        self.assertEqual(chart['series'][0]['values'], VALUES[0])
        self.assertEqual(chart['series'][0]['values_ref'], "'Datos'!$B$2:$B$5")
        self.assertEqual(chart['series'][0]['color'], COLORS[0])
        self.assertEqual(chart['position']['from']['col'], 5)
        self.assertEqual(chart['position']['from']['row'], 1)
        self.assertEqual(len(load_workbook(self.xlsx).active._charts), 1)
        expected = self.expected('xlsx')
        expected['chart'] = chart['id']
        self.assertTrue(self.checks(self.xlsx, 'xlsx', expected)[0]['ok'])

    def test_real_pptx_native_workbook_data_position_and_selector(self):
        chart = self.inspect_chart(self.pptx, 'pptx')
        self.assertTrue(chart['complete'], chart)
        self.assertTrue(chart['editable'])
        self.assertEqual(chart['type'], 'column')
        self.assertEqual(chart['series'][2]['values'], VALUES[2])
        self.assertEqual(chart['series'][1]['color'], COLORS[1])
        self.assertEqual(chart['position'], {'x_mm': 20.0, 'y_mm': 25.0, 'w_mm': 180.0, 'h_mm': 110.0})
        self.assertTrue(Presentation(self.pptx).slides[0].shapes[0].has_chart)
        expected = self.expected('pptx')
        expected['position'] = chart['position']
        self.assertTrue(self.checks(self.pptx, 'pptx', expected)[0]['ok'])

    def test_wrong_type_data_palette_categories_position_or_missing_chart_fails(self):
        for fmt, path in [('xlsx', self.xlsx), ('pptx', self.pptx)]:
            for mutation in ['type', 'values', 'color', 'categories', 'missing', 'position']:
                with self.subTest(fmt=fmt, mutation=mutation):
                    expected = self.expected(fmt)
                    if mutation == 'type': expected['type'] = 'line'
                    if mutation == 'values': expected['series'][0]['values'] = [120, 135, 128, 999]
                    if mutation == 'color': expected['series'][1]['color'] = '000000'
                    if mutation == 'categories': expected['categories'] = ['A', 'B', 'C', 'D']
                    if mutation == 'missing': expected['chart'] = 2
                    if mutation == 'position': expected['position'] = {'w_mm': 2}
                    self.assertFalse(self.checks(path, fmt, expected)[0]['ok'])

    def test_missing_native_chart_cannot_be_certified(self):
        dst = self.path('.xlsx')
        wb = Workbook()
        wb.active.title = 'Datos'
        wb.save(dst)
        self.assertFalse(self.checks(dst, 'xlsx', self.expected('xlsx'))[0]['ok'])

    def test_picture_is_not_an_editable_native_chart(self):
        from PIL import Image
        from openpyxl.drawing.image import Image as XlImage
        png = self.path('.png')
        Image.new('RGB', (200, 100), '#1F4E78').save(png)
        dst = self.path('.xlsx')
        wb = Workbook()
        wb.active.title = 'Datos'
        wb.active.add_image(XlImage(png), 'F2')
        wb.save(dst)
        self.assertFalse(self.checks(dst, 'xlsx', self.expected('xlsx'))[0]['ok'])
        dst = self.path('.pptx')
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(png, Mm(20), Mm(25))
        prs.save(dst)
        self.assertFalse(self.checks(dst, 'pptx', self.expected('pptx'))[0]['ok'])

    def test_pptx_cache_must_match_embedded_workbook(self):
        dst = self.path('.pptx')
        rewrite(self.pptx, dst, lambda p: xml_change(p, 'ppt/charts/chart1.xml',
            lambda r: setattr(r.find('.//c:ser/c:val/c:numRef/c:numCache/c:pt/c:v', so.NS), 'text', '999')))
        chart = self.inspect_chart(dst, 'pptx')
        self.assertFalse(chart['complete'])
        self.assertIn('no coinciden', chart['error'])
        self.assertFalse(self.checks(dst, 'pptx', self.expected('pptx'))[0]['ok'])

    def test_external_and_escape_relationships_never_certify(self):
        for mode, target in [('External', 'file:///etc/passwd'), ('External', '../embeddings/Microsoft_Excel_Sheet1.xlsx'),
                             ('Internal', '../../../../outside.xlsx'), ('Internal', 'https://example.com/a.xlsx'),
                             ('Internal', '%2E%2E/x.xlsx')]:
            with self.subTest(mode=mode, target=target):
                dst = self.path(str(len(target)) + mode + '.pptx')
                def change(parts):
                    def edit(root):
                        for rel in root:
                            rel.set('TargetMode', mode)
                            rel.set('Target', target)
                    xml_change(parts, 'ppt/charts/_rels/chart1.xml.rels', edit)
                rewrite(self.pptx, dst, change)
                chart = self.inspect_chart(dst, 'pptx')
                self.assertFalse(chart['editable'])
                self.assertFalse(self.checks(dst, 'pptx', self.expected('pptx'))[0]['ok'])

    def test_external_drawing_and_external_series_are_not_followed(self):
        for kind in ['drawing', 'values']:
            dst = self.path(kind + '.xlsx')
            def change(parts):
                if kind == 'drawing':
                    xml_change(parts, 'xl/worksheets/_rels/sheet1.xml.rels', lambda r: r[0].set('TargetMode', 'External'))
                else:
                    xml_change(parts, 'xl/charts/chart1.xml', lambda r: setattr(r.find('.//c:val/c:numRef/c:f', so.NS), 'text', "'[external.xlsx]Datos'!$B$2:$B$5"))
            rewrite(self.xlsx, dst, change)
            self.assertFalse(self.checks(dst, 'xlsx', self.expected('xlsx'))[0]['ok'])

    def test_formula_without_cache_does_not_invent_a_value(self):
        dst = self.path('.xlsx')
        wb = load_workbook(self.xlsx)
        wb.active['B2'] = '=100+20'
        wb.save(dst)
        chart = self.inspect_chart(dst, 'xlsx')
        self.assertFalse(chart['complete'])
        self.assertIn('sin recalcular', chart['error'])
        self.assertFalse(self.checks(dst, 'xlsx', self.expected('xlsx'))[0]['ok'])

    def test_zero_negative_and_missing_values_remain_distinct(self):
        dst = self.path('.xlsx')
        wb = load_workbook(self.xlsx)
        ws = wb.active
        ws['B2'], ws['B3'], ws['B4'], ws['B5'] = 0, -5, None, 10
        wb.save(dst)
        chart = self.inspect_chart(dst, 'xlsx')
        self.assertEqual(chart['series'][0]['values'], [0, -5, None, 10])
        expected = self.expected('xlsx')
        expected['series'][0]['values'] = [0, -5, None, 10]
        self.assertTrue(self.checks(dst, 'xlsx', expected)[0]['ok'])
        expected['series'][0]['values'] = [0, -5, 0, 10]
        self.assertFalse(self.checks(dst, 'xlsx', expected)[0]['ok'])

    def test_native_pie_per_point_colors(self):
        dst = self.path('.xlsx')
        wb = load_workbook(self.xlsx)
        ws = wb.active
        ws._charts = []
        chart = PieChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
        colors = COLORS + ['A5A5A5']
        for i, color in enumerate(colors):
            point = DataPoint(idx=i)
            point.graphicalProperties.solidFill = color
            chart.series[0].data_points.append(point)
        ws.add_chart(chart, 'F2')
        wb.save(dst)
        actual = self.inspect_chart(dst, 'xlsx')
        self.assertEqual(actual['series'][0]['point_colors'], colors)
        expected = {'sheet': 'Datos', 'chart': 1, 'type': 'pie', 'series': [{'point_colors': colors, 'values': VALUES[0]}]}
        self.assertTrue(self.checks(dst, 'xlsx', expected)[0]['ok'])
        expected['series'][0]['point_colors'][0] = '000000'
        self.assertFalse(self.checks(dst, 'xlsx', expected)[0]['ok'])

    def test_line_color_requires_the_line_not_a_series_fill(self):
        dst = self.path('.xlsx')
        wb = load_workbook(self.xlsx)
        ws = wb.active
        ws._charts = []
        chart = LineChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
        chart.series[0].graphicalProperties.solidFill = COLORS[0]
        ws.add_chart(chart, 'F2')
        wb.save(dst)
        actual = self.inspect_chart(dst, 'xlsx')
        self.assertIsNone(actual['series'][0]['color'])
        expected = {'sheet': 'Datos', 'chart': 1, 'type': 'line', 'series': [{'color': COLORS[0]}]}
        self.assertFalse(self.checks(dst, 'xlsx', expected)[0]['ok'])
        chart.series[0].graphicalProperties.line.solidFill = COLORS[0]
        wb.save(dst)
        self.assertTrue(self.checks(dst, 'xlsx', expected)[0]['ok'])

    def test_conflicting_fill_cannot_certify_requested_color(self):
        dst = self.path('.xlsx')
        def change(parts):
            def edit(root):
                properties = root.find('.//c:ser/c:spPr', so.NS)
                etree.SubElement(properties, so.A('noFill'))
            xml_change(parts, 'xl/charts/chart1.xml', edit)
        rewrite(self.xlsx, dst, change)
        chart = self.inspect_chart(dst, 'xlsx')
        self.assertFalse(chart['complete'])
        self.assertIn('ambiguo', chart['error'])
        self.assertFalse(self.checks(dst, 'xlsx', self.expected('xlsx'))[0]['ok'])

    def test_series_color_checks_visible_point_overrides_unless_explicit(self):
        dst = self.path('.xlsx')
        wb = load_workbook(self.xlsx)
        point = DataPoint(idx=0)
        point.graphicalProperties.solidFill = '0000FF'
        wb.active._charts[0].series[0].data_points.append(point)
        wb.save(dst)
        expected = self.expected('xlsx')
        self.assertFalse(self.checks(dst, 'xlsx', expected)[0]['ok'])
        expected['series'][0]['point_colors'] = ['0000FF', COLORS[0], COLORS[0], COLORS[0]]
        self.assertTrue(self.checks(dst, 'xlsx', expected)[0]['ok'])

    def test_embedded_zip_expansion_is_bounded(self):
        dst = self.path('.pptx')
        def change(parts):
            target = next(n for n in parts if n.startswith('ppt/embeddings/'))
            with zipfile.ZipFile(io.BytesIO(parts[target])) as z:
                data = {n: z.read(n) for n in z.namelist()}
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
                for name, value in data.items():
                    z.writestr(name, value)
                z.writestr('oversize.xml', b'x' * (8 * 1024 * 1024 + 1))
            parts[target] = buf.getvalue()
        rewrite(self.pptx, dst, change)
        actual = self.inspect_chart(dst, 'pptx')
        self.assertFalse(actual['editable'])
        self.assertIn('descompresión', actual['error'])

    def test_inventory_bound_is_explicit_and_partial_arrays_cannot_verify(self):
        dst = self.path('.xlsx')
        wb = load_workbook(self.xlsx)
        ws = wb.active
        for row in range(2, 202):
            ws.cell(row, 1, '\\"' * 240 + str(row))
            ws.cell(row, 2, row)
        ws._charts = []
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=201), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=201))
        ws.add_chart(chart, 'F2')
        wb.save(dst)
        actual = self.inspect_chart(dst, 'xlsx')
        self.assertFalse(actual['complete'])
        self.assertTrue(actual['truncated'])
        self.assertLess(len(json.dumps(actual)), 2000)
        self.assertFalse(self.checks(dst, 'xlsx', {'sheet': 'Datos', 'chart': 1, 'type': 'column'})[0]['ok'])

    def test_complete_chart_in_truncated_inventory_is_not_certified(self):
        dst = self.path('.xlsx')
        wb = load_workbook(self.xlsx)
        ws = wb.active
        for i in range(20):
            ws.add_chart(copy.deepcopy(ws._charts[0]), f'F{30 + i * 20}')
        wb.save(dst)
        inventory = so.native_charts(so.OfficePackage(dst), 'xlsx', sheet='Datos')
        self.assertTrue(inventory['charts'][0]['complete'])
        self.assertTrue(inventory['truncated'])
        self.assertFalse(self.checks(dst, 'xlsx', self.expected('xlsx'))[0]['ok'])

    @unittest.skipUnless(HAS_RENDER, 'LibreOffice y Poppler requeridos para recálculo real')
    def test_final_recalculated_file_can_verify_formula_chart(self):
        src = self.path('.xlsx')
        wb = load_workbook(self.xlsx)
        wb.active['B2'] = '=100+20'
        wb.save(src)
        self.assertFalse(self.checks(src, 'xlsx', self.expected('xlsx'))[0]['ok'])
        final = so.soffice_convert(src, self.path('-recalculated'), 'xlsx')
        self.assertEqual(load_workbook(final, data_only=True).active['B2'].value, 120)
        report = self.checks(final, 'xlsx', self.expected('xlsx'))
        self.assertTrue(report[0]['ok'], report)

    def test_oversized_data_fails_without_silent_truncation(self):
        dst = self.path('.xlsx')
        rewrite(self.xlsx, dst, lambda p: xml_change(p, 'xl/charts/chart1.xml',
            lambda r: setattr(r.find('.//c:val/c:numRef/c:f', so.NS), 'text', "'Datos'!$B$2:$B$999999")))
        chart = self.inspect_chart(dst, 'xlsx')
        self.assertFalse(chart['complete'])
        self.assertFalse(self.checks(dst, 'xlsx', self.expected('xlsx'))[0]['ok'])
        self.assertLess(len(json.dumps(chart)), 2000)

    @unittest.skipUnless(HAS_RENDER, 'LibreOffice y Poppler requeridos para render real')
    def test_verify_real_render_fails_native_mismatch_without_any_vision(self):
        for fmt, path in [('xlsx', self.xlsx), ('pptx', self.pptx)]:
            with self.subTest(fmt=fmt):
                expected = self.expected(fmt)
                expected['type'] = 'line'
                report = so.verify(None, path, self.path('-render-' + fmt), dpi=50, expect={'charts': [expected]})
                self.assertGreater(report['visual']['page_count_after'], 0)
                self.assertFalse(report['ok'])
                self.assertFalse(report['checks'][-1]['ok'])
                self.assertIn('type', report['checks'][-1]['detail'])


if __name__ == '__main__':
    unittest.main()
