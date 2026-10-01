#!/usr/bin/env python3
"""sira_design — deterministic professional restyle of an EXISTING Office file.

Used by the AgentRunner DESIGN WORKFLOW («agrégale más diseño a la ppt»,
«mejora el formato del excel», «hazlo más profesional el word»). It changes
the LOOK of the same file and never its content: every text, number,
formula, image, slide/page/sheet and its order is kept. The output is a new
version of the same format: <stem>-v2.<ext> (or -v(N+1) when the name already
ends in -vN).

    import sys, json; sys.path.insert(0, '/workspace/tmp'); import sira_design as sd
    print(json.dumps(sd.restyle('uploads/deck.pptx')))

The theme defaults to /workspace/tmp/sira_theme.json (written by the runner
from document-pipeline/pptx-design-system.js) and falls back to «aurora».

Dependencies: python-pptx, python-docx, openpyxl (all in the sandbox image).
Each file type imports its library lazily, so a missing library only affects
that format. Best-effort per element: a shape/table the helper cannot style
is left as it was and reported in `warnings`.
"""

import json
import os
import re
import sys

DEFAULT_THEME = {
    'id': 'aurora',
    'fonts': {'display': 'Calibri', 'body': 'Calibri'},
    'palette': {
        'bg': 'F8FAFC', 'surface': 'FFFFFF', 'surfaceAlt': 'EFF6FF', 'ink': '0F172A',
        'body': '334155', 'muted': '64748B', 'line': 'E2E8F0', 'accent': '2563EB',
        'accent2': '06B6D4', 'chipLine': 'BFDBFE', 'coverBg': 'EEF6FF', 'coverInk': '0F172A',
        'coverMuted': '334155', 'sectionBg': '0F172A', 'sectionInk': 'FFFFFF',
        'sectionMuted': 'CBD5E1', 'inverse': 'FFFFFF',
    },
    'chartColors': ['2563EB', '06B6D4', '8B5CF6', '10B981', 'F59E0B'],
    'coverStyle': 'light',
}

# Aptos only ships with recent Office builds and the render sandbox has no
# metric-compatible substitute: Calibri renders the same in PowerPoint, Word,
# Excel and LibreOffice (Carlito), so the verification image matches.
FONT_FALLBACKS = {'aptos display': 'Calibri', 'aptos': 'Calibri'}

HEX_RE = re.compile(r'^[0-9A-Fa-f]{6}$')
DECO_PREFIX = 'SiraDeco'
DECO_THEME_RE = re.compile(r'^SiraDeco\[([^\]]+)\]')
KEEP_PREFIXES = ('SiraCard', 'SiraChip', 'SiraKpi')
CLOSING_RE = re.compile(r'^\s*(gracias|muchas gracias|thank(s| you)|preguntas|q\s*&\s*a|fin)\b', re.I)
# A KPI is a real METRIC followed by its label with a plain space («35 %
# reducción…», «$2,4 M en ahorro»). Numbered outlines («1. Planificación»),
# bare counts («5 estrategias») and «value: label» forms stay text.
KPI_RE = re.compile(
    r'^\s*((?:S/\.?\s?|US\$\s?|[$€£]\s?)?\d[\d.,]*(?:\s?(?:%|k|m|mm|mil|millones|mill|bn|b|x|pts|puntos|pp))?)\s+(\S.{2,})$',
    re.I,
)
METRIC_RE = re.compile(r'%|[$€£]|S/|US\$|\d[.,]\d|\d\s?(?:k|m|mm|mil|millones|mill|bn|b|x|pts|puntos|pp)$', re.I)
LIST_MARKER_RE = re.compile(r'^\d{1,2}[.)]$')
YEAR_RE = re.compile(r'^(19|20)\d{2}$')
FIGURE_ONLY_RE = re.compile(r'^(?:S/\.?|[$€£])?\s?[\d.,]+\s?(?:%|k|m|mm|bn|b|x)?$', re.I)

# Alternative themes (same tokens as document-pipeline/pptx-design-system.js)
# used when the source is ALREADY a SiraGPT redesign with the chosen theme:
# a second «más diseño» must look different, not produce an identical -v3.
BUILTIN_ALTERNATES = [
    {
        'id': 'consulting', 'fonts': {'display': 'Calibri', 'body': 'Calibri'}, 'coverStyle': 'light',
        'palette': {
            'bg': 'FFFFFF', 'surface': 'FFFFFF', 'surfaceAlt': 'F1F5F9', 'ink': '0C2340', 'body': '334155',
            'muted': '6B7280', 'line': 'D8DEE9', 'accent': '1E3A5F', 'accent2': '2E75B6', 'chipLine': 'C7D2E0',
            'coverBg': 'FFFFFF', 'coverInk': '0C2340', 'coverMuted': '46596E', 'sectionBg': '0C2340',
            'sectionInk': 'FFFFFF', 'sectionMuted': 'B7C4D6', 'inverse': 'FFFFFF',
        },
        'chartColors': ['1E3A5F', '2E75B6', '6B93B8', '94A9C0', 'C4CFDC'],
    },
    {
        'id': 'editorial', 'fonts': {'display': 'Georgia', 'body': 'Calibri'}, 'coverStyle': 'light',
        'palette': {
            'bg': 'FAF7F2', 'surface': 'FFFFFF', 'surfaceAlt': 'F2EBE0', 'ink': '1C1917', 'body': '44403C',
            'muted': '78716C', 'line': 'E7DFD2', 'accent': '15803D', 'accent2': 'C2571B', 'chipLine': 'D6CCBB',
            'coverBg': 'F2EBE0', 'coverInk': '1C1917', 'coverMuted': '57534E', 'sectionBg': '14532D',
            'sectionInk': 'FDFCF9', 'sectionMuted': 'BBF7D0', 'inverse': 'FFFFFF',
        },
        'chartColors': ['15803D', 'C2571B', '0F766E', 'A16207', '57534E'],
    },
    {
        'id': 'boardroom', 'fonts': {'display': 'Calibri', 'body': 'Calibri'}, 'coverStyle': 'dark',
        'palette': {
            'bg': '0B1220', 'surface': '111B2E', 'surfaceAlt': '16233B', 'ink': 'F8FAFC', 'body': 'CBD5E1',
            'muted': '9FB0C8', 'line': '1F2E48', 'accent': 'D9A441', 'accent2': '5EA0EF', 'chipLine': '2C3E5D',
            'coverBg': '0B1220', 'coverInk': 'F8FAFC', 'coverMuted': '9FB0C8', 'sectionBg': '060B14',
            'sectionInk': 'F8FAFC', 'sectionMuted': '8FA1B8', 'inverse': '0B1220',
        },
        'chartColors': ['D9A441', '5EA0EF', '34D399', 'F472B6', '94A3B8'],
    },
]


# ── paths / theme ─────────────────────────────────────────────────────────

def _ws(path):
    """Map a /workspace path to the local workspace when /workspace is absent."""
    p = str(path or '')
    if p.startswith('/workspace/') and not os.path.isdir('/workspace'):
        return p[len('/workspace/'):]
    return p


def _hex(value, fallback):
    s = str(value or '').strip().lstrip('#')
    return s.upper() if HEX_RE.match(s) else fallback


def _font(name, fallback='Calibri'):
    n = str(name or '').strip()
    if not n:
        return fallback
    return FONT_FALLBACKS.get(n.lower(), n)


def _read_json(candidates):
    for candidate in candidates:
        path = _ws(candidate)
        if path and os.path.isfile(path):
            try:
                with open(path, 'r', encoding='utf-8') as fh:
                    return json.load(fh)
            except Exception:
                continue
    return None


def load_theme(theme=None):
    """dict | path | None → complete theme dict (missing tokens from aurora)."""
    if isinstance(theme, dict):
        raw = theme
    elif isinstance(theme, str) and theme and not theme.lower().endswith('.json'):
        # A theme id: an alternate by name («consulting», «boardroom»…).
        raw = next((t for t in _alternate_sources() if str(t.get('id')) == theme), None) or {}
    else:
        raw = _read_json(([theme] if theme else []) + ['tmp/sira_theme.json', '/workspace/tmp/sira_theme.json'])
    raw = raw if isinstance(raw, dict) else {}
    palette = dict(DEFAULT_THEME['palette'])
    for key, value in (raw.get('palette') or {}).items():
        palette[key] = _hex(value, palette.get(key, '000000'))
    fonts = raw.get('fonts') or {}
    charts = [_hex(c, None) for c in (raw.get('chartColors') or [])]
    charts = [c for c in charts if c] or list(DEFAULT_THEME['chartColors'])
    theme_id = str(raw.get('id') or DEFAULT_THEME['id'])
    return {
        'id': theme_id,
        'fonts': {
            'display': _font(fonts.get('display'), 'Calibri'),
            'body': _font(fonts.get('body'), 'Calibri'),
        },
        'palette': palette,
        'chartColors': charts,
        'coverStyle': raw.get('coverStyle') or ('dark' if _is_dark(palette['coverBg']) else 'light'),
        # A color the user asked for is the background of EVERY slide.
        'colorLocked': bool(raw.get('colorLocked')) or theme_id.startswith('user-color:'),
        # Chosen by the user's words: never rotated away on a repeat.
        'pinned': bool(raw.get('pinned')),
    }


def _alternate_sources():
    raw = _read_json(['tmp/sira_theme_alternates.json', '/workspace/tmp/sira_theme_alternates.json'])
    items = [t for t in (raw if isinstance(raw, list) else []) if isinstance(t, dict) and t.get('id')]
    known = {str(t.get('id')) for t in items}
    return items + [t for t in BUILTIN_ALTERNATES if t['id'] not in known]


def alternate_theme(current_id=None):
    """Another professional theme than `current_id` (for a repeated redesign)."""
    for raw in _alternate_sources():
        if str(raw.get('id')) != str(current_id or ''):
            return load_theme(raw)
    return load_theme(None)


# WCAG 2.x contrast: text colors are chosen by contrast ratio, never by a
# brightness threshold (coral used to get near-white text at 2.4:1).
def _lin(channel):
    c = channel / 255.0
    return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4


def _lum(hex_value):
    try:
        n = int(str(hex_value), 16)
    except Exception:
        return 1.0
    r, g, b = (n >> 16) & 255, (n >> 8) & 255, n & 255
    return 0.2126 * _lin(r) + 0.7152 * _lin(g) + 0.0722 * _lin(b)


def _contrast(a, b):
    la, lb = _lum(a), _lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def _is_dark(hex_value):
    """Light text reads better than dark text on this fill."""
    return _contrast(hex_value, 'F8FAFC') > _contrast(hex_value, '111827')


def _on(fill_hex):
    """Readable text color on a filled chip / header cell."""
    return max(('FFFFFF', '111827', '000000'), key=lambda c: _contrast(c, fill_hex))


def _text_for(fill_hex, preferred, minimum=4.5):
    """`preferred` when it reads on `fill_hex`, else the best ink."""
    return preferred if _contrast(preferred, fill_hex) >= minimum else _on(fill_hex)


def next_version_name(filename):
    """deck.pptx → deck-v2.pptx; deck-v2.pptx → deck-v3.pptx; x_editado.docx → x-v2.docx."""
    base = os.path.basename(str(filename or 'documento'))
    stem, ext = os.path.splitext(base)
    m = re.match(r'^(.*?)[-_ ]v(\d{1,3})$', stem, re.I)
    if m and m.group(1):
        return '%s-v%d%s' % (m.group(1), int(m.group(2)) + 1, ext)
    clean = re.sub(r'(?:[-_ ](?:editado|edited|corregido|actualizado|rediseno|redisenado))+$', '', stem, flags=re.I) or stem
    return '%s-v2%s' % (clean, ext)


def _default_dst(src):
    out_dir = 'outputs'
    if os.path.isdir('/workspace/outputs'):
        out_dir = '/workspace/outputs'
    elif not os.path.isdir(out_dir):
        out_dir = os.path.dirname(os.path.abspath(src))
    return os.path.join(out_dir, next_version_name(src))


# ── entry point ───────────────────────────────────────────────────────────

def restyle(src, dst=None, theme=None):
    """Restyle src (pptx/docx/xlsx) into dst. Returns a JSON-able report.

    When `theme` is not given and the source is already a SiraGPT redesign
    with the chosen theme («más diseño» twice), another theme is used so the
    new version is visibly different (report: theme_rotated_from). A color
    the user asked for is never rotated away.
    """
    src_path = _ws(src)
    if not os.path.isfile(src_path):
        return {'ok': False, 'error': 'source not found: %s' % src}
    ext = os.path.splitext(src_path)[1].lower().lstrip('.')
    if ext not in ('pptx', 'docx', 'xlsx'):
        return {
            'ok': False,
            'error': 'unsupported format: .%s (pptx, docx, xlsx). Restyle it with your own python '
                     'on a copy saved as outputs/%s' % (ext, next_version_name(src_path)),
        }
    dst_path = _ws(dst) if dst else _default_dst(src_path)
    if os.path.splitext(dst_path)[1].lower().lstrip('.') != ext:
        return {'ok': False, 'error': 'the output must keep the .%s format' % ext}
    if os.path.abspath(dst_path) == os.path.abspath(src_path):
        return {'ok': False, 'error': 'write a new version, never overwrite the source'}
    parent = os.path.dirname(os.path.abspath(dst_path))
    if parent and not os.path.isdir(parent):
        os.makedirs(parent, exist_ok=True)
    t = load_theme(theme)
    rotated_from = None
    if theme is None and not t.get('colorLocked') and not t.get('pinned'):
        try:
            if _styled_with(src_path, ext, t):
                rotated_from = t['id']
                t = alternate_theme(t['id'])
        except Exception:
            rotated_from = None
    if ext == 'pptx':
        report = restyle_pptx(src_path, dst_path, t)
    elif ext == 'docx':
        report = restyle_docx(src_path, dst_path, t)
    else:
        report = restyle_xlsx(src_path, dst_path, t)
    report.update({'source': src, 'output': dst_path, 'theme': t['id'], 'format': ext})
    if rotated_from:
        report['theme_rotated_from'] = rotated_from
    return report


def _styled_with(src, ext, theme):
    """True when src is already a SiraGPT redesign with this theme."""
    accent = theme['palette']['accent']
    if ext == 'pptx':
        from pptx import Presentation
        prs = Presentation(src)
        for slide in prs.slides:
            for shape in slide.shapes:
                m = DECO_THEME_RE.match(str(shape.name))
                if m:
                    return m.group(1) == theme['id']
        return False
    if ext == 'docx':
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(src)
        try:
            ppr = doc.styles['Title'].element.pPr
            bottom = ppr.find(qn('w:pBdr')).find(qn('w:bottom')) if ppr is not None else None
            return bottom is not None and str(bottom.get(qn('w:color')) or '').upper() == accent
        except Exception:
            return False
    from openpyxl import load_workbook
    wb = load_workbook(src)
    for ws in wb.worksheets[:3]:
        for row in ws.iter_rows(min_row=ws.min_row, max_row=min(ws.max_row, ws.min_row + 10)):
            for cell in row:
                rgb = str(getattr(cell.fill.fgColor, 'rgb', '') or '') if cell.fill is not None and cell.fill.fill_type else ''
                if rgb.upper().endswith(accent) and cell.font is not None and cell.font.bold:
                    return True
    return False


# ── PPTX ──────────────────────────────────────────────────────────────────

def restyle_pptx(src, dst, theme):
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(src)
    W, H = int(prs.slide_width), int(prs.slide_height)
    pal, fonts = theme['palette'], theme['fonts']
    locked = bool(theme.get('colorLocked'))
    deco = '%s[%s]' % (DECO_PREFIX, theme['id'])
    slides = list(prs.slides)
    total = len(slides)
    titles, changes, warnings = [], [], []
    for index, slide in enumerate(slides):
        try:
            _remove_named(slide, DECO_PREFIX)
            removed = _remove_background_rects(slide, W, H, warnings, index + 1)
            title = _find_title(slide, H, 0.7 if index == 0 else 0.3)
            title_text = title.text_frame.text.strip() if title is not None else ''
            titles.append(title_text)
            body_chars = sum(len(s.text_frame.text.strip()) for s in _text_shapes(slide) if not _same(s, title))
            role = 'content'
            if index == 0:
                role = 'cover'
            elif CLOSING_RE.match(title_text or '') or (title is not None and body_chars == 0 and not _has_media(slide)):
                role = 'section'
            if locked:
                # The requested color is the background of EVERY slide;
                # dividers are marked with an accent band instead.
                bg = pal['bg']
                dark = _is_dark(bg)
                ink = _text_for(bg, pal['ink'])
                body = _text_for(bg, pal['body'])
                muted = _text_for(bg, pal['muted'])
            else:
                dark = role in ('cover', 'section') and (role == 'section' or theme.get('coverStyle') == 'dark')
                bg = pal['sectionBg'] if role == 'section' else (pal['coverBg'] if role == 'cover' else pal['bg'])
                if role == 'cover' and dark:
                    bg = pal['sectionBg']
                ink = pal['sectionInk'] if dark else (pal['coverInk'] if role == 'cover' else pal['ink'])
                body = pal['sectionMuted'] if dark else (pal['coverMuted'] if role == 'cover' else pal['body'])
                ink = _text_for(bg, ink)
                body = _text_for(bg, body)
                muted = _text_for(bg, pal['sectionMuted'] if dark else pal['muted'])
            _set_background(slide, bg)
            # Decorations always go to the BACK so they can never hide text.
            if role == 'cover':
                _add_back_rect(slide, 0, 0, Emu(int(W * 0.028)), H, pal['accent'], deco, 'Cover bar')
                _add_back_rect(slide, Emu(int(W * 0.028)), 0, Emu(int(W * 0.006)), H, pal['accent2'], deco, 'Cover bar 2')
            else:
                _add_back_rect(slide, 0, 0, Emu(int(W * 0.009)), H, pal['accent'], deco, 'Side bar')
            if locked and role == 'section':
                _add_back_rect(slide, 0, int(H - H * 0.03), W, int(H * 0.03), pal['accent'], deco, 'Section band')
            if title is not None:
                # Placeholders inherit their size from the layout (a 44 pt
                # title stays 44 pt); only free text boxes get a minimum.
                min_size = None if _is_placeholder(title) else (36 if role == 'cover' else 26)
                _style_text(title.text_frame, fonts['display'], ink, bold=True, min_size=min_size)
                if role == 'section' and body_chars == 0 and title.top is not None and title.top < H * 0.3:
                    # Title-only divider / closing slide: centre it vertically.
                    title.top = int((H - int(title.height or 0)) / 2)
                rule_color = pal['accent'] if not dark else pal['accent2']
                if role == 'cover':
                    # Above the title: a wrapped cover title never collides with it.
                    rule_y = int(title.top) - Emu(137160)
                    if rule_y > 0:
                        rule_x = int(title.left) + Emu(91440)
                        if _centered(title):
                            rule_x = int(title.left + (int(title.width or 0) - int(W * 0.12)) / 2)
                        _add_back_rect(slide, rule_x, rule_y, Emu(int(W * 0.12)), Emu(64008), rule_color, deco, 'Title rule')
                else:
                    underline_y = int(title.top + _text_height(title) + Emu(45720))
                    rule_w = int(W * 0.085)
                    rule_x = int(title.left) + Emu(91440)
                    if _centered(title):
                        rule_x = int(title.left + (int(title.width or 0) - rule_w) / 2)
                    if underline_y < H * 0.8:
                        _add_back_rect(slide, rule_x, underline_y, Emu(rule_w), Emu(54864), rule_color, deco, 'Title rule')
            cards = 0
            if role == 'content':
                cards = _cards_from_bullets(slide, title, theme, W, H)
            for shape in _text_shapes(slide):
                if _same(shape, title) or str(shape.name).startswith(KEEP_PREFIXES):
                    continue
                _style_text(shape.text_frame, fonts['body'], body)
                _color_bullets(shape.text_frame, pal['accent'] if not dark else pal['accent2'])
            _restyle_cards(slide, theme)
            tables = _restyle_tables(slide, theme)
            charts = _restyle_charts(slide, theme, ink, muted, pal['line'])
            if role != 'cover' and total > 1:
                _add_page_number(slide, index + 1, total, W, H, muted, fonts['body'], deco)
            changes.append({
                'slide': index + 1, 'role': role, 'cards': cards, 'tables': tables,
                'charts': charts, 'background_rects_removed': removed,
            })
        except Exception as exc:  # best-effort: the slide stays as it was
            warnings.append('slide %d: %s' % (index + 1, exc))
    prs.save(dst)
    check = Presentation(dst)
    after_titles = []
    for c_index, slide in enumerate(check.slides):
        t = _find_title(slide, int(check.slide_height), 0.7 if c_index == 0 else 0.3)
        after_titles.append(t.text_frame.text.strip() if t is not None else '')
    ok = len(check.slides) == total
    return {
        'ok': ok, 'slides': total, 'slides_after': len(check.slides), 'titles': titles,
        'titles_after': after_titles, 'changes': changes, 'warnings': warnings,
    }


def _rgb(hex_value):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_value)


def _remove_named(slide, prefix):
    for shape in list(slide.shapes):
        if str(shape.name).startswith(prefix):
            el = shape._element
            el.getparent().remove(el)


def _remove_background_rects(slide, W, H, warnings=None, number=0):
    """Full-slide OPAQUE solid rectangles without text only paint the old
    background (flat pptxgenjs-style decks): they go. Picture-filled,
    gradient, pattern or semi-transparent full-slide shapes are CONTENT (a
    photo background, an overlay that keeps text readable): they stay."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.ns import qn
    removed = 0
    for shape in list(slide.shapes):
        try:
            if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                continue
            if shape.left is None or shape.width is None:
                continue
            if not (shape.left <= W * 0.02 and shape.top <= H * 0.02 and shape.width >= W * 0.96 and shape.height >= H * 0.96):
                continue
            sp_pr = shape._element.find(qn('p:spPr'))
            geom = sp_pr.find(qn('a:prstGeom')) if sp_pr is not None else None
            solid = sp_pr.find(qn('a:solidFill')) if sp_pr is not None else None
            other_fill = sp_pr is not None and any(
                sp_pr.find(qn(tag)) is not None for tag in ('a:blipFill', 'a:gradFill', 'a:pattFill', 'a:grpFill'))
            transparent = solid is not None and solid.find('.//' + qn('a:alpha')) is not None
            if geom is None or geom.get('prst') != 'rect' or solid is None or other_fill or transparent:
                if warnings is not None:
                    warnings.append('slide %d: full-slide shape «%s» kept (image / gradient / transparency)' % (number, shape.name))
                continue
            el = shape._element
            el.getparent().remove(el)
            removed += 1
        except Exception:
            continue
    return removed


def _text_shapes(slide):
    out = []
    for shape in slide.shapes:
        if getattr(shape, 'has_text_frame', False) and shape.has_text_frame and shape.text_frame.text.strip():
            out.append(shape)
    return out


def _has_media(slide):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in slide.shapes:
        try:
            if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.TABLE):
                return True
            if getattr(shape, 'has_table', False) and shape.has_table:
                return True
            if getattr(shape, 'has_chart', False) and shape.has_chart:
                return True
        except Exception:
            continue
    return False


def _max_font_pt(shape):
    sizes = []
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                sizes.append(r.font.size.pt)
    return max(sizes) if sizes else 0


def _same(a, b):
    """python-pptx returns a new proxy per access: compare the XML elements."""
    return a is not None and b is not None and a._element is b._element


def _find_title(slide, H, limit=0.3):
    try:
        title = slide.shapes.title
    except Exception:
        title = None
    if title is not None and title.has_text_frame and title.text_frame.text.strip():
        return title
    # Figures («80%», «$2,4 M») are KPI values, never titles; footer text
    # (deck name, «NN / TT») lives in the bottom band.
    usable = [
        s for s in _text_shapes(slide)
        if not str(s.name).startswith((DECO_PREFIX,) + KEEP_PREFIXES)
        and not FIGURE_ONLY_RE.match(s.text_frame.text.strip())
        and not (s.top is not None and s.top >= H * 0.85)
    ]
    candidates = [s for s in usable if s.top is not None and s.top < H * limit]
    if not candidates:
        # A title-only slide whose title sits lower (closing slide, or
        # centred by an earlier pass).
        candidates = usable if len(usable) == 1 else []
    if not candidates:
        return None
    candidates.sort(key=lambda s: (-_effective_size_pt(s), s.top))
    return candidates[0]


def _is_placeholder(shape):
    try:
        return bool(shape.is_placeholder)
    except Exception:
        return False


def _inherited_sources(shape):
    """Layout + master placeholders (and the master text style) a placeholder
    inherits its alignment and size from."""
    from pptx.oxml.ns import qn
    from pptx.enum.shapes import PP_PLACEHOLDER
    sources = []
    if not _is_placeholder(shape):
        return sources, None
    try:
        fmt = shape.placeholder_format
        is_title = fmt.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE)
        layout = shape.part.slide.slide_layout
        for ph in layout.placeholders:
            if ph.placeholder_format.idx == fmt.idx:
                sources.append(ph)
                break
        master = layout.slide_master
        master_type = PP_PLACEHOLDER.TITLE if is_title else fmt.type
        for ph in master.placeholders:
            if ph.placeholder_format.type == master_type:
                sources.append(ph)
                break
        style_tag = 'p:titleStyle' if is_title else 'p:bodyStyle'
        tx_styles = master._element.find(qn('p:txStyles'))
        style = tx_styles.find(qn(style_tag)) if tx_styles is not None else None
        return sources, style
    except Exception:
        return sources, None


def _lvl1(el):
    """(algn, sz) of a txBody's lstStyle lvl1 / first paragraph."""
    from pptx.oxml.ns import qn
    algn, sz = None, None
    if el is None:
        return algn, sz
    body = el if el.tag == qn('p:txBody') else el.find('.//' + qn('p:txBody'))
    if body is None:
        return algn, sz
    lst = body.find(qn('a:lstStyle'))
    lvl1 = lst.find(qn('a:lvl1pPr')) if lst is not None else None
    if lvl1 is not None:
        algn = lvl1.get('algn')
        d = lvl1.find(qn('a:defRPr'))
        if d is not None and d.get('sz'):
            sz = int(d.get('sz')) / 100.0
    para = body.find(qn('a:p'))
    if para is not None:
        ppr = para.find(qn('a:pPr'))
        if algn is None and ppr is not None:
            algn = ppr.get('algn')
        if sz is None:
            for rpr in para.iter(qn('a:rPr'), qn('a:endParaRPr')):
                if rpr.get('sz'):
                    sz = int(rpr.get('sz')) / 100.0
                    break
    return algn, sz


def _inherited(shape):
    from pptx.oxml.ns import qn
    algn, sz = None, None
    sources, style = _inherited_sources(shape)
    for src in sources:
        a, z = _lvl1(src._element)
        algn = algn or a
        sz = sz or z
    if style is not None:
        lvl1 = style.find(qn('a:lvl1pPr'))
        if lvl1 is not None:
            algn = algn or lvl1.get('algn')
            d = lvl1.find(qn('a:defRPr'))
            if sz is None and d is not None and d.get('sz'):
                sz = int(d.get('sz')) / 100.0
    return algn, sz


def _effective_size_pt(shape):
    size = _max_font_pt(shape)
    if size:
        return size
    _, inherited = _inherited(shape)
    return inherited or 18


def _centered(shape):
    """Explicit paragraph alignment first, else the layout/master's."""
    try:
        from pptx.enum.text import PP_ALIGN
        explicit = [p.alignment for p in shape.text_frame.paragraphs if p.text.strip() and p.alignment is not None]
        if explicit:
            return any(a == PP_ALIGN.CENTER for a in explicit)
        algn, _ = _inherited(shape)
        return algn == 'ctr'
    except Exception:
        return False


def _estimated_lines(text, size_pt, width_emu):
    """Wrapped line count for `text` at `size_pt` in a box `width_emu` wide."""
    width_in = max(float(width_emu) / 914400.0 - 0.2, 0.5)
    chars_per_line = max(8, int(width_in * 72.0 / (size_pt * 0.52)))
    lines = 0
    for part in str(text or '').split('\n'):
        lines += max(1, -(-len(part) // chars_per_line))
    return max(lines, 1)


def _text_height(shape):
    """Rendered height of a text box: its frame, or more when the text wraps past it."""
    size = _effective_size_pt(shape)
    lines = _estimated_lines(shape.text_frame.text, size, shape.width or 0)
    needed = int(lines * size * 1.25 / 72.0 * 914400) + 91440
    return max(int(shape.height or 0), needed) if lines > 1 else int(shape.height or needed)


def _set_background(slide, hex_value):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(hex_value)


def _to_back(slide, shape):
    el = shape._element
    tree = el.getparent()
    tree.remove(el)
    tree.insert(2, el)  # after nvGrpSpPr + grpSpPr: behind every existing shape


def _add_back_rect(slide, x, y, w, h, hex_value, deco, label):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(h))
    shape.name = '%s %s' % (deco, label)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hex_value)
    shape.line.fill.background()
    shape.shadow.inherit = False
    _to_back(slide, shape)
    return shape


def _style_text(text_frame, font_name, hex_value, bold=None, min_size=None):
    from pptx.util import Pt
    for p in text_frame.paragraphs:
        for r in p.runs:
            r.font.name = font_name
            r.font.color.rgb = _rgb(hex_value)
            if bold is not None:
                r.font.bold = bold
            if min_size and r.font.size is None:
                r.font.size = Pt(min_size)


BULLET_TAGS = ('buSzTx', 'buSzPct', 'buSzPts', 'buFontTx', 'buFont', 'buNone', 'buAutoNum', 'buChar', 'buBlip')


def _color_bullets(text_frame, hex_value):
    from pptx.oxml.ns import qn
    from lxml import etree
    for p in text_frame.paragraphs:
        pPr = p._p.find(qn('a:pPr'))
        if pPr is None:
            continue
        if pPr.find(qn('a:buChar')) is None and pPr.find(qn('a:buAutoNum')) is None:
            continue
        for old in pPr.findall(qn('a:buClr')):
            pPr.remove(old)
        clr = etree.SubElement(pPr, qn('a:buClr'))
        srgb = etree.SubElement(clr, qn('a:srgbClr'))
        srgb.set('val', hex_value)
        pPr.remove(clr)
        anchor = None
        for tag in BULLET_TAGS:
            anchor = pPr.find(qn('a:%s' % tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(clr)
        else:
            pPr.append(clr)


def kpi_parts(text):
    """(value, label) of «35 % reducción…»; None for plain text."""
    m = KPI_RE.match(str(text or ''))
    if not m:
        return None
    value = m.group(1).strip()
    if YEAR_RE.match(value) or LIST_MARKER_RE.match(value) or not METRIC_RE.search(value):
        return None
    return value, m.group(2).strip()


def _cards_from_bullets(slide, title, theme, W, H):
    """One body text box with 2–6 short items → cards with numbered chips / KPI tiles."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Pt, Emu

    if _has_media(slide):
        return 0
    if any(str(s.name).startswith(KEEP_PREFIXES) for s in slide.shapes):
        return 0  # already laid out as cards (an earlier pass or SiraGPT's own builder)
    body = [s for s in _text_shapes(slide) if not _same(s, title) and not str(s.name).startswith((DECO_PREFIX,) + KEEP_PREFIXES)]
    if len(body) != 1:
        return 0
    source = body[0]
    items = []
    for p in source.text_frame.paragraphs:
        text = str(getattr(p, 'text', '') or '').replace('\v', ' ').strip()
        if not text:
            continue
        if p.level and p.level > 0:
            return 0
        items.append(text)
    if len(items) < 2 or len(items) > 6 or any(len(t) > 140 for t in items):
        return 0
    pal, fonts = theme['palette'], theme['fonts']
    kpis = [kpi_parts(t) for t in items]
    kpi_mode = sum(1 for k in kpis if k) >= max(2, len(items) - 1)

    margin = int(W * 0.055)
    top = int(title.top + _text_height(title) + Emu(320040)) if title is not None else int(H * 0.24)
    top = max(top, int(H * 0.2))
    bottom = int(H - H * 0.11)
    if bottom - top < H * 0.3:
        return 0
    n = len(items)
    cols = n if n <= 3 else (2 if n == 4 else 3)
    rows = 1 if n <= 3 else 2
    gap = int(W * 0.018)
    card_w = int((W - 2 * margin - gap * (cols - 1)) / cols)
    free_h = int((bottom - top - gap * (rows - 1)) / rows)
    pad = int(W * 0.014)
    chip = max(min(int(min(card_w, free_h) * 0.22), Emu(420000)), Emu(300000))
    longest = max(len(t) for t in items)
    size = 20 if longest <= 60 else (18 if longest <= 100 else 16)
    if cols == 3 and longest > 45 and size > 18:
        size = 18
    value_h = Emu(int(914400 * 0.85))

    def kpi_needed():
        labels = [k[1] if k else t for k, t in zip(kpis, items)]
        lines = max(_estimated_lines(label, 16, card_w - 2 * pad) for label in labels)
        return int(pad * 2 + value_h + lines * 16 * 1.3 / 72.0 * 914400 + Emu(91440))

    def card_needed(pt):
        lines = max(_estimated_lines(t, pt, card_w - 2 * pad) for t in items)
        return int(chip + pad * 2.8 + lines * pt * 1.3 / 72.0 * 914400 + Emu(182880))

    # Tiles grow with their text; text that cannot fit the free area falls
    # back to cards, then smaller type, then the slide is left as it was.
    if kpi_mode and kpi_needed() > free_h:
        kpi_mode = False
    if not kpi_mode:
        while card_needed(size) > free_h and size > 14:
            size -= 2
        if card_needed(size) > free_h:
            return 0
    needed = kpi_needed() if kpi_mode else card_needed(size)
    card_h = min(free_h, max(needed, int(H * (0.34 if rows == 1 else 0.26))))
    for i, text in enumerate(items):
        r, c = divmod(i, cols)
        x = margin + c * (card_w + gap)
        y = top + r * (card_h + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.name = 'SiraCard %d' % (i + 1)
        try:
            card.adjustments[0] = 0.08
        except Exception:
            pass
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(pal['surface'])
        card.line.color.rgb = _rgb(pal['chipLine'])
        card.line.width = Pt(0.75)
        card.shadow.inherit = False
        kpi = kpis[i] if kpi_mode else None
        if kpi_mode and not kpi:
            # The odd non-figure item of a KPI row keeps its full text.
            box = slide.shapes.add_textbox(x + pad, y + pad, card_w - 2 * pad, card_h - 2 * pad)
            box.name = 'SiraCard %d text' % (i + 1)
            btf = box.text_frame
            btf.word_wrap = True
            btf.vertical_anchor = MSO_ANCHOR.MIDDLE
            brun = btf.paragraphs[0].add_run()
            brun.text = text
            brun.font.size = Pt(16)
            brun.font.name = fonts['body']
            brun.font.color.rgb = _rgb(_text_for(pal['surface'], pal['body']))
            continue
        if kpi:
            num = slide.shapes.add_textbox(x + pad, y + pad, card_w - 2 * pad, value_h)
            num.name = 'SiraKpi %d value' % (i + 1)
            tf = num.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.BOTTOM
            run = tf.paragraphs[0].add_run()
            run.text = kpi[0]
            run.font.size = Pt(34 if len(run.text) <= 8 else 26)
            run.font.bold = True
            run.font.name = fonts['display']
            run.font.color.rgb = _rgb(_text_for(pal['surface'], pal['accent'], 3.0))
            label_top = y + pad + value_h
            label = slide.shapes.add_textbox(x + pad, label_top, card_w - 2 * pad, max(card_h - (label_top - y) - pad, Emu(300000)))
            label.name = 'SiraKpi %d label' % (i + 1)
            ltf = label.text_frame
            ltf.word_wrap = True
            lrun = ltf.paragraphs[0].add_run()
            lrun.text = kpi[1]
            lrun.font.size = Pt(16)
            lrun.font.name = fonts['body']
            lrun.font.color.rgb = _rgb(_text_for(pal['surface'], pal['body']))
            continue
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + pad, y + pad, chip, chip)
        dot.name = 'SiraChip %d' % (i + 1)
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb(pal['accent'])
        dot.line.fill.background()
        dot.shadow.inherit = False
        dtf = dot.text_frame
        dtf.margin_left = dtf.margin_right = dtf.margin_top = dtf.margin_bottom = 0
        dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
        dp = dtf.paragraphs[0]
        dp.alignment = PP_ALIGN.CENTER
        drun = dp.add_run()
        drun.text = str(i + 1)
        drun.font.size = Pt(13)
        drun.font.bold = True
        drun.font.name = fonts['display']
        drun.font.color.rgb = _rgb(_on(pal['accent']))
        text_top = y + pad + chip + int(pad * 0.6)
        box = slide.shapes.add_textbox(x + pad, text_top, card_w - 2 * pad, max(card_h - (text_top - y) - pad, Emu(300000)))
        box.name = 'SiraCard %d text' % (i + 1)
        btf = box.text_frame
        btf.word_wrap = True
        brun = btf.paragraphs[0].add_run()
        brun.text = text
        brun.font.size = Pt(size)
        brun.font.name = fonts['body']
        brun.font.color.rgb = _rgb(_text_for(pal['surface'], pal['body']))
    el = source._element
    el.getparent().remove(el)
    return n


def _restyle_cards(slide, theme):
    """A second pass (v2 → v3) recolors the cards it made before."""
    pal = theme['palette']
    body = _text_for(pal['surface'], pal['body'])
    for shape in slide.shapes:
        name = str(shape.name)
        try:
            if name.startswith('SiraCard') and not name.endswith('text'):
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb(pal['surface'])
                shape.line.color.rgb = _rgb(pal['chipLine'])
            elif name.startswith('SiraChip'):
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb(pal['accent'])
                _style_text(shape.text_frame, theme['fonts']['display'], _on(pal['accent']), bold=True)
            elif name.startswith('SiraCard') and name.endswith('text'):
                _style_text(shape.text_frame, theme['fonts']['body'], body)
            elif name.startswith('SiraKpi') and name.endswith('value'):
                _style_text(shape.text_frame, theme['fonts']['display'], _text_for(pal['surface'], pal['accent'], 3.0), bold=True)
            elif name.startswith('SiraKpi'):
                _style_text(shape.text_frame, theme['fonts']['body'], body)
        except Exception:
            continue


def _set_cell_fill(cell, hex_value):
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(hex_value)


def _restyle_tables(slide, theme):
    pal, fonts = theme['palette'], theme['fonts']
    count = 0
    for shape in slide.shapes:
        if not (getattr(shape, 'has_table', False) and shape.has_table):
            continue
        table = shape.table
        for r_idx, row in enumerate(table.rows):
            for cell in row.cells:
                try:
                    if r_idx == 0:
                        _set_cell_fill(cell, pal['accent'])
                        _style_text(cell.text_frame, fonts['display'], _on(pal['accent']), bold=True)
                    else:
                        fill = pal['surfaceAlt'] if r_idx % 2 == 0 else pal['surface']
                        _set_cell_fill(cell, fill)
                        _style_text(cell.text_frame, fonts['body'], _text_for(fill, pal['ink']))
                except Exception:
                    continue
        count += 1
    return count


def _font_color(font, hex_value):
    try:
        font.color.rgb = _rgb(hex_value)
    except Exception:
        pass


def _restyle_charts(slide, theme, ink, muted, line):
    """Series take the theme's chart colors; every chart TEXT (title, axis
    labels, legend) and the gridlines take the slide's ink / line so a chart
    stays readable on a dark background."""
    colors = theme['chartColors']
    count = 0
    for shape in slide.shapes:
        if not (getattr(shape, 'has_chart', False) and shape.has_chart):
            continue
        try:
            chart = shape.chart
            for plot in chart.plots:
                pie_like = plot.__class__.__name__ in ('PiePlot', 'DoughnutPlot')
                for s_idx, series in enumerate(plot.series):
                    if pie_like:
                        for p_idx in range(len(list(series.values))):
                            point = series.points[p_idx]
                            point.format.fill.solid()
                            point.format.fill.fore_color.rgb = _rgb(colors[p_idx % len(colors)])
                    else:
                        series.format.fill.solid()
                        series.format.fill.fore_color.rgb = _rgb(colors[s_idx % len(colors)])
                        try:
                            series.format.line.color.rgb = _rgb(colors[s_idx % len(colors)])
                        except Exception:
                            pass
            _font_color(chart.font, ink)
            try:
                if chart.has_title:
                    for p in chart.chart_title.text_frame.paragraphs:
                        for r in p.runs:
                            _font_color(r.font, ink)
            except Exception:
                pass
            for axis_name in ('category_axis', 'value_axis'):
                try:
                    axis = getattr(chart, axis_name)
                except Exception:
                    continue  # pie / doughnut: no axes
                try:
                    _font_color(axis.tick_labels.font, muted)
                    axis.format.line.color.rgb = _rgb(line)
                    if axis_name == 'value_axis' and axis.has_major_gridlines:
                        axis.major_gridlines.format.line.color.rgb = _rgb(line)
                except Exception:
                    pass
            try:
                if chart.has_legend:
                    _font_color(chart.legend.font, ink)
            except Exception:
                pass
            count += 1
        except Exception:
            continue
    return count


def _add_page_number(slide, number, total, W, H, hex_value, font_name, deco):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    w, h = int(W * 0.12), int(H * 0.05)
    box = slide.shapes.add_textbox(int(W - w - W * 0.03), int(H - h - H * 0.02), w, h)
    box.name = '%s Page number' % deco
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = '%02d / %02d' % (number, total)
    run.font.size = Pt(10)
    run.font.name = font_name
    run.font.color.rgb = _rgb(hex_value)


# ── DOCX ──────────────────────────────────────────────────────────────────

def _doc_colors(theme):
    """Documents are printed on white paper: dark themes keep only their accent."""
    pal = theme['palette']
    if _is_dark(pal['bg']):
        return {'ink': '111827', 'body': '374151', 'muted': '6B7280', 'accent': pal['accent'],
                'band': 'F1F5F9', 'line': 'CBD5E1'}
    band = pal['surfaceAlt'] if not _is_dark(pal['surfaceAlt']) else 'F1F5F9'
    line = pal['line'] if not _is_dark(pal['line']) else 'CBD5E1'
    return {'ink': pal['ink'], 'body': pal['body'], 'muted': pal['muted'], 'accent': pal['accent'],
            'band': band, 'line': line}


# Fonts whose glyphs ARE content: symbol fonts map private-use code points,
# equations use a math font, code uses a monospace one. Never replaced.
KEEP_FONT_RE = re.compile(
    r'^(symbol|wingdings.*|webdings|marlett|.*\bmath\b.*|cambria math|consolas|courier.*|menlo|monaco|'
    r'lucida console|lucida sans typewriter|source code.*|fira code|fira mono|jetbrains mono|dejavu sans mono|'
    r'liberation mono|andale mono|inconsolata|roboto mono|sf mono|ubuntu mono|noto sans mono.*|.*\bmono\b.*)$',
    re.I,
)
ACADEMIC_TEXT_RE = re.compile(
    r'\b(tesis|tesina|monograf[ií]a|marco te[oó]rico|referencias bibliogr[aá]ficas|bibliograf[ií]a|abstract|'
    r'palabras clave|estado del arte|trabajo de (?:investigaci[oó]n|grado)|hip[oó]tesis)\b',
    re.I,
)


def _keeps_font(run):
    """True for runs whose font is part of the content (symbols, math, code)."""
    from docx.oxml.ns import qn
    try:
        el = run._r
        if el.find(qn('w:sym')) is not None:
            return True
        names = [run.font.name or '']
        rpr = el.rPr
        rfonts = rpr.find(qn('w:rFonts')) if rpr is not None else None
        if rfonts is not None:
            names += [rfonts.get(qn(a)) or '' for a in ('w:ascii', 'w:hAnsi', 'w:cs')]
        return any(n and KEEP_FONT_RE.match(n.strip()) for n in names)
    except Exception:
        return False


def _is_academic_docx(doc):
    """Theses / papers follow institutional norms (fonts, spacing, black
    headings): double spacing or 2+ academic section markers."""
    try:
        spacing = doc.styles['Normal'].paragraph_format.line_spacing
        if isinstance(spacing, float) and spacing >= 1.5:
            return True
    except Exception:
        pass
    hits = set()
    for p in doc.paragraphs[:600]:
        for m in ACADEMIC_TEXT_RE.finditer(p.text[:300]):
            hits.add(m.group(1).lower())
        if len(hits) >= 2:
            return True
    return False


def restyle_docx(src, dst, theme):
    from docx import Document
    from docx.shared import Pt

    doc = Document(src)
    fonts = theme['fonts']
    colors = _doc_colors(theme)
    academic = _is_academic_docx(doc)
    before_paragraphs = len(doc.paragraphs)
    before_tables = len(doc.tables)
    changes, warnings = [], []
    if academic:
        # Academic profile: fonts, sizes and line spacing stay (they follow the
        # institution's norm); headings in black, no colored rule; only table
        # polish and page numbers.
        colors = dict(colors, accent='000000', ink='000000', band='F2F2F2', line='808080')
        for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
            try:
                doc.styles[name].font.color.rgb = _docx_rgb('000000')
                changes.append('style:%s' % name)
            except KeyError:
                continue
    else:
        spec = [
            ('Normal', fonts['body'], None, colors['body'], None),
            ('Title', fonts['display'], 26, colors['ink'], True),
            ('Subtitle', fonts['body'], 14, colors['muted'], None),
            ('Heading 1', fonts['display'], 16, colors['accent'], True),
            ('Heading 2', fonts['display'], 13, colors['ink'], True),
            ('Heading 3', fonts['display'], 12, colors['body'], True),
            ('List Paragraph', fonts['body'], None, colors['body'], None),
        ]
        for name, font, size, color, bold in spec:
            if _docx_style(doc, name, font, size, color, bold):
                changes.append('style:%s' % name)
        try:
            normal = doc.styles['Normal'].paragraph_format
            # An explicit line spacing is the author's choice (and the page
            # count depends on it): only an unset one gets 1.15.
            if normal.line_spacing is None:
                normal.line_spacing = 1.15
            if normal.space_after is None:
                normal.space_after = Pt(6)
            for name, before, after in (('Heading 1', 18, 6), ('Heading 2', 12, 4), ('Heading 3', 10, 4)):
                try:
                    pf = doc.styles[name].paragraph_format
                    pf.space_before = Pt(before)
                    pf.space_after = Pt(after)
                    pf.keep_with_next = True
                except KeyError:
                    continue
            _docx_bottom_rule(doc.styles['Title'].element, colors['accent'])
        except Exception as exc:
            warnings.append('paragraph styles: %s' % exc)
    headings = []
    for p in doc.paragraphs:
        style = (p.style.name if p.style is not None else '') or ''
        text = p.text.strip()
        try:
            if style == 'Title' or style.startswith('Heading'):
                if text:
                    headings.append(text)
                color = colors['accent'] if style == 'Heading 1' else colors['ink']
                for r in p.runs:
                    if not academic and not _keeps_font(r):
                        r.font.name = fonts['display']
                        _docx_run_fonts(r, fonts['display'])
                    if r.font.color is not None and r.font.color.rgb is not None:
                        r.font.color.rgb = _docx_rgb(color)
            elif not academic:
                for r in p.runs:
                    if r.font.name and r.font.name != fonts['body'] and not _keeps_font(r):
                        r.font.name = fonts['body']
                        _docx_run_fonts(r, fonts['body'])
        except Exception as exc:
            warnings.append('paragraph: %s' % exc)
    tables = 0
    for table in doc.tables:
        try:
            _docx_table(table, fonts, colors, keep_fonts=academic)
            tables += 1
        except Exception as exc:
            warnings.append('table: %s' % exc)
    footers = 0
    for s_index, section in enumerate(doc.sections):
        try:
            if _docx_page_numbers(section, s_index, fonts['body'] if not academic else None, colors['muted']):
                footers += 1
        except Exception as exc:
            warnings.append('footer: %s' % exc)
    doc.save(dst)
    check = Document(dst)
    ok = len(check.paragraphs) >= before_paragraphs and len(check.tables) == before_tables
    return {
        'ok': ok, 'profile': 'academic' if academic else 'professional',
        'paragraphs': before_paragraphs, 'paragraphs_after': len(check.paragraphs),
        'tables': before_tables, 'tables_styled': tables, 'footers_numbered': footers,
        'titles': headings[:60], 'changes': changes, 'warnings': warnings,
        # New fonts / spacing legitimately move page breaks: verify with
        # expect.same_page_count=false for documents.
        'page_count_may_change': True,
    }


def _docx_rgb(hex_value):
    from docx.shared import RGBColor
    return RGBColor.from_string(hex_value)


def _docx_run_fonts(run, font_name):
    from docx.oxml.ns import qn
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(attr), font_name)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme'):
        key = qn(attr)
        if key in rfonts.attrib:
            del rfonts.attrib[key]


def _docx_style(doc, name, font, size, color, bold):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    try:
        style = doc.styles[name]
    except KeyError:
        return False
    style.font.name = font
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), font)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        key = qn(attr)
        if key in rfonts.attrib:
            del rfonts.attrib[key]
    if size:
        style.font.size = Pt(size)
    if color:
        style.font.color.rgb = _docx_rgb(color)
    if bold is not None:
        style.font.bold = bold
    return True


def _docx_bottom_rule(style_el, hex_value):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    ppr = style_el.get_or_add_pPr()
    for old in ppr.findall(qn('w:pBdr')):
        ppr.remove(old)
    bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '6')
    bottom.set(qn('w:color'), hex_value)
    bdr.append(bottom)
    # pBdr sits after keepNext/keepLines/pageBreakBefore/framePr/widowControl/numPr/suppressLineNumbers.
    anchor = None
    for tag in ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
                'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
                'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
                'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId',
                'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange'):
        anchor = ppr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(bdr)
    else:
        ppr.append(bdr)


def _docx_cell_shading(cell, hex_value):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn('w:shd')):
        tcpr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_value)
    # shd follows tcW/gridSpan/vMerge/tcBorders in CT_TcPr.
    anchor = None
    for tag in ('w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark'):
        anchor = tcpr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(shd)
    else:
        tcpr.append(shd)


def _docx_table_borders(table, hex_value):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn('w:tblBorders')):
        tblpr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:%s' % edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), hex_value)
        borders.append(el)
    anchor = None
    for tag in ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription'):
        anchor = tblpr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        tblpr.append(borders)


def _docx_table(table, fonts, colors, keep_fonts=False):
    _docx_table_borders(table, colors['line'])
    header_fill = colors['accent'] if not keep_fonts else colors['band']
    header_fg = _on(header_fill)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if r_idx == 0:
                _docx_cell_shading(cell, header_fill)
            elif r_idx % 2 == 0 and not keep_fonts:
                _docx_cell_shading(cell, colors['band'])
            for p in cell.paragraphs:
                for r in p.runs:
                    if not keep_fonts and not _keeps_font(r):
                        r.font.name = fonts['display'] if r_idx == 0 else fonts['body']
                        _docx_run_fonts(r, r.font.name)
                    if r_idx == 0:
                        r.font.bold = True
                        r.font.color.rgb = _docx_rgb(header_fg)


def _docx_page_numbers(section, index, font_name, hex_value):
    """«N / T» in an EMPTY footer only; an existing footer is left untouched."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    footer = section.footer
    if index > 0 and footer.is_linked_to_previous:
        return False  # inherits the first section's footer
    xml = footer._element.xml if hasattr(footer, '_element') else ''
    if any(p.text.strip() for p in footer.paragraphs) or 'fldSimple' in xml or 'instrText' in xml:
        return False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def field(instr):
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), instr)
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = '1'
        r.append(t)
        fld.append(r)
        return fld

    p._p.append(field('PAGE'))
    sep = p.add_run(' / ')
    if font_name:
        sep.font.name = font_name
    sep.font.color.rgb = _docx_rgb(hex_value)
    p._p.append(field('NUMPAGES'))
    return True


# ── XLSX ──────────────────────────────────────────────────────────────────

TOTAL_RE = re.compile(r'^\s*(sub)?total(es)?\b', re.I)
MONEY_HEADER_RE = re.compile(r'\b(total|monto|importe|valor|ventas?|costo|coste|precio|ingresos?|gastos?|presupuesto|subtotal|amount|revenue|cost|price|sales|s/|usd|eur|\$)', re.I)
ID_HEADER_RE = re.compile(r'\b(a[nñ]o|year|id|c[oó]digo|codigo|dni|ruc|tel[eé]fono|telefono|cp|zip|n[°º]|nro|numero de documento)\b', re.I)


def restyle_xlsx(src, dst, theme):
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.utils import get_column_letter
    from sira_charts import add_xlsx_chart, apply_xlsx_chart_colors

    wb = load_workbook(src)  # formulas stay formulas
    fonts = theme['fonts']
    colors = _doc_colors(theme)
    sheets, warnings, headers_out = [], [], []
    before = {ws.title: _sheet_signature(ws) for ws in wb.worksheets}
    thin = Side(style='thin', color=colors['line'])
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill('solid', fgColor=colors['accent'])
    band_fill = PatternFill('solid', fgColor=colors['band'])
    header_fg = _on(colors['accent'])
    for ws in wb.worksheets:
        info = {'sheet': ws.title, 'header_row': None, 'chart': False, 'charts_recolored': 0, 'data_bar': None, 'title_row': None}
        try:
            if ws.max_row <= 1 and ws.max_column <= 1 and ws['A1'].value is None:
                sheets.append(info)
                continue
            min_col, max_col = ws.min_column, ws.max_column
            header_row = _detect_header(ws, min_col, max_col)
            info['header_row'] = header_row
            last_row = ws.max_row
            start = (header_row + 1) if header_row else ws.min_row
            if header_row:
                for c in range(min_col, max_col + 1):
                    cell = ws.cell(row=header_row, column=c)
                    if cell.value is None:
                        continue
                    cell.fill = header_fill
                    cell.font = Font(name=fonts['display'], bold=True, color=header_fg, size=cell.font.size or 11)
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = border
                    headers_out.append(str(cell.value))
                ws.row_dimensions[header_row].height = max(ws.row_dimensions[header_row].height or 0, 24)
                if ws.freeze_panes is None:
                    ws.freeze_panes = ws.cell(row=header_row + 1, column=min_col).coordinate
                info['title_row'] = _style_title_row(ws, header_row, min_col, max_col, fonts, colors)
            label_col = _label_column(ws, min_col, max_col, start, last_row)
            temporal = re.compile(r'\b(mes|meses|month|months|a[nñ]o|year|fecha|date|trimestre|quarter|semana|week)\b', re.I)
            if header_row:
                time_cols = [c for c in range(min_col, max_col + 1)
                             if temporal.search(str(ws.cell(header_row, c).value or ''))]
                if time_cols:
                    label_col = time_cols[0]
            for existing_chart in getattr(ws, '_charts', []):
                apply_xlsx_chart_colors(existing_chart, theme['chartColors'])
                info['charts_recolored'] += 1
            total_rows = set()
            if label_col:
                for r in range(start, last_row + 1):
                    if TOTAL_RE.match(str(ws.cell(row=r, column=label_col).value or '')):
                        total_rows.add(r)
            numeric_cols, formula_cols, float_cols = {}, {}, set()
            for r in range(start, last_row + 1):
                band = (r - start) % 2 == 1
                for c in range(min_col, max_col + 1):
                    cell = ws.cell(row=r, column=c)
                    if cell.value is None:
                        continue
                    old = cell.font
                    cell.font = Font(name=fonts['body'], bold=True if r in total_rows else old.bold, italic=old.italic,
                                     size=old.size, color=old.color, underline=old.underline, strike=old.strike)
                    cell.border = border
                    if r in total_rows:
                        cell.border = Border(left=thin, right=thin, bottom=thin, top=Side(style='medium', color=colors['accent']))
                        if cell.fill is None or cell.fill.fill_type is None:
                            cell.fill = band_fill
                    elif band and (cell.fill is None or cell.fill.fill_type is None):
                        cell.fill = band_fill
                    if isinstance(cell.value, bool):
                        continue
                    if isinstance(cell.value, str) and cell.value.startswith('='):
                        formula_cols.setdefault(c, []).append(r)
                        continue
                    header_text = str(ws.cell(row=header_row, column=c).value or '') if header_row else ''
                    if isinstance(cell.value, (int, float)):
                        numeric_cols.setdefault(c, []).append(r)
                        if isinstance(cell.value, float) and not float(cell.value).is_integer():
                            float_cols.add(c)
                        if cell.number_format == 'General' and not ID_HEADER_RE.search(header_text):
                            if isinstance(cell.value, float) and not float(cell.value).is_integer():
                                cell.number_format = '#,##0.00'
                            elif abs(cell.value) >= 1000:
                                cell.number_format = '#,##0'
            for c in float_cols:  # one format per decimal column (55 → 55.00 next to 28.50)
                for r in numeric_cols.get(c, []):
                    cell = ws.cell(row=r, column=c)
                    if cell.number_format in ('General', '#,##0'):
                        cell.number_format = '#,##0.00'
            # Formula columns (totals, amounts) get the money format of the
            # figures they compute; a formula never changes, only its format.
            # Only a money header or a column that already holds decimals gets
            # 2 decimals: an integer «Unidades» total stays «735».
            for c, rows in formula_cols.items():
                header_text = str(ws.cell(row=header_row, column=c).value or '') if header_row else ''
                if ID_HEADER_RE.search(header_text):
                    continue
                if MONEY_HEADER_RE.search(header_text) or c in float_cols:
                    for r in rows:
                        cell = ws.cell(row=r, column=c)
                        if cell.number_format == 'General':
                            cell.number_format = '#,##0.00'
            _autowidth(ws, min_col, max_col, get_column_letter)
            candidates = {}
            for c in set(numeric_cols) | set(formula_cols):
                header_text = str(ws.cell(row=header_row, column=c).value or '') if header_row else ''
                if ID_HEADER_RE.search(header_text) or c == label_col:
                    continue
                rows = sorted(set(numeric_cols.get(c, [])) | set(formula_cols.get(c, [])))
                rows = [r for r in rows if r not in total_rows]
                if len(rows) >= 2:
                    candidates[c] = rows
            if candidates:
                # The main measure: a money/total header, else the right-most figure column.
                named = [c for c in candidates if MONEY_HEADER_RE.search(str(ws.cell(row=header_row, column=c).value or '') if header_row else '')]
                col = max(named) if named else max(candidates)
                rows = candidates[col]
                letter = get_column_letter(col)
                if len(rows) >= 3:
                    rng = '%s%d:%s%d' % (letter, rows[0], letter, rows[-1])
                    if not _has_data_bar(ws, rng):
                        ws.conditional_formatting.add(rng, DataBarRule(start_type='min', end_type='max', color=colors['accent']))
                    info['data_bar'] = rng
                contiguous = rows == list(range(rows[0], rows[-1] + 1))
                if header_row and label_col and contiguous and rows[0] == header_row + 1 and 2 <= len(rows) <= 40 and not getattr(ws, '_charts', None):
                    # Comparable measures share one native chart. Do not mix
                    # quantity / unit-price columns into a monetary total.
                    compatible = sorted(c for c in candidates if candidates[c] == rows)
                    money = [c for c in compatible if MONEY_HEADER_RE.search(str(ws.cell(header_row, c).value or ''))
                             and not re.search(r'unitari|unit price|precio|price', str(ws.cell(header_row, c).value or ''), re.I)]
                    chart_cols = money or compatible
                    chart_type = 'line' if temporal.search(str(ws.cell(header_row, label_col).value or '')) else 'column'
                    ranges = ['%s%d:%s%d' % (get_column_letter(c), header_row, get_column_letter(c), rows[-1]) for c in chart_cols]
                    # Right of the table AND of every existing image / chart.
                    anchor_col = max(max_col + 2, _drawings_right_col(ws) + 2)
                    add_xlsx_chart(ws, chart_type=chart_type, data_range=ranges,
                                   category_range='%s%d:%s%d' % (get_column_letter(label_col), rows[0], get_column_letter(label_col), rows[-1]),
                                   titles_from_data=True, colors=theme['chartColors'],
                                   title=ws.title if len(chart_cols) > 1 else str(ws.cell(header_row, chart_cols[0]).value or ws.title),
                                   legend='b' if len(chart_cols) > 1 else False,
                                   anchor='%s%d' % (get_column_letter(anchor_col), header_row), width=16, height=8)
                    info['chart'] = True
            ws.sheet_view.showGridLines = False
            _fit_to_page(ws, max_col - min_col + 1, info['chart'])
        except Exception as exc:
            warnings.append('%s: %s' % (ws.title, exc))
        sheets.append(info)
    wb.save(dst)
    check = load_workbook(dst)
    after = {ws.title: _sheet_signature(ws) for ws in check.worksheets}
    ok = before == after
    return {'ok': ok, 'sheets': sheets, 'titles': headers_out[:80], 'values_preserved': ok, 'warnings': warnings}


def _style_title_row(ws, header_row, min_col, max_col, fonts, colors):
    """A single text cell above the header (often merged) is the sheet title."""
    from openpyxl.styles import Font
    for r in range(ws.min_row, header_row):
        cells = [ws.cell(row=r, column=c) for c in range(min_col, max_col + 1)]
        filled = [c for c in cells if c.value is not None and str(c.value).strip()]
        if len(filled) != 1 or not isinstance(filled[0].value, str) or filled[0].value.startswith('='):
            continue
        cell = filled[0]
        old = cell.font
        cell.font = Font(name=fonts['display'], bold=True, italic=old.italic, size=max(float(old.size or 11), 14),
                         color=colors['ink'], underline=old.underline)
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 24)
        return r
    return None


def _drawings_right_col(ws):
    """Right-most column (1-based) covered by an image or chart of the sheet."""
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    right = 0
    for obj in list(getattr(ws, '_images', []) or []) + list(getattr(ws, '_charts', []) or []):
        anchor = getattr(obj, 'anchor', None)
        try:
            if isinstance(anchor, str):
                col = column_index_from_string(coordinate_from_string(anchor)[0])
                width_px = float(getattr(obj, 'width', 0) or 0)
                if not width_px and getattr(obj, 'width', None) is None:
                    width_px = 15 * 37.8  # openpyxl chart default: 15 cm
                right = max(right, col + int(width_px // 64) + 1)
                continue
            to = getattr(anchor, 'to', None)
            if to is not None:
                right = max(right, int(to.col) + 1)
                continue
            frm = getattr(anchor, '_from', None)
            ext = getattr(anchor, 'ext', None)
            if frm is not None:
                width_emu = float(getattr(ext, 'width', None) or getattr(ext, 'cx', 0) or 0)
                right = max(right, int(frm.col) + 1 + int(width_emu / 9525 // 64) + 1)
        except Exception:
            continue
    return right


def _fit_to_page(ws, width_cols, has_chart):
    """Print / render the styled table on one page width."""
    try:
        from openpyxl.worksheet.properties import PageSetupProperties
        props = ws.sheet_properties
        if props.pageSetUpPr is None:
            props.pageSetUpPr = PageSetupProperties(fitToPage=True)
        elif not props.pageSetUpPr.fitToPage:
            props.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        if not ws.page_setup.orientation and (width_cols > 6 or has_chart):
            ws.page_setup.orientation = 'landscape'
    except Exception:
        pass


def _has_data_bar(ws, rng):
    try:
        for cf in ws.conditional_formatting:
            if str(cf.sqref) == rng and any(getattr(rule, 'type', None) == 'dataBar' for rule in cf.rules):
                return True
    except Exception:
        return False
    return False


def _sheet_signature(ws):
    """Every value/formula of the sheet: the restyle must not change any."""
    out = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                out.append((cell.coordinate, str(cell.value)))
    return out


def _detect_header(ws, min_col, max_col):
    width = max_col - min_col + 1
    for r in range(ws.min_row, min(ws.min_row + 10, ws.max_row + 1)):
        values = [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
        filled = [v for v in values if v is not None and str(v).strip() != '']
        if not filled:
            continue
        texts = [v for v in filled if isinstance(v, str) and not str(v).startswith('=')]
        if len(filled) >= max(2, int(width * 0.5)) and len(texts) >= 0.6 * len(filled):
            return r
    return None


def _label_column(ws, min_col, max_col, start, last_row):
    for c in range(min_col, max_col + 1):
        values = [ws.cell(row=r, column=c).value for r in range(start, last_row + 1)]
        texts = [v for v in values if isinstance(v, str) and v.strip() and not v.startswith('=')]
        if len(texts) >= max(2, int(0.6 * len([v for v in values if v is not None]))):
            return c
    return None


def _autowidth(ws, min_col, max_col, get_column_letter):
    for c in range(min_col, max_col + 1):
        letter = get_column_letter(c)
        longest = 0
        for r in range(ws.min_row, ws.max_row + 1):
            value = ws.cell(row=r, column=c).value
            if value is None:
                continue
            text = str(value)
            longest = max(longest, 12 if text.startswith('=') else len(text))
        width = min(60, max(10, longest + 3))
        current = ws.column_dimensions[letter].width
        if not current or current < width:
            ws.column_dimensions[letter].width = width


def main(argv):
    if len(argv) < 2:
        print(json.dumps({'ok': False, 'error': 'usage: sira_design.py <src> [dst] [theme.json]'}))
        return 2
    src = argv[1]
    dst = argv[2] if len(argv) > 2 and argv[2] not in ('-', '') else None
    theme = argv[3] if len(argv) > 3 else None
    report = restyle(src, dst, theme)
    print(json.dumps(report, ensure_ascii=False))
    return 0 if report.get('ok') else 1


if __name__ == '__main__':
    sys.exit(main(sys.argv))
