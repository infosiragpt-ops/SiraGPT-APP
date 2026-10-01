"""Native Excel charts from an already-open openpyxl worksheet.

No data are generated, coerced, copied or overwritten. The caller chooses
the worksheet, real cell ranges and presentation, then saves its workbook.
"""

import math
import re
from numbers import Real

DEFAULT_COLORS = ('2563EB', '06B6D4', '8B5CF6', '10B981', 'F59E0B', 'E11D48')
CHART_TYPES = ('column', 'bar', 'line', 'area', 'pie', 'doughnut', 'scatter')
_RANGE = re.compile(r'^\$?[A-Za-z]{1,3}\$?[1-9]\d*(?::\$?[A-Za-z]{1,3}\$?[1-9]\d*)?$')


def _palette(colors):
    colors = DEFAULT_COLORS if colors is None else colors
    if not isinstance(colors, (list, tuple)) or not colors:
        raise ValueError('colors must be a nonempty list of RRGGBB colors')
    result = []
    for color in colors:
        value = str(color).lstrip('#')
        if not re.fullmatch(r'[0-9a-fA-F]{6}', value):
            raise ValueError('invalid chart color: %s' % color)
        result.append(value.upper())
    return result


def _bounds(ws, cell_range):
    from openpyxl.utils.cell import range_boundaries
    if not isinstance(cell_range, str) or not _RANGE.fullmatch(cell_range):
        raise ValueError('use a real range on the selected worksheet, e.g. B1:C6')
    c1, r1, c2, r2 = range_boundaries(cell_range)
    if c1 > c2 or r1 > r2 or c2 > min(ws.max_column, 16384) or r2 > min(ws.max_row, 1048576):
        raise ValueError('chart range is reversed or outside the worksheet data')
    if r2 - r1 > 10000 or c2 - c1 >= 32:
        raise ValueError('chart exceeds 10000 data rows or 32 series; select an explicit smaller range')
    return c1, r1, c2, r2


def _numeric(value):
    # Formula references are retained verbatim; Excel calculates their values.
    if isinstance(value, Real) and not isinstance(value, bool):
        try:
            return math.isfinite(value)
        except OverflowError:
            return False
    return isinstance(value, str) and value.startswith('=') and len(value) > 1


def _point_count(series):
    from openpyxl.utils.cell import range_boundaries
    data = getattr(series, 'val', None) or getattr(series, 'yVal', None)
    if data is None:
        return 0
    ref = getattr(data, 'numRef', None)
    if ref is not None and ref.f:
        try:
            c1, r1, c2, r2 = range_boundaries(ref.f.rsplit('!', 1)[-1])
            return (c2 - c1 + 1) * (r2 - r1 + 1)
        except (TypeError, ValueError):
            pass
    cache = getattr(ref, 'numCache', None) if ref is not None else getattr(data, 'numLit', None)
    return int(getattr(cache, 'ptCount', 0) or 0)


def _solid_color(properties, color):
    # DrawingML fill choices are mutually exclusive. Setting solidFill alone
    # leaves an old noFill/gradient/pattern serialized beside it in openpyxl.
    properties.noFill = None
    properties.gradFill = None
    properties.pattFill = None
    properties.solidFill = color


def apply_xlsx_chart_colors(chart, colors=None, *, color_by='auto'):
    """Recolor an existing native chart without changing type, data or refs.

    auto uses points for pie/doughnut and series for other plots. Point-level
    overrides inherit the requested series color when color_by is series.
    """
    from openpyxl.chart import PieChart, DoughnutChart
    from openpyxl.chart.marker import DataPoint
    palette = _palette(colors)
    if color_by not in ('auto', 'series', 'point'):
        raise ValueError('color_by must be auto, series or point')
    count = 0
    for plot in getattr(chart, '_charts', [chart]):
        for series in plot.series:
            color = palette[count % len(palette)]
            by_point = color_by == 'point' or color_by == 'auto' and isinstance(plot, (PieChart, DoughnutChart))
            _solid_color(series.graphicalProperties, color)
            _solid_color(series.graphicalProperties.line, color)
            marker = getattr(series, 'marker', None)
            if marker is not None:
                _solid_color(marker.graphicalProperties, color)
                _solid_color(marker.graphicalProperties.line, color)
            points = {point.idx: point for point in series.data_points}
            if by_point:
                n = _point_count(series)
                if not n or n > 10000:
                    raise ValueError('cannot safely determine chart point count for color_by=point')
                for i in range(n):
                    points.setdefault(i, DataPoint(idx=i))
            for i, point in points.items():
                point_color = palette[i % len(palette)] if by_point else color
                _solid_color(point.graphicalProperties, point_color)
                _solid_color(point.graphicalProperties.line, point_color)
            series.data_points = list(points.values())
            count += 1
    return count


def add_xlsx_chart(ws, *, chart_type, data_range, category_range=None, x_range=None,
                   titles_from_data=True, colors=None, color_by='auto', title=None,
                   legend='r', data_labels=False, anchor='E2', width=16, height=9,
                   style=None):
    """Add one editable native chart; return it without saving the workbook.

    data_range is B1:C6, or aligned column ranges ['B1:B6', 'D1:D6'].
    category_range is A2:A6. Scatter requires x_range with numeric X values.
    Empty Y cells remain gaps; text, invalid ranges and incompatible options
    raise ValueError before a chart is attached. width/height are centimeters.
    data_labels=True shows values; a dict supports showVal, showCatName,
    showSerName, showPercent and showLegendKey booleans.
    """
    from openpyxl.chart import (BarChart, LineChart, AreaChart, PieChart,
                               DoughnutChart, ScatterChart, Reference, Series)
    from openpyxl.chart.label import DataLabelList
    from openpyxl.utils.cell import coordinate_to_tuple
    if chart_type not in CHART_TYPES:
        raise ValueError('unsupported chart_type; use ' + ', '.join(CHART_TYPES))
    if not isinstance(titles_from_data, bool):
        raise ValueError('titles_from_data must be boolean')
    palette = _palette(colors)
    ranges = [data_range] if isinstance(data_range, str) else data_range
    if not isinstance(ranges, (list, tuple)) or not ranges:
        raise ValueError('data_range must name one or more real ranges')
    columns = []
    for value in ranges:
        c1, r1, c2, r2 = _bounds(ws, value)
        columns.extend((c, r1, r2) for c in range(c1, c2 + 1))
    if len(columns) > 32 or len({item[0] for item in columns}) != len(columns):
        raise ValueError('select at most 32 distinct data columns')
    row_bounds = {(r1, r2) for _, r1, r2 in columns}
    if len(row_bounds) != 1:
        raise ValueError('all chart series must use the same source rows')
    r1, r2 = next(iter(row_bounds))
    first_data = r1 + int(titles_from_data)
    if first_data > r2:
        raise ValueError('data_range has no data rows')
    if r2 - first_data + 1 > 10000:
        raise ValueError('chart exceeds 10000 data rows; select an explicit smaller range')
    for column, _, _ in columns:
        if titles_from_data and not str(ws.cell(r1, column).value or '').strip():
            raise ValueError('each series requires an existing header when titles_from_data=True')
        values = [ws.cell(r, column).value for r in range(first_data, r2 + 1)]
        if not any(_numeric(value) for value in values) or any(value is not None and not _numeric(value) for value in values):
            raise ValueError('chart series must contain real numeric cells or formulas, with optional empty gaps')
        if chart_type in ('pie', 'doughnut') and any(value is None or isinstance(value, Real) and value < 0 for value in values):
            raise ValueError('pie/doughnut requires complete nonnegative data')
        if chart_type in ('pie', 'doughnut') and all(isinstance(value, Real) for value in values) and sum(values) <= 0:
            raise ValueError('pie/doughnut requires a positive total, not an empty set of portions')
    if chart_type in ('pie', 'doughnut') and len(columns) != 1:
        raise ValueError('pie/doughnut supports exactly one series; select the intended column')
    if chart_type == 'scatter':
        if not x_range or category_range is not None:
            raise ValueError('scatter requires x_range, not category_range')
        axis_range = x_range
    else:
        if x_range is not None:
            raise ValueError('x_range is only supported for scatter')
        if category_range is None:
            raise ValueError('category_range is required; category labels are never invented')
        axis_range = category_range
    ac1, ar1, ac2, ar2 = _bounds(ws, axis_range)
    if ac1 != ac2 or (ar1, ar2) != (first_data, r2):
        raise ValueError('category/X range must be one column aligned to the data rows, excluding headers')
    axis_values = [ws.cell(r, ac1).value for r in range(ar1, ar2 + 1)]
    if chart_type == 'scatter' and not all(_numeric(value) for value in axis_values):
        raise ValueError('scatter X cells must be numeric or formulas')
    if any(value is None or value == '' for value in axis_values):
        raise ValueError('category/X labels cannot be missing')
    if not isinstance(anchor, str) or ':' in anchor or not _RANGE.fullmatch(anchor):
        raise ValueError('anchor must be a worksheet cell such as E2')
    anchor_row, anchor_column = coordinate_to_tuple(anchor.replace('$', '').upper())
    if anchor_row > 1048576 or anchor_column > 16384:
        raise ValueError('chart anchor is outside Excel worksheet bounds')
    for dimension in (width, height):
        if not isinstance(dimension, Real) or isinstance(dimension, bool) or not math.isfinite(dimension) or not 0 < dimension <= 100:
            raise ValueError('chart width and height must be between 0 and 100 cm')
    if style is not None and (not isinstance(style, int) or isinstance(style, bool) or not 1 <= style <= 48):
        raise ValueError('style must be a native Excel style number from 1 to 48')
    if legend is not None and legend is not False and legend not in ('r', 'l', 't', 'b', 'tr'):
        raise ValueError('legend must be r, l, t, b, tr or False')
    if title is not None and not isinstance(title, str):
        raise ValueError('title must be text')
    labels = {}
    if data_labels is True:
        labels = {'showVal': True}
    elif isinstance(data_labels, dict):
        allowed = {'showVal', 'showCatName', 'showSerName', 'showPercent', 'showLegendKey'}
        if not set(data_labels).issubset(allowed) or any(not isinstance(v, bool) for v in data_labels.values()):
            raise ValueError('unsupported data_labels key or value')
        if data_labels.get('showPercent') and chart_type not in ('pie', 'doughnut'):
            raise ValueError('percentage labels are only supported for pie/doughnut')
        labels = data_labels
    elif data_labels is not False:
        raise ValueError('data_labels must be a boolean or a supported label dictionary')

    classes = {'column': BarChart, 'bar': BarChart, 'line': LineChart, 'area': AreaChart,
               'pie': PieChart, 'doughnut': DoughnutChart, 'scatter': ScatterChart}
    chart = classes[chart_type]()
    if chart_type in ('column', 'bar'):
        chart.type = 'col' if chart_type == 'column' else 'bar'
    if chart_type == 'scatter':
        chart.scatterStyle = 'marker'
    if chart_type == 'line':
        # Some Office renderers interpolate curves when smooth is omitted.
        # A normal line chart connects the supplied observations directly.
        chart.smooth = False
    axis = Reference(ws, min_col=ac1, min_row=ar1, max_row=ar2)
    for column, _, _ in columns:
        values = Reference(ws, min_col=column, min_row=r1, max_row=r2)
        if chart_type == 'scatter':
            series = Series(values, axis, title_from_data=titles_from_data)
            series.marker.symbol = 'circle'
            chart.series.append(series)
        else:
            chart.add_data(values, titles_from_data=titles_from_data)
            if chart_type == 'line':
                chart.series[-1].smooth = False
    if chart_type != 'scatter':
        chart.set_categories(axis)
    chart.display_blanks = 'gap'
    if title is not None:
        chart.title = title
    chart.style = style
    if legend is None or legend is False:
        chart.legend = None
    else:
        chart.legend.position = legend
    if labels:
        chart.dataLabels = DataLabelList(**labels)
    chart.width, chart.height = float(width), float(height)
    apply_xlsx_chart_colors(chart, palette, color_by=color_by)
    ws.add_chart(chart, anchor.replace('$', '').upper())
    return chart
